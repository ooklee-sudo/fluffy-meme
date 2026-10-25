"""Builds the Web Appendix (separate PDF, AMA format) for the JIM submission."""
import os, re, subprocess, json
import pandas as pd, docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__)); OUT = os.path.join(HERE, "..")
RES = os.path.join(HERE, "..", "..", "..", "influencer_replication", "results")
PB = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'
TITLE = "Do Influencers Pay? Limited Attention, Engagement Elasticity, and Simulation-Based Pre-Tests for Seeding Decisions"
toc_pages = json.load(open(os.path.join(HERE, "toc_pages.json"))) if os.path.exists(os.path.join(HERE, "toc_pages.json")) else {"A": 2, "B": 3, "C": 4}

t7 = pd.read_csv(os.path.join(RES, "table7_estimates.csv"))
rows = []
want = [("Pooled, f in [10,1e3)", "Pooled, 10 ≤ f < 10^3^ (selected sample; not interpreted)"), ("abc_reading, f>=1e5", "abc_reading, f ≥ 10^5^"),
        ("electric_toothbrush, f>=1e5", "electric_toothbrush, f ≥ 10^5^"), ("intelligent_floor_scrubber, f>=1e5", "intelligent_floor_scrubber, f ≥ 10^5^"),
        ("ruby_face_cream, f>=1e5", "ruby_face_cream, f ≥ 10^5^"), ("spark_thinking, f>=1e5", "spark_thinking, f ≥ 10^5^"),
        ("supor_boosted_showerhead, f>=1e5", "supor_boosted_showerhead, f ≥ 10^5^"),
        ("Comments only (f>=10)", "Pooled, comments only (f ≥ 10)"), ("All accounts, f floored at 10", "Pooled, all accounts with f set to at least 10"),
        ("Pooled, all accounts, clusters merged across campaigns", "Pooled, clusters merged across campaigns")]
for key, label in want:
    r = t7[t7["sample"] == key].iloc[0]
    rows.append(f"| {label} | {int(r['posts']):,} | {int(r['authors']):,} | {r['beta_ppml']:.2f} ({r['se']:.3f}) [{r['lo']:.2f}, {r['hi']:.2f}] | {r['beta_ztp']:.2f} ({r['se_ztp']:.3f}) [{r['lo_ztp']:.2f}, {r['hi_ztp']:.2f}] |")
w3 = "| Sample | Posts | Authors | β̂ PPML (SE) [95% CI] | β̂ ZTP (SE) [95% CI] |\n|---|---|---|---|---|\n" + "\n".join(rows)
v = pd.read_csv(os.path.join(RES, "tableB1_validation.csv"))
vr = []
for _, r in v.iterrows():
    se = (r["ppml_hi"] - r["ppml_lo"]) / 3.92
    vr.append(f"| {r['true_beta']:g} | {r['ppml']:.3f} ({se:.3f}) [{r['ppml_lo']:.3f}, {r['ppml_hi']:.3f}] | {r['ols_ln1p']:.3f} | {r['fitted_evening_peak']:.1f} hours | {r['heldout_tv']:.3f} ({r['flat_tv']:.3f}) | {r['heldout_r']:.3f} |")
w2 = "| True β | PPML β̂ (SE) [95% CI] | OLS on ln(1 + y) β̂ | Fitted evening peak | Held-out TV (flat profile) | Held-out r |\n|---|---|---|---|---|---|\n" + "\n".join(vr)

md = f"""::: {{custom-style="Title"}}
Web Appendix
:::

**{TITLE}**

*Contents*

Web Appendix A: Simulation Settings and Assumptions ... {toc_pages['A']}

Web Appendix B: Validation of the Estimation Procedure on Synthetic Data ... {toc_pages['B']}

Web Appendix C: Additional Estimates of the Engagement Elasticity ... {toc_pages['C']}

*These materials have been supplied by the authors to aid in the understanding of their paper. The AMA is sharing these materials at the request of the authors.*
{PB}
# Web Appendix A: Simulation Settings and Assumptions

**Table W1. Simulation Parameters and Assumptions**

| Parameter | Value | Remarks |
|---|---|---|
| Stimulus scale κ | 0.16 | Set manually |
| Stimulus decay τ | 1 hour | Single exponential |
| Stimulus cap S~max~ | 3.0 | Synthetic saturation |
| Base exposure probability | 0.5 | Multiplied by (max{{f, 10}}/10)^−β^ |
| Engagement elasticity β | 0 (base); 0 to 0.66 | Applied to all accounts |
| Comment visibility factor | 0.4 | Multiplies exposure probability |
| Feed capacity | 40 items | Most recent items |
| Recency weight | exp(−age/2 hours) | For comment and retweet targets |
| Circadian profile | 0.12 + evening Gaussian at chronotype (SD 2.2 hours) + 0.45 × midday Gaussian at 12.5 hours (SD 1.8 hours) | Normalized to mean one; 15-minute resolution |
| Emotion at each urge evaluation | excited / angry / happy / indifferent with probabilities .20 / .15 / .30 / .35 | Assumption |
| Relevance level | interest < 0.10: 0.25; < 0.30: 0.50; otherwise 0.75 | Assumption |
| Network | 2,000 agents, preferential attachment, 3 follows per new agent | |
| Horizon | 18-hour warm-up; 48 hours after injection | Injection at 18:00 unless stated |
| Distributed seed pool | Accounts outside the top 5% by followers | Random order until the budget is spent |
| Volume-matching factor | 2.00 | Alternative schedulers only; ratio of total events in pilot runs without a campaign |
| Seed price floor | One follower-equivalent | Accounts without followers are priced as one follower |
| Adoption | First comment or retweet of campaign content | Seeds are not counted as adopters |

*Notes:* Likes create no items and no notifications. The code in the supplementary archive reproduces every value.

*Rule-based urge function.* With h the hour of day, leisure = 1 if h ≥ 19 or h < 1, 0.5 if 11 ≤ h < 14, and 0 otherwise; n the number of feed items in the last hour on topics with interest weight ≥ 0.10; ℓ ∈ {{0, 1, 2}} the relevance level of the latest item; ν = 1 for a notification and 0 otherwise; and e an emotion bonus of 0.8 (excited), 0.6 (angry), 0.4 (happy), or −0.6 (indifferent), the rating is I = min(max(1.0 + 1.0 × leisure + 0.4 × min(n, 3) + 0.5 × ℓ + 0.8 × ν + e, 0), 5). For notifications ℓ is set to 1.
{PB}
# Web Appendix B: Validation of the Estimation Procedure on Synthetic Data

To check that the procedure in the main text recovers known parameters, we generated synthetic campaign data with the structure of an influencer-campaign dataset with timestamps: 400 authors with follower counts drawn log-uniformly between 10^5^ and 10^7^, assigned to six products, 15 posts per author over 60 days, author and product random effects (SD 0.8 and 0.3), post-level noise (SD 0.5), log-normal reply delays (median 0.8 hours, log-SD 1.2), and interaction times thinned by a known daily profile with an evening peak at 21:30. Table W2 reports the results for two true values of β. The random-effect scales are our choices, because the structure rather than the scale is what the exercise needs.

**Table W2. Poisson Pseudo-Maximum Likelihood Recovers the True Elasticity, Whereas OLS on the Log Count Does Not**

{w2}

*Notes:* 6,000 posts per row. PPML = Poisson pseudo-maximum likelihood; SE = standard error implied by the confidence interval (clustered by author). TV = total variation distance between the fitted profile from three calibration products and the observed hourly distribution in three held-out products; the value for a flat profile is shown in parentheses. The fitted reply-delay median was 0.80 and 0.79 hours in the two cases, against a generating value of 0.8 hours.

The PPML confidence interval contained the true β in both cases. OLS on the log-transformed count was close to the truth when the true β was 0 but far off when it was 0.35 (0.83), because adding one to the count distorts the relation between interactions and audience size. The daily profile fitted on calibration products reproduced the held-out profile closely in both cases (r = 0.99), far more closely than a flat profile. The synthetic data include posts with no interactions and only accounts with 10^5^ to 10^7^ followers, so this exercise does not test the estimator under the zero truncation and account selection present in SAGraph, which the main text addresses with the zero-truncated model and a discussion of bias. Because the public SAGraph release lacks timestamps, only the elasticity part of the procedure could be applied to real data.
{PB}
# Web Appendix C: Additional Estimates of the Engagement Elasticity

**Table W3. Elasticity Estimates Are Similar Across Robustness Checks, but Imprecise for the Largest Accounts Within a Campaign**

{w3}

*Notes:* Standard errors (SE) are clustered by author within campaign, except in the last row, where accounts that appear in several campaigns (identical description and following count, same order of magnitude of followers) share one cluster. A post is a unique pair of author and post text; the dependent variable counts comments and reposts, except in the comments-only row. For accounts with 10 to 10^3^ followers the estimate is close to 1, but such accounts enter the data only when one of their posts received an interaction, which inflates their measured engagement; we do not interpret this estimate. PPML = Poisson pseudo-maximum likelihood; ZTP = zero-truncated Poisson.
"""
open(os.path.join(HERE, "web_appendix.md"), "w", encoding="utf-8").write(md)
subprocess.run(["pandoc", os.path.join(HERE, "web_appendix.md"), "-o", os.path.join(OUT, "Web_Appendix.docx"), "--reference-doc", os.path.join(HERE, "ref_jim.docx")], check=True)
d = docx.Document(os.path.join(OUT, "Web_Appendix.docx"))
for s in d.sections:
    p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run()
    for typ in ("begin", "instr", "end"):
        if typ == "instr":
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = " PAGE "
        else:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), typ)
        r._r.append(e)
d.core_properties.author = ""; d.core_properties.last_modified_by = ""; d.core_properties.title = "Web Appendix"
d.save(os.path.join(OUT, "Web_Appendix.docx"))
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT, os.path.join(OUT, "Web_Appendix.docx")], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
print("web appendix built")
