"""Small helpers to typeset an ISR-style manuscript with python-docx."""
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Inches(1))
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    # page number in the footer
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)
    r.font.size = Pt(11); r.font.name = FONT
    return doc


TOKEN = re.compile(r"(\*\*[^*]+\*\*|_\{(?:[^{}]|\{[^{}]*\})*\}|\^\{(?:[^{}]|\{[^{}]*\})*\}|\*[^*]+\*)")


def rich(par, text, size=None, bold=False, italic=False):
    """Write text with _{sub}, ^{sup} and *italic* markup into a paragraph."""
    for part in TOKEN.split(text):
        if not part:
            continue
        r = None
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("_{"):
            r = par.add_run(re.sub(r"[_^]\{([^{}]*)\}", r"\1", part[2:-1])); r.font.subscript = True
        elif part.startswith("^{"):
            r = par.add_run(re.sub(r"[_^]\{([^{}]*)\}", r"\1", part[2:-1])); r.font.superscript = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            r = par.add_run(part)
        r.font.name = FONT
        r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        if size:
            r.font.size = Pt(size)
        if bold:
            r.bold = True
        if italic and not (part.startswith("*") and part.endswith("*") and not part.startswith("**")):
            r.italic = True
    return par


def body(doc, text, indent=True, double=True, after=0, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 2.0 if double else 1.0
    pf.space_after = Pt(after)
    if indent:
        pf.first_line_indent = Inches(0.5)
    p.alignment = align
    rich(p, text)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.keep_with_next = True
    rich(p, text, bold=True)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.keep_with_next = True
    rich(p, text, bold=True, italic=True)
    return p


def eq(doc, text, num):
    """Centered equation with a right-aligned number via a tab stop."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.tab_stops.add_tab_stop(Inches(3.25), alignment=1)  # center
    pf.tab_stops.add_tab_stop(Inches(6.5), alignment=2)   # right
    p.add_run("\t")
    rich(p, text, italic=True)
    p.add_run("\t(%s)" % num).font.name = FONT
    return p


def caption(doc, text, bold_label=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    m = re.match(r"^((?:Table|Figure) \d+\.?)(.*)$", text)
    if m and bold_label:
        rich(p, m.group(1), size=11, bold=True)
        rich(p, m.group(2), size=11)
    else:
        rich(p, text, size=11)
    return p


def note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    rich(p, text, size=10)
    return p


def _border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for edge, val in kw.items():
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), val); e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)


def table(doc, header, rows, widths, size=10, align_first_left=True, bold_rows=(), left_all=False):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        pp = c.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if (i == 0 and align_first_left) else WD_ALIGN_PARAGRAPH.CENTER
        rich(pp, h, size=size, bold=True)
        _border(c, top="single", bottom="single")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            pp = cells[i].paragraphs[0]
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ((i == 0 and align_first_left) or left_all) else WD_ALIGN_PARAGRAPH.CENTER
            rich(pp, str(v), size=size, bold=(ri in bold_rows))
            if ri == len(rows) - 1:
                _border(cells[i], bottom="single")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    return t


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
