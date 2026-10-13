"""Cover letter (with statement of contribution) and title page for the ISR submission."""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
FONT = "Times New Roman"
TITLE = "Price Competition and Capacity Investment in LLM Inference Markets: A Dynamic Bertrand–Edgeworth Game"
BR = re.compile(r"(\[[^\]]*\]|\*\*[^*]+\*\*|\*[^*]+\*)")


def new(size=11):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, m, Inches(1))
    st = d.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(0)
    cp = d.core_properties
    cp.author = ""; cp.last_modified_by = ""; cp.title = ""; cp.comments = ""
    return d


def para(d, text="", bold=False, align=None, after=8, size=None, indent=None, spacing=1.08, keep=False):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = spacing
    if keep:
        p.paragraph_format.keep_with_next = True
    if indent is not None:
        p.paragraph_format.left_indent = Inches(indent)
    if align:
        p.alignment = align
    for part in BR.split(text):
        if not part:
            continue
        if part.startswith("[") and part.endswith("]"):
            r = p.add_run(part); r.font.highlight_color = WD_COLOR_INDEX.YELLOW        # placeholders to fill in
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        r.font.name = FONT
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)
    return p


# ====================================== COVER LETTER ======================================
d = new(11)
para(d, "October 13, 2026", after=12)
para(d, "Editor-in-Chief", after=0)
para(d, "*Information Systems Research*", after=0)
para(d, "INFORMS", after=12)
para(d, "**Re: Submission of “" + TITLE + "”**", after=12)
para(d, "Dear Editor,")
para(d, "We submit the enclosed manuscript for consideration as a research article in *Information Systems Research*. The paper studies how providers of large language model "
        "(LLM) inference, who compete on posted prices but must build GPU capacity with a lag and at a sunk cost, choose prices and capacity in a dynamic oligopoly. "
        "It develops a Markov perfect equilibrium model, proves a set of analytical results, and supports them with a systematic numerical analysis of 20 parameter "
        "settings, a controlled experiment on timing, and an exploratory comparison with public data.")
para(d, "Why this matters for information systems. Inference capacity is becoming an infrastructure on which much of the digital economy depends, and a small number "
        "of providers set its price. The IS literature on computing services has analyzed pricing and capacity mostly for a single provider or in static settings "
        "(for example Mendelson 1985, Dewan and Mendelson 1990, Cachon and Harker 2002). The industrial organization literature has the tools for dynamic capacity "
        "games but has not been applied to this market. The paper brings the two together, and it speaks to questions that platform operators and regulators now face: "
        "when competition keeps prices down, whether disclosing capacity helps or hurts, and what portability and switching rules can and cannot achieve.")

para(d, "**Statement of contribution**", after=4, keep=True)
para(d, "**1. A separation result.** When prices are not state variables, pricing and capacity decouple exactly in a logit Markov perfect equilibrium: prices solve a static "
        "game in each state, and capacity solves a dynamic game fed by the resulting profits (Proposition 1). This simplifies the analysis and makes sensitivity analysis feasible.", indent=0.25, after=5)
para(d, "**2. When competition disciplines prices.** With scarce capacity the reservation price is a dominant strategy (Proposition 2); with ample capacity, pure-strategy "
        "markups cannot exceed twice the price-grid step (Proposition 3). In the baseline model, capacity is therefore the only source of market power.", indent=0.25, after=5)
para(d, "**3. Multiplicity and welfare.** Across 20 parameter settings, equilibrium capacity lies between a joint-profit benchmark and a market-share contest. Equilibria are "
        "multiple in 11 of the 20 settings, and the equilibrium that consumers prefer is not the one that maximizes total surplus.", indent=0.25, after=5)
para(d, "**4. Price commitment as a driver of capacity.** In a two-price benchmark with alternating moves, we prove that commitment sustains tacit collusion for intermediate "
        "discount factors (Proposition 4) and makes capacity a strategic complement, which creates a high-capacity equilibrium that simultaneous pricing rules out "
        "(Proposition 5). A controlled timing experiment shows that it is price commitment, and not the timing of capacity decisions, that raises equilibrium capacity "
        "(by 11 to 15 request-units in five of six scenarios, against changes of at most 2.4 from timing alone).", indent=0.25, after=5)
para(d, "**5. Implications.** We draw implications for capacity disclosure, switching and portability rules, and repricing frequency, and we separate those supported by the "
        "computed results from those that are only suggestive.", indent=0.25, after=8)

para(d, "**Scope and limitations.** We want to be transparent about what the paper does not do. The baseline pricing stage is static, so it has no price cycles. In the "
        "dynamic extension, price cycles are a numerical finding that depends on the demand process, and we do not prove them. That extension is solved only at moderate "
        "payoff noise, and some scenarios did not converge, which the manuscript flags. The analytical results for dynamic pricing are for a two-price benchmark with "
        "explicit equilibrium selection. Parameters are illustrative, and only one price ratio is disciplined by public data. An exploratory test of the repricing prediction on "
        "public list prices found a sign opposite to the prediction, statistically insignificant and based on a noisy measure; we report it as inconclusive.")
para(d, "**Fit with the journal.** The paper contributes to the economics of information technology and digital markets, using analytical modeling and computation. "
        "We believe it is of interest to readers who study IT infrastructure, platform competition, and the governance of digital markets.")
para(d, "**Originality and compliance.** The manuscript has not been published and is not under consideration elsewhere. All authors have approved the submission. "
        "[We have no conflicts of interest to declare. / Describe any conflicts.] The title page is provided as a separate file and the manuscript file contains no author information. "
        "The code, parameter files, and result files are provided as an anonymized supplementary archive; [we will make them public on acceptance – confirm]. "
        "[If applicable, insert a statement on the use of AI-assisted tools in preparing the manuscript, code, or analysis, in line with INFORMS policy.]")
para(d, "[Optional: suggested reviewers and any reviewers to exclude, with affiliations and e-mail addresses.]")
para(d, "Thank you for considering our submission. We look forward to hearing from you.", after=12)
para(d, "Sincerely,", after=18)
para(d, "[Corresponding author name]", after=0)
para(d, "[Affiliation]", after=0)
para(d, "ooklee@hanyang.ac.kr", after=0)
para(d, "on behalf of all authors", after=0)
d.save(os.path.join(OUT, "01_Cover_Letter.docx"))

# ====================================== TITLE PAGE ======================================
t = new(11)
para(t, "Title Page (not shared with reviewers)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=14, size=11)
para(t, "**Title**", after=2)
para(t, TITLE, after=12)
para(t, "**Authors and affiliations**", after=2)
para(t, "[Author 1 full name], [Department], [University], [City, Country], [ORCID]", after=2)
para(t, "[Author 2 full name], [Department], [University], [City, Country], [ORCID]", after=2)
para(t, "[Add or delete lines as needed]", after=12)
para(t, "**Corresponding author**", after=2)
para(t, "[Name], [Postal address], ooklee@hanyang.ac.kr, [telephone]", after=12)
para(t, "**Manuscript type**", after=2)
para(t, "Research article (analytical modeling with numerical analysis)", after=12)
para(t, "**Keywords**", after=2)
para(t, "cloud computing; generative AI; capacity investment; Bertrand–Edgeworth competition; Markov perfect equilibrium; dynamic oligopoly; IT infrastructure markets", after=12)
para(t, "**Manuscript length**", after=2)
para(t, "Abstract: 178 words. Main text including appendices, excluding tables, figures, and references: approximately 11,500 words. 9 tables and 3 figures.", after=12)
para(t, "**Acknowledgments**", after=2)
para(t, "[Thank colleagues, seminar participants, and anonymous reviewers of earlier versions, if any. Leave blank if none.]", after=12)
para(t, "**Funding**", after=2)
para(t, "[State funding sources and grant numbers, or “This research received no specific funding.”]", after=12)
para(t, "**Conflicts of interest**", after=2)
para(t, "[None declared. / Describe.]", after=12)
para(t, "**Data and code availability**", after=2)
para(t, "The analysis uses simulated equilibria of the model described in the paper and publicly available price and capital-expenditure information cited in the text. "
        "The code, parameter files, and result files are provided as an anonymized supplementary archive. [State whether and where they will be made public.]", after=12)
para(t, "**Use of AI-assisted tools**", after=2)
para(t, "[Describe, in line with INFORMS policy, any AI-assisted tools used in preparing the manuscript, code, or analysis, and the authors’ responsibility for the content.]", after=12)
para(t, "**Prior versions and presentations**", after=2)
para(t, "[None. / List.]", after=0)
t.save(os.path.join(OUT, "02_Title_Page.docx"))
print("saved cover letter and title page")
