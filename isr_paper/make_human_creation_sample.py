"""Stratified, blinded sample of O*NET task statements for human validation of the visual-content-creation score.
Strata (by LLM score s and rule-based score): A: s=2 (n=45 of 180); B: s=1 (n=55 of 679); C: s=0 and rule-based>0 (n=40, possible LLM misses); D: s=0 and rule-based=0 (n=60, random).
Coder file shows only occupation + task text (shuffled); the key file holds LLM score, rule score, stratum and design weight (N_stratum / n_stratum)."""
import csv, json, random, pandas as pd
import classify_onet as C, rule_creation as R
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
rng=random.Random(20261004)
T=pd.DataFrame(C.load()); S=pd.DataFrame([json.loads(l) for l in open("data/all_creation_sonnet55.jsonl")])[["id","s"]]
t=T.merge(S,on="id"); t["rule"]=t.text.map(R.score)
titles={r["O*NET-SOC Code"]:r["Title"] for r in csv.DictReader(open("data/Occupation_Data.txt",encoding="utf-8"),delimiter="\t")}
t["occupation"]=t.soc.map(titles)
def stratum(r):
    if r.s==2: return "A"
    if r.s==1: return "B"
    return "C" if r.rule>0 else "D"
t["stratum"]=t.apply(stratum,axis=1)
N=t.stratum.value_counts().to_dict(); n={"A":45,"B":55,"C":40,"D":60}; print("population strata:",N)
parts=[]
for k,m in n.items():
    d=t[t.stratum==k]; take=d.sample(min(m,len(d)),random_state=rng.randint(0,10**6)).copy(); take["weight"]=N[k]/len(take); parts.append(take)
smp=pd.concat(parts).sample(frac=1,random_state=7).reset_index(drop=True)
smp["row"]=range(1,len(smp)+1)
smp[["row","id","soc","occupation","text","s","rule","stratum","weight"]].to_csv("data/human_creation_sample_key.csv",index=False)
wb=Workbook(); ws=wb.active; ws.title="coding"
ws.append(["row","occupation","task","score_0_1_2","notes"])
for _,r in smp.iterrows(): ws.append([int(r.row),r.occupation,r.text,None,None])
for c,w in zip("ABCDE",[7,40,100,14,32]): ws.column_dimensions[c].width=w
for row in ws.iter_rows():
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
for c in ws[1]: c.font=Font(bold=True)
for r in range(2,len(smp)+2): ws.cell(r,4).fill=PatternFill("solid",fgColor="FFF2CC")
dv=DataValidation(type="list",formula1='"0,1,2"',allow_blank=True,showErrorMessage=True,error="0, 1 또는 2만 입력"); ws.add_data_validation(dv); dv.add(f"D2:D{len(smp)+1}"); ws.freeze_panes="A2"
g=wb.create_sheet("guide")
for l in ["질문: 이 과업을 수행할 때 이미지·그래픽·사진·영상·애니메이션·레이아웃 같은 시각 콘텐츠를 만들거나 편집하는가? (과업 문장과 직업명만 보고 판단)",
 "2 = 시각 콘텐츠를 만들거나 편집하는 것이 과업의 핵심 (예: 그래픽·일러스트 디자인, 영상 편집, 사진 촬영·보정, 애니메이션, 시각 광고물 제작)",
 "1 = 시각 콘텐츠가 과업의 일부이거나 보조 산출물 (예: 이미지 선택·승인, 시각 제작 감독, 발표 자료·도표 작성, 더 큰 업무 중 단순 시각물 제작)",
 "0 = 시각 콘텐츠를 만들거나 편집하지 않음 (글쓰기, 데이터 분석, 관리, 영업, 기계 조작, 돌봄 등)",
 "주의: (1) 보거나 검사하거나 읽는 것은 점수에 포함하지 않음 (2) 글쓰기와 음향 제작은 포함하지 않음 (3) 기술 도면·설계도는 도면 자체를 그리는 과업일 때만 인정",
 "노란 칸(D열)에 0/1/2를 고르고, 애매하면 notes(E열)에 메모하세요. 모델 점수는 보지 마세요. 코더가 두 명이면 서로 상의하지 말고 각자 파일을 채우세요.",
 "완료 후 파일 이름을 human_creation_sample_filled_coderA.xlsx (코더 B는 _coderB)로 저장해 올려 주세요."]: g.append([l])
g.column_dimensions["A"].width=160
for r in g.iter_rows():
    for c in r: c.alignment=Alignment(wrap_text=True,vertical="top")
wb.save("data/human_creation_sample.xlsx"); print("sample size",len(smp),smp.stratum.value_counts().to_dict(),"| weights",smp.groupby("stratum").weight.first().round(1).to_dict())
