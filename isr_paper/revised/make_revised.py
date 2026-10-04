"""Builds the Europe/Canada-only revision of the uploaded manuscript (UK data removed).
usage: python revised/make_revised.py <uploaded.docx> <out.docx>"""
import sys, re, copy, docx
from docx.oxml.ns import qn
src,out=sys.argv[1],sys.argv[2]
d=docx.Document(src); P=list(d.paragraphs); orig={i:p.text for i,p in enumerate(P)}
def base_rpr(p):
    r=p.runs[0]._r; rp=r.find(qn("w:rPr")); return copy.deepcopy(rp) if rp is not None else None
def settext(p,text):
    rp=base_rpr(p) if p.runs else None
    for r in list(p.runs): r._r.getparent().remove(r._r)
    for part in re.split(r"(\{i\}.*?\{/i\}|\{b\}.*?\{/b\})",text):
        if not part: continue
        it=part.startswith("{i}"); bd=part.startswith("{b}"); t=re.sub(r"\{/?[ib]\}","",part)
        r=p.add_run(t)
        if rp is not None:
            old=r._r.find(qn("w:rPr"))
            if old is not None: r._r.remove(old)
            nr=copy.deepcopy(rp); 
            for tag in ("w:i","w:b"):
                e=nr.find(qn(tag)); 
                if e is not None: nr.remove(e)
            r._r.insert(0,nr)
        if it: r.italic=True
        if bd: r.bold=True
def sub(i,old,new):
    t=orig[i]; assert old in t,(i,old); settext(P[i],t.replace(old,new))
def delete(i): e=P[i]._p; e.getparent().remove(e)
def cell(c,text):
    p=c.paragraphs[0]; rp=base_rpr(p) if p.runs else None
    for r in list(p.runs): r._r.getparent().remove(r._r)
    r=p.add_run(text)
    if rp is not None:
        old=r._r.find(qn("w:rPr")); 
        if old is not None: r._r.remove(old)
        r._r.insert(0,copy.deepcopy(rp))
def setrow(t,ri,vals):
    for c,v in zip(t.rows[ri].cells,vals): cell(c,v)
# ---------- abstract
settext(P[5],"Statistics showing that firms use generative visual AI are increasingly read as evidence that creative work is being displaced. Information systems research separates acquiring an innovation from assimilating it into work. We model a firm that reassigns a visual task to a generator only when output acceptance exceeds a threshold set by generation, review, and human production costs. Review cost depends on organizational complements, so near-free acquisition can run ahead of assimilation. Eurostat data show that 9.6 percent of EU firms use AI that generates pictures, video, or audio, and that firms which considered AI cite missing expertise (70 percent) and legal uncertainty (54 percent) far more than lack of usefulness (18 percent). A detectability bound shows that European and Canadian job-advert designs could not detect even complete assimilation at platform-sized effects; the small gradients they find are no larger than earlier trends. We specify the firm-level evidence needed.")
# ---------- introduction
sub(11,"while the UK design cannot detect a reduction smaller than about 30 percent for that occupation. The exposure gradients that the designs do report are present before generative tools were released and are larger than any plausible assimilation effect. They reflect something other than the technology.","while the European design cannot detect a reduction smaller than about 24 percent for that occupation. The exposure gradients that the designs do report are small and imprecise, about 1 to 2 percent fewer adverts per standard deviation of exposure in 2023 and 2024, and the European gradient was already moving in the same direction before generative tools were released.")
sub(11,"roughly 2 to 18 percent","roughly 2 to 17 percent") if False else None
t=orig[11].replace("roughly 2 to 18 percent","roughly 2 to 17 percent").replace("while the UK design cannot detect a reduction smaller than about 30 percent for that occupation. The exposure gradients that the designs do report are present before generative tools were released and are larger than any plausible assimilation effect. They reflect something other than the technology.","while the European design cannot detect a reduction smaller than about 24 percent for that occupation. The exposure gradients that the designs do report are small and imprecise, about 1 to 2 percent fewer adverts per standard deviation of exposure in 2023 and 2024, and the European gradient was already moving in the same direction before generative tools were released.")
settext(P[11],t)
sub(13,"and it shows that a common empirical strategy currently attributes to the technology declines that the technology cannot have produced.","and it shows how the bound changes the reading of exposure gradients: the 2023 European gradient is no larger than the trend that preceded the technology, and the 2024 gradient is too imprecise to separate the scenarios.")
# ---------- data
settext(P[50],"Because the acquisition measure covers EU firms, the hiring series is Eurostat’s experimental online job advertisements by ISCO-08 three-digit occupation, 2019–2024, with 2022 as the base year (604 country-occupation cells, 96 occupations, 16 countries, after a floor of 1,000 base-year adverts). Its population matches the acquisition data, but it offers only two post-2022 years. The Canadian series is the Job Vacancy and Wage Survey by NOC 2021 unit group, 2015–2026, restricted to 308 unit groups with at least 1,000 vacancies in the base window; the quarter 2020Q2 was not collected, so 2020 and 2021 cannot be used. European Labour Force Survey employment by industry and US sector postings from Indeed Hiring Lab are used descriptively.")
t=orig[52].replace("mapped to UK occupations through ISCO-08","mapped to European occupations through ISCO-08 (three-digit prefix of the ESCO crosswalk)")
i0=t.index("Low reliability"); t=t[:i0]+"With the broad measure the European coefficients are −0.009 in 2023 and −0.004 in 2024 (standard errors 0.007 and 0.015). Low reliability attenuates exposure coefficients toward zero, which lowers power further and so strengthens the detectability result; we do not correct for it."
settext(P[52],t)
# ---------- theory rule (iii): define gamma**
t=orig[44]; assert "γ**" in t
# paragraph 44 has italics in several runs; keep as plain text with gamma markers
t=orig[44].replace("γ**","γ**")
settext(P[44],"(i) If |{i}γ{/i}*| is below the minimum detectable effect, the design is uninformative about assimilation, whatever the estimated coefficient. A null result does not support the gap, and a negative result is not evidence of substitution. (ii) If |{i}γ{/i}*| exceeds the minimum detectable effect and the estimate is significantly smaller in magnitude than {i}γ{/i}*, the data reject complete assimilation with net substitution but do not distinguish the gap from offsetting scale effects. (iii) If the estimate is significantly larger in magnitude than the largest plausible {i}γ{/i}**, which we take to be the coefficient implied by an effect of {i}β{/i} = −0.5 at the acquisition rate of the highest-acquisition sector, or if the gradient is already present before the technology was available, the decline is attributed to other causes and counts as evidence for none of the accounts.")

# ---------- theory extension: bottleneck, decomposition, attrition (new Sections 3.4-3.6, Appendix C)
from docx.text.paragraph import Paragraph as _Par
def clone_after(anchor_p,model_p,text):
    e=copy.deepcopy(model_p._p); anchor_p._p.addnext(e); q=_Par(e,anchor_p._parent); settext(q,text); return q
def reorig(p): return p
anchor=P[34]
H=P[31]; BODY=P[32]; EQ=P[33]
a=clone_after(anchor,H,"3.4 Review as a Bottleneck Task")
b=clone_after(a,BODY,"Equation (1) says when a firm assimilates a task. To see what assimilation does to labor demand, let a unit of the visual asset require {i}h{/i} hours of human production, so that {i}θ{/i} = {i}wh{/i}, and let review require {i}u{/i} hours, so that {i}r{/i} = {i}wu{/i}. Write {i}ρ{/i} = {i}r{/i}/{i}θ{/i} for review cost relative to human production and {i}γ{/i} = {i}c{/i}/{i}θ{/i} for generation cost relative to human production, so that {i}a{/i}* = {i}γ{/i} + {i}ρ{/i}. With the generator, a unit of output needs review labor and, with probability 1 − {i}a{/i}, human production labor, so labor per unit falls from {i}h{/i} to {i}u{/i} + (1 − {i}a{/i}){i}h{/i}, a ratio of 1 − {i}a{/i} + {i}ρ{/i}, while cost per unit falls by the ratio {i}κ{/i} = {i}γ{/i} + {i}ρ{/i} + 1 − {i}a{/i}. The ratio {i}κ{/i} is below one exactly when {i}a{/i} > {i}a{/i}*. Review labor is paid the same wage as production labor; relaxing this changes {i}ρ{/i} but not the argument.")
c=clone_after(b,BODY,"{b}Proposition 1 (review floor).{/b} Labor per unit of output at an assimilating firm cannot fall below the review share {i}ρ{/i}, however high the acceptance rate {i}a{/i}. If review hours do not fall with model quality, improvements in the generator reduce labor per unit only down to this floor, and when {i}ρ{/i} is large relative to {i}a{/i} the labor saved is small even though the task has been assimilated.")
d_=clone_after(c,BODY,"The floor is a bottleneck in the sense of Baumol (1967) and Aghion, Jones, and Jones (2019): when a task that can be automated is complementary to one that cannot, the task that cannot be automated becomes the binding constraint, and its share of the remaining labor rises as the generator improves. Because {i}ρ{/i} depends on the complements {i}K{/i} through review cost, the same complements that govern assimilation also govern how much labor assimilation can save.")
e_=clone_after(d_,H,"3.5 Displacement, Reinstatement, and Scale")
f_=clone_after(e_,BODY,"Equation (2) leaves {i}β{/i} unspecified. Let {i}ε{/i} be the price elasticity of demand for the occupation’s visual output at the assimilating firm, so that output rises by {i}κ{/i}^(−{i}ε{/i}) when cost per unit falls by the ratio {i}κ{/i}. The net effect on demand for the occupation’s work at an assimilating employer is")
g_=clone_after(f_,EQ,"{i}β{/i}  =  ln(1 − {i}a{/i} + {i}ρ{/i})  −  {i}ε{/i} ln {i}κ{/i}.          (3)")
h_=clone_after(g_,BODY,"The first term combines displacement, because the generator takes over a share {i}a{/i} of human production, with reinstatement, because review and direction create labor {i}ρ{/i} that did not exist before; these are the displacement and reinstatement effects of the task-based framework (Acemoglu and Restrepo 2018, 2019). The second term is the productivity or scale effect, which is positive because cheaper output raises the quantity demanded and with it the demand for review and direction. The values {i}β{/i} = −0.19 and {i}β{/i} = −0.5 used below are therefore reduced-form: each bundles {i}a{/i}, {i}ρ{/i}, {i}γ{/i}, and {i}ε{/i}. For example, with {i}a{/i} = 0.8, {i}ρ{/i} = 0.3, and {i}γ{/i} = 0.05, labor per unit falls to half and cost per unit to 0.55 of its earlier level, and {i}β{/i} is −0.39 if {i}ε{/i} = 0.5, −0.10 if {i}ε{/i} = 1, and +0.02 if {i}ε{/i} = 1.2.")
i_=clone_after(h_,BODY,"{b}Proposition 2 (scale offset).{/b} The net effect {i}β{/i} is negative if and only if {i}ε{/i} < {i}ε{/i}* = ln(1 − {i}a{/i} + {i}ρ{/i}) / ln {i}κ{/i}, and {i}ε{/i}* > 1 whenever {i}γ{/i} > 0. As generation cost falls toward zero, {i}ε{/i}* approaches 1: complete assimilation reduces demand for the occupation if demand for visual output is inelastic and can raise it only if demand is elastic.")
j_=clone_after(i_,H,"3.6 Adjustment Through Attrition")
k_=clone_after(j_,BODY,"Job adverts measure hiring, and hiring is how most employers reduce a workforce without layoffs. Let {i}L{/i}(t) be an employer’s employment in the occupation, {i}δ{/i} the separation rate, and {i}H{/i}(t) hires, so that {i}L{/i}(t + 1) = (1 − {i}δ{/i}){i}L{/i}(t) + {i}H{/i}(t). Suppose assimilation lowers target employment to {i}L{/i}* = (1 − {i}f{/i}){i}L{/i}₀ and the employer closes a fraction {i}λ{/i} ≤ {i}δ{/i} of the remaining gap each period by adjusting hiring:")
l_=clone_after(k_,EQ,"{i}H{/i}(t)  =  {i}δL{/i}(t)  +  {i}λ{/i}({i}L{/i}*  −  {i}L{/i}(t)).          (4)")
m_=clone_after(l_,BODY,"{b}Proposition 3 (attrition lag).{/b} Under equation (4), (i) hiring in the first period falls by the fraction ({i}λ{/i}/{i}δ{/i}){i}f{/i} ≤ {i}f{/i} and converges to a fall of {i}f{/i}, so the decline in adverts never exceeds the decline in target employment, and equation (5) below, which gives the coefficient implied by {i}f{/i}, is an upper bound for adverts as well as for employment; (ii) employment closes only a fraction 1 − (1 − {i}λ{/i})^{i}T{/i} of the gap after {i}T{/i} periods. With {i}λ{/i} = 0.10, for example, 19 percent of the gap is closed after two years and 41 percent after five.")
n_=clone_after(m_,BODY,"Two consequences follow. Employment series, such as those in Section 7.3, can show no contraction for years after assimilation has begun, so adverts are the earlier signal, though not a larger one. And a decline that began before the technology, because of structural change in the industries that employ the occupation, enters as a lower target employment at an earlier date and follows the same dynamics. A pre-existing trend in adverts and a target cut caused by the technology cannot be told apart without firm-level data that observe use of the tool.")
# old equation (3) -> (5), old (3)->(5) label paragraph and renamed subsection
settext(P[42],"{i}γ{/i}*  ≈  {i}A{/i}ᵘ · {i}β{/i} / {i}k{/i}.          (5)")
settext(P[35],"3.7 Competing Accounts")
# literature, introduction
sub(20,"Exposure measures describe where a technology could matter, not where it has been acquired or assimilated.","Exposure measures describe where a technology could matter, not where it has been acquired or assimilated. Task-based models separate displacement, reinstatement, and productivity effects of automation (Acemoglu and Restrepo 2018, 2019), and models of unbalanced growth show that tasks that cannot be automated become bottlenecks (Baumol 1967; Aghion, Jones, and Jones 2019); we use both ideas in Section 3.")

# ---------- section 6.3 (UK depth) removed
for i in (66,65,64): delete(i)
# ---------- section 7
settext(P[69],"For each country-occupation cell in 16 European countries we regress the change in log online adverts relative to 2022 on standardized visual-creation exposure and standardized language-model exposure, with country-year fixed effects and standard errors clustered by ISCO occupation (Table 4). A standard deviation of exposure is associated with 1.6 percent fewer adverts in 2023 (standard error 0.6, randomization p = 0.016) and 1.2 percent fewer in 2024 (standard error 1.7, p = 0.19). Without the language control the estimates are −3.7 and −5.9 percent, because exposure to visual content creation and exposure to language models are positively correlated across occupations (0.27).")
settext(P[70],"Table 4. Europe: log-advert change per standard deviation of visual-creation exposure, relative to 2022")
settext(P[71],"Notes. 604 country-occupation cells (96 occupations, 16 countries); country-year fixed effects; standard errors clustered by ISCO three-digit occupation in parentheses. The deviation row estimates a linear trend in the exposure gradient on 2019–2021 (−0.0072 per year, standard error 0.0047) and reports the deviation of 2023 and 2024 from its extrapolation. Minimum detectable effect = 2.8 × standard error (two-sided 5 percent test, 80 percent power) of the preferred specification.")
settext(P[72],"The coefficients before 2022 are positive (+0.024, +0.008, and +0.011 for 2019, 2020, and 2021), which means that exposed occupations were already declining relative to others before generative tools were released. The trend is weak (−0.007 per year, standard error 0.005), but it runs in the same direction as the post-2022 coefficients: relative to it, the 2023 and 2024 coefficients are −0.009 (standard error 0.007) and +0.002 (0.023). The series has only two post-2022 years. In Canada, within NOC two-digit groups, the coefficient is −2.2 percent in the first post-2022 year (standard error 0.9, randomization p = 0.10) and indistinguishable from zero afterward (−0.9, +1.6, and +1.1 percent in 2024, 2025, and 2026, with standard errors of 2.5, 2.7, and 3.0), in a market where total vacancies fell 47 percent.")
settext(P[74],"In the European exposure distribution the most exposed occupations are creative and performing artists (ISCO 265, 6.2 standard deviations above the mean) and architects, planners, surveyors, and designers (ISCO 216, which includes graphic designers, 5.7 standard deviations). We set {i}k{/i} = 5.7, the value more favorable to detection. Table 5 reports the coefficient per standard deviation that complete assimilation would produce under equation (5) for a range of acquisition rates and substitution effects, alongside the minimum detectable effect and the observed coefficients.")
settext(P[75],"Table 5. Coefficient per standard deviation implied by complete assimilation (Aᵃ = Aᵘ), Europe, k = 5.7")
settext(P[76],"Notes. γ* = Aᵘ·β/k. The upper bound Aᵃ = Aᵘ maximizes the implied response; any assimilation gap makes it smaller. Acquisition rates: Eurostat 2025 (all enterprises; publishing, film, television, and music) and Eurostat 2023 text generation as a proxy for early acquisition. The minimum detectable effect and the observed coefficients are for 2024, the year with the least power; the last column multiplies by k.")
settext(P[77],"Two results follow. First, every complete-assimilation scenario for 2024 implies a coefficient below the minimum detectable effect of 0.048; the largest, −0.032, requires a 39 percent fall among assimilating employers in a sector where 36 percent have acquired the tool, and assimilation by all of them. Under rule (i), the 2024 European estimate is uninformative about whether acquisition has been assimilated, and its small and imprecise coefficient supports neither the gap nor substitution. The preferred estimate of −0.012 equals what the platform effect would produce if every employer in a high-acquisition sector had assimilated, but it is 0.7 standard errors from zero. The estimate without the language control (−0.059) is larger than any scenario, but it is not significantly larger than {i}γ{/i}** (−0.032; z = 1.3) and, as the difference between the specifications shows, partly reflects exposure to language models. Second, the 2023 estimate has more power (minimum detectable effect 0.016). Text-generation acquisition stood at 2.1 percent in 2023. If acquisition of visual generators was similarly low, complete assimilation with {i}β{/i} = −0.5 implies a coefficient of about −0.002, and the observed −0.016 differs from it significantly (z = 2.6); if acquisition had already reached its 2025 level, the implied coefficient would be −0.008 and the difference would not be significant. Either way, the 2023 gradient is of the size predicted by the trend that preceded the technology, since relative to that trend it is −0.009 (standard error 0.007). Under rule (iii), the gradient is better attributed to other causes than to generative visual AI.")
settext(P[78],"The Canadian series has standard errors of 0.025 to 0.030 per standard deviation after 2023, so its minimum detectable effect is about 0.07 to 0.08 per standard deviation. For the most exposed unit group, 9.0 standard deviations above the mean, this corresponds to a log change of about 0.6 to 0.8. The Canadian series is therefore uninformative about assimilation at any plausible acquisition rate, and its first-year coefficient, which does not persist, leads to the same classification.")
# 7.3 quality requirements removed (UK-only analysis); renumber
delete(80); delete(79)
settext(P[81],"7.3 Employment Where Acquisition Is Highest")
settext(P[83],"7.4 Acquisition Interaction in Europe")
settext(P[84],"A European dose-response test interacts occupational exposure with country-level acquisition of picture, video, or audio generation (Eurostat 2025), which is available for 12 of the 16 countries. Complete assimilation predicts a steeper exposure gradient where acquisition is higher, whereas the gap and offsetting-scale accounts predict no such interaction. The interaction coefficient per standard deviation of acquisition (5.1 percentage points) is −0.008 (standard error 0.008) in 2023 and +0.003 (0.016) in 2024; permutation p-values across countries are 0.46 and 0.85. In 2021, before the base year, it is +0.008 (0.007). Complete assimilation with {i}β{/i} = −0.5 predicts an interaction of only about −0.004 (−0.5 × 0.051 / 5.7), far smaller than the standard errors, so the test cannot discriminate among the accounts. We report it because it is the only test here that uses variation in acquisition directly. It becomes informative only with more countries, more post-2022 years, and acquisition measured before the outcome. The detectability result does not depend on it, because that result depends only on acquisition rates, reference effect sizes, and standard errors.")
# ---------- section 8
settext(P[86],"Equation (5) also tells us when the European design would gain the power to detect complete assimilation. Holding the 2024 standard error fixed, detection for an occupation at {i}k{/i} = 5.7 requires {i}A{/i}ᵃ · {i}β{/i} ≤ −0.28. With {i}β{/i} = −0.19 this cannot occur even at full assimilation. With {i}β{/i} = −0.5 it requires {i}A{/i}ᵃ ≥ 0.55; with {i}β{/i} = −0.8 (a 55 percent fall among assimilators), {i}A{/i}ᵃ ≥ 0.34. Table 6 reports the first year in which acquisition would reach these shares if its log-odds rose at the 2023–2025 pace of text-generation AI (about 0.75 per year) or at two-thirds or one-third of that pace. Because {i}A{/i}ᵃ ≤ {i}A{/i}ᵘ, these are the earliest possible detection dates; any persistent assimilation gap pushes them later. The standard error will change as post-2022 years accumulate, so the dates are conditional on the precision of the current design.")
settext(P[87],"Table 6. Earliest year in which complete assimilation would become detectable, European design")
settext(P[88],"Notes. Paths start in 2025 at 9.6 percent (all firms) or 36.2 percent (publishing, film, and broadcasting) with log-odds slopes of 0.75 (fast), 0.50 (medium), and 0.25 (slow) per year. Detection threshold: 2.8 × the 2024 standard error (0.017) × k = 5.7 ≈ 0.28. Each entry is the year in which the required share is first reached. Dates assume Aᵃ = Aᵘ and are lower bounds.")
sub(89,"The media entries for β = −0.8 show that only very large effects concentrated in high-acquisition sectors could already be within reach, and even these assume that every acquiring employer has assimilated the tool.","The media entries for β = −0.8 are already within reach in 2025, because that sector’s acquisition rate exceeds the required share. They show that only very large effects concentrated in high-acquisition sectors could already be detectable, and even these assume that every acquiring employer has assimilated the tool.")
# ---------- firm-level design, discussion, limitations, conclusion
t=orig[91]; a=t.index("In the United Kingdom"); b=t.index("Eurostat ICT-usage microdata"); settext(P[91],t[:a]+t[b:])
sub(96,"and that gradients they do detect can be too large, and too early, to be caused by the technology.","and that gradients they do detect can be no larger than a trend that preceded the technology.")


# ---------- contributions sentence, references, Appendix C
sub(13,"It extends assimilation theory to generative AI by identifying review cost as the place where complementary assets enter, and by explaining why near-zero acquisition costs should make assimilation gaps wider for generative tools than for the enterprise systems on which the theory was built.","It extends assimilation theory to generative AI by identifying review cost as the place where complementary assets enter, and by explaining why near-zero acquisition costs should make assimilation gaps wider for generative tools than for the enterprise systems on which the theory was built. Embedding review cost in a task-based model of labor demand shows that review sets a floor on the labor that assimilation can save, that cheaper generation raises employment only if demand for visual output is elastic, and that employers who adjust through attrition make adverts the earlier and employment the later signal.")
def addref(after_p,parts):
    e=copy.deepcopy(after_p._p); after_p._p.addnext(e); q=_Par(e,after_p._parent)
    rp=q.runs[0]._r.find(qn("w:rPr")); rp=copy.deepcopy(rp)
    for r in list(q.runs): r._r.getparent().remove(r._r)
    for t,it in parts:
        r=q.add_run(t); old=r._r.find(qn("w:rPr"))
        if old is not None: r._r.remove(old)
        nr=copy.deepcopy(rp); ie=nr.find(qn("w:i"))
        if ie is not None: nr.remove(ie)
        r._r.insert(0,nr); r.italic=True if it else None
    return q
x=addref(P[112],[("Acemoglu, D., and Restrepo, P. 2018. The race between man and machine: Implications of technology for growth, factor shares, and employment. ",0),("American Economic Review",1),(" 108(6):1488–1542.",0)])
y=addref(x,[("Acemoglu, D., and Restrepo, P. 2019. Automation and new tasks: How technology displaces and reinstates labor. ",0),("Journal of Economic Perspectives",1),(" 33(2):3–30.",0)])
addref(y,[("Aghion, P., Jones, B. F., and Jones, C. I. 2019. Artificial intelligence and economic growth. In ",0),("The Economics of Artificial Intelligence: An Agenda",1),(", edited by A. Agrawal, J. Gans, and A. Goldfarb. Chicago: University of Chicago Press.",0)])
addref(P[115],[("Baumol, W. J. 1967. Macroeconomics of unbalanced growth: The anatomy of urban crisis. ",0),("American Economic Review",1),(" 57(3):415–426.",0)])
hc=clone_after(P[110],P[109],"Appendix C. Derivations for Section 3")
c1=clone_after(hc,P[110].__class__(P[32]._p,P[32]._parent) if False else P[32],"{b}Labor and cost ratios.{/b} Human production of a unit takes {i}h{/i} hours. With the generator, a unit requires {i}u{/i} hours of review and, with probability 1 − {i}a{/i}, {i}h{/i} hours of human production, so expected labor per unit is {i}u{/i} + (1 − {i}a{/i}){i}h{/i} = {i}h{/i}({i}ρ{/i} + 1 − {i}a{/i}) with {i}ρ{/i} = {i}u{/i}/{i}h{/i} = {i}r{/i}/{i}θ{/i} at equal wages. Expected cost per unit is {i}c{/i} + {i}r{/i} + (1 − {i}a{/i}){i}θ{/i} = {i}θκ{/i}. Output scales by {i}κ{/i}^(−{i}ε{/i}) under constant elasticity {i}ε{/i}, so labor demand scales by (1 − {i}a{/i} + {i}ρ{/i}){i}κ{/i}^(−{i}ε{/i}), which gives equation (3) after taking logarithms.")
c2=clone_after(c1,P[32],"{b}Proposition 2.{/b} {i}β{/i} < 0 if and only if ln(1 − {i}a{/i} + {i}ρ{/i}) < {i}ε{/i} ln {i}κ{/i}. When {i}γ{/i} > 0 and {i}a{/i} > {i}a{/i}*, 1 − {i}a{/i} + {i}ρ{/i} < {i}κ{/i} < 1, so both logarithms are negative and |ln(1 − {i}a{/i} + {i}ρ{/i})| > |ln {i}κ{/i}|. Dividing by ln {i}κ{/i} reverses the inequality: {i}β{/i} < 0 if and only if {i}ε{/i} < {i}ε{/i}* = ln(1 − {i}a{/i} + {i}ρ{/i}) / ln {i}κ{/i}, and {i}ε{/i}* > 1. As {i}γ{/i} → 0, {i}κ{/i} → 1 − {i}a{/i} + {i}ρ{/i} and {i}ε{/i}* → 1.")
c3=clone_after(c2,P[32],"{b}Proposition 3.{/b} Substituting equation (4) into the employment equation gives {i}L{/i}(t + 1) − {i}L{/i}* = (1 − {i}λ{/i})({i}L{/i}(t) − {i}L{/i}*), so {i}L{/i}(t) − {i}L{/i}* = (1 − {i}λ{/i})^{i}t{/i} ({i}L{/i}₀ − {i}L{/i}*) and the share of the gap closed after {i}T{/i} periods is 1 − (1 − {i}λ{/i})^{i}T{/i}. Hires in the first period are {i}H{/i}(0) = {i}δL{/i}₀ − {i}λfL{/i}₀, so {i}H{/i}(0)/({i}δL{/i}₀) = 1 − ({i}λ{/i}/{i}δ{/i}){i}f{/i}. As {i}t{/i} → ∞, {i}L{/i}(t) → {i}L{/i}* and {i}H{/i}(t) → {i}δL{/i}*, so {i}H{/i}(t)/({i}δL{/i}₀) → 1 − {i}f{/i}. Hires are non-negative in every period when {i}λ{/i} ≤ {i}δ{/i}, because {i}λf{/i} ≤ {i}δ{/i} for {i}f{/i} ≤ 1. The upper bound in part (i) holds when hiring is adjusted smoothly as in equation (4); a temporary hiring freeze could produce a larger transitory fall in adverts.")


# ---------- 7.5 sector-level acquisition and employment
h75=clone_after(P[84],P[83],"7.5 Sector-Level Acquisition and Employment")
p75=clone_after(h75,P[32],"Firm-level links between use and hiring are not public, so we use sectors as the nearest proxy. In 260 country-sector cells (30 countries, 9 NACE sections), sectors with higher acquisition of picture, video, or audio generation in 2025 had higher employment growth between 2019–2021 and 2023–2025: 0.051 log points per standard deviation of acquisition (10.6 percentage points), with country fixed effects (standard error 0.010). The same relationship holds before the technology: for 2016–2018 to 2019–2021 the coefficient is 0.049 (0.008), and the difference between the two periods is 0.002 (0.012). Acquisition is not a separate signal here, because it correlates 0.92 with acquisition of text generation and 0.95 with use of any AI across cells; adding the text-generation control leaves a post-minus-pre difference of 0.041 (0.020), which is imprecise because of this collinearity. In the United States, the sector share of businesses reporting AI use in the Business Trends and Outlook Survey correlates 0.59 across 20 sectors with the sector’s reported employee-change index (descriptive only). Complete assimilation with {i}β{/i} = −0.5 would imply a coefficient of about −0.053 times {i}s{/i}, the share of sector employment in visual-production occupations, per standard deviation of acquisition. Because {i}s{/i} is well below one, the implied effect is far smaller than the standard errors. Sector-level data therefore show that acquisition marks growing, digitally intensive sectors, not that it displaces labor, and they cannot reveal the assimilation margin; this is the limit that firm-level data would remove.")
addref(P[134],[("U.S. Census Bureau. Business Trends and Outlook Survey (BTOS), sector estimates. Public-use tables.",0)])

# ---------- structural-change paragraph (after paragraph 96)
import copy as _c
newp=_c.deepcopy(P[96]._p); P[96]._p.addnext(newp)
from docx.text.paragraph import Paragraph
np_=Paragraph(newp,P[96]._parent)
settext(np_,"A related reading is that falling demand for visual-production occupations reflects structural change in industries, such as shifts in how marketing and media content is produced and bought, and that generative AI accelerates change already under way rather than starting it. Our evidence is compatible with this reading. The European gradient was already moving in the same direction before 2022 (−0.007 per year, not statistically significant), and the estimates depend on the language-exposure control. In the model, acceleration enters through the acceptance threshold: the technology lowers the cost of review and generation for tasks in industries that were already reducing their demand for them, so assimilation is most likely where structural pressure exists. Country-year fixed effects absorb economy-wide shocks but not industry-specific restructuring, and the industry employment series in Section 7.3 show no contraction in the creative industries. Public data cannot show what happened to visual-production occupations within those industries, so the within-industry reading is a hypothesis rather than a finding. Separating the two requires firm-level data that observe an employer’s industry, its use of the tool, and its hiring together.")
sub(82,"Neither series isolates visual generation,","The finest public occupation-by-industry cross-tabulation (one-digit occupations by industry section) shows professionals’ employment growing by 0.26 log points in information and communication, 0.17 in professional and technical activities, and 0.14 in arts and entertainment, against a median of 0.19 across 19 sections; these groups are dominated by occupations other than visual production, and two-digit occupation data (cultural professionals, +0.09) are too coarse to isolate designers or artists. Neither series isolates visual generation,")
sub(100,"UK acquisition comes from a voluntary survey with a response rate of about 27 percent and does not identify the employers of particular occupations.","Acquisition rates are sector averages and do not identify the employers of particular occupations.")
sub(102,"and the exposure gradients they show are not effects of the technology.","and the exposure gradients they show are small, imprecise, and no larger than the trend that preceded the technology.")
t=orig[104].replace("Eurostat, Office for National Statistics, Statistics Canada","Eurostat, Statistics Canada, U.S. Census Bureau")+""
settext(P[104],t)
delete(108)   # quality rubric (unused after removing UK heterogeneity analysis)
# ---------- appendix B
settext(P[110],"Let exposure {i}e{/i}ⱼ be proportional to the visual share of work, {i}e{/i}ⱼ = {i}λs{/i}ⱼ, and let standardized exposure be {i}z{/i}ⱼ = ({i}e{/i}ⱼ − ē)/σₑ. If the most exposed occupation has {i}s{/i}ⱼ ≈ 1 and {i}z{/i}ⱼ = {i}k{/i}, and if exposure near the mean is small relative to its maximum, then {i}s{/i}ⱼ ≈ {i}z{/i}ⱼ/{i}k{/i} over the relevant range. Substituting into equation (2) with {i}A{/i}ᵃ = {i}A{/i}ᵘ gives Δ ln {i}N{/i}ⱼ ≈ ({i}A{/i}ᵘ{i}β{/i}/{i}k{/i})·{i}z{/i}ⱼ, so the regression coefficient on {i}z{/i}ⱼ under complete assimilation is {i}γ{/i} = {i}A{/i}ᵘ{i}β{/i}/{i}k{/i}. The minimum detectable effect for a two-sided test of size 0.05 with power 0.8 is (1.96 + 0.84)·SE ≈ 2.8·SE. Detection for the most exposed occupation requires |{i}A{/i}ᵃ{i}β{/i}| ≥ 2.8·SE·{i}k{/i}. With SE = 0.017 (Europe, 2024) and {i}k{/i} = 5.7 the threshold is 0.28. If mean exposure is not small relative to its maximum, {i}k{/i} overstates the effective spread and the implied {i}γ{/i}* is somewhat larger, but for the values in Table 5 to cross the detection threshold the effective {i}k{/i}* would need to fall below about 4 under the most aggressive scenario.")
# ---------- references: remove ONS
for i in (133,132): delete(i)
# ---------- tables
T=d.tables
t4=T[2]; setrow(t4,0,["Specification","2019","2020","2021","2023","2024"])
setrow(t4,1,["Country-year FE, LLM control (preferred)","+0.024 (0.014)","+0.008 (0.013)","+0.011 (0.009)","−0.016 (0.006)","−0.012 (0.017)"])
setrow(t4,2,["Country-year FE, no LLM control","+0.008 (0.013)","−0.002 (0.013)","+0.008 (0.008)","−0.037 (0.009)","−0.059 (0.021)"])
setrow(t4,3,["Deviation from linear pre-trend, LLM control","","","","−0.009 (0.007)","+0.002 (0.023)"])
setrow(t4,4,["Randomization p (preferred)","","","","0.016","0.19"])
setrow(t4,5,["Minimum detectable effect (preferred)","0.038","0.037","0.024","0.016","0.048"])
t5=T[3]
setrow(t5,0,["Acquisition rate Aᵘ","β = −0.19 (platform)","β = −0.50 (aggressive)","Implied log change, most exposed occupation (β = −0.50)"])
setrow(t5,1,["9.6% (EU27, all firms)","−0.003","−0.008","−0.048"])
setrow(t5,2,["36.2% (publishing, film, broadcasting)","−0.012","−0.032","−0.181"])
setrow(t5,3,["2.1% (text generation, 2023, proxy for early acquisition)","−0.001","−0.002","−0.010"])
setrow(t5,4,["Minimum detectable effect, 2024","0.048","0.048","−0.28"])
setrow(t5,5,["Observed (preferred), 2024","−0.012","−0.012","−0.07"])
newrow=copy.deepcopy(t5.rows[5]._tr); t5._tbl.append(newrow)
setrow(t5,6,["Observed (no language control), 2024","−0.059","−0.059","−0.34"])
t6=T[4]
setrow(t6,1,["−0.19 (−17%)","above 100%","never","never","never"])
setrow(t6,2,["−0.50 (−39%)","55%","2029 / 2027","2030 / 2027","2035 / 2029"])
setrow(t6,3,["−0.80 (−55%)","34%","2028 / 2025","2029 / 2025","2032 / 2025"])
d.save(out); print("saved",out)
