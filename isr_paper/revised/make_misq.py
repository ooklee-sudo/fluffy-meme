"""Applies MISQ manuscript conventions (as supplied by the author) to the revised manuscript:
APA 7th in-text citations (information first, parenthetical, '&'), APA reference list (single-spaced, hanging indent, alphabetical),
heading hierarchy (1 MAJOR HEAD bold caps centered; 1.1 Subhead bold title case centered), left-aligned double-spaced Times New Roman 12 body,
title page (title, abstract, keywords) followed by page break.
usage: python revised/make_misq.py <in.docx> <out.docx>"""
import sys, re, copy, docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
src,out=sys.argv[1],sys.argv[2]
d=docx.Document(src)
def ps(): return d.paragraphs
def markup(p):
    s=""
    for r in p.runs:
        t=r.text
        if r.italic: s+="{i}"+t+"{/i}"
        elif r.bold: s+="{b}"+t+"{/b}"
        else: s+=t
    return s.replace("{/i}{i}","").replace("{/b}{b}","")
def base_rpr(p):
    rp=p.runs[0]._r.find(qn("w:rPr")); return copy.deepcopy(rp) if rp is not None else None
def settext(p,text):
    rp=base_rpr(p) if p.runs else None
    for r in list(p.runs): r._r.getparent().remove(r._r)
    for part in re.split(r"(\{i\}.*?\{/i\}|\{b\}.*?\{/b\})",text):
        if not part: continue
        it=part.startswith("{i}"); bd=part.startswith("{b}"); t=re.sub(r"\{/?[ib]\}","",part)
        r=p.add_run(t)
        if rp is not None:
            old=r._r.find(qn("w:rPr"))
            if old is not None: r._r.remove(old)
            nr=copy.deepcopy(rp)
            for tag in ("w:i","w:b"):
                e=nr.find(qn(tag))
                if e is not None: nr.remove(e)
            r._r.insert(0,nr)
        if it: r.italic=True
        if bd: r.bold=True
ri=[k for k,p in enumerate(ps()) if p.text=="References"][-1]
body=ps()[:ri]
# ---- 1. narrative citations -> information first, parenthetical
NARR=[
("Fichman and Kemerer (1999) showed that organizations can acquire an innovation well before they deploy it, so diffusion statistics overstate use; they called the wedge an assimilation gap and found it widest where knowledge barriers were high.","Organizations can acquire an innovation well before they deploy it, so diffusion statistics overstate use; this wedge has been called an assimilation gap and was widest where knowledge barriers were high (Fichman & Kemerer, 1999)."),
("Cooper and Zmud (1990) described IT implementation as a staged process in which adoption is followed by adaptation, acceptance, routinization, and infusion, and in which the factors that explain early stages differ from those that explain later ones.","IT implementation is a staged process in which adoption is followed by adaptation, acceptance, routinization, and infusion, and the factors that explain early stages differ from those that explain later ones (Cooper & Zmud, 1990)."),
("Attewell (1992) argued that complex technologies spread slowly not because firms lack information about them but because they lack the know-how to use them, and that mediating institutions emerge to lower that knowledge barrier.","Complex technologies spread slowly not because firms lack information about them but because they lack the know-how to use them, and mediating institutions emerge to lower that knowledge barrier (Attewell, 1992)."),
("Fichman and Kemerer (1999) turned the distinction into a measurement problem: cumulative acquisition of software-process innovations ran well ahead of cumulative deployment, and the resulting assimilation gap was largest for innovations with high knowledge barriers.","The distinction becomes a measurement problem: cumulative acquisition of software-process innovations ran well ahead of cumulative deployment, and the resulting assimilation gap was largest for innovations with high knowledge barriers (Fichman & Kemerer, 1999)."),
("Zhu and Kraemer (2005) showed that post-adoption usage and value vary widely across adopters of e-business, again pointing to organizational rather than technological determinants.","Post-adoption usage and value vary widely across adopters of e-business, again pointing to organizational rather than technological determinants (Zhu & Kraemer, 2005)."),
("General-purpose technologies display long lags for the same reason (David 1990; Bresnahan and Trajtenberg 1995), and Brynjolfsson, Rock, and Syverson (2021) formalize the lag as a productivity J-curve in which intangible complements are built before measured gains appear. Agrawal, Gans, and Goldfarb (2018) argue that when prediction becomes cheap the value of human judgment rises.","General-purpose technologies display long lags for the same reason (Bresnahan & Trajtenberg, 1995; David, 1990), and the lag can be formalized as a productivity J-curve in which intangible complements are built before measured gains appear (Brynjolfsson et al., 2021). When prediction becomes cheap, the value of human judgment rises (Agrawal et al., 2018)."),
("Acemoglu, Autor, Hazell, and Restrepo (2022) find that establishments exposed to AI change their hiring but detect little aggregate effect at the occupation or industry level.","Establishments exposed to AI change their hiring, but little aggregate effect is detectable at the occupation or industry level (Acemoglu et al., 2022)."),
("Hui, Reshef, and Zhou (2024) document employment and earnings losses for freelancers exposed to ChatGPT and image generators, and Demirci, Hannane, and Zhu (2025) find that image-creation posts fell by about 17 percent after image generators were released, relative to manual-intensive jobs.","Freelancers exposed to ChatGPT and image generators lost employment and earnings (Hui et al., 2024), and image-creation posts fell by about 17 percent after image generators were released, relative to manual-intensive jobs (Demirci et al., 2025)."),
("The floor is a bottleneck in the sense of Baumol (1967) and Aghion, Jones, and Jones (2019): when a task","The floor is a bottleneck: when a task"),
("rises as the generator improves. Because","rises as the generator improves (Aghion et al., 2019; Baumol, 1967). Because"),
("as Eloundou et al. (2023) assess, the experience","(Eloundou et al., 2023), the experience"),
# data sources
("We use Eurostat’s survey on ICT usage in enterprises, which covers","We use Eurostat’s survey on ICT usage in enterprises (Eurostat, 2026), which covers"),
("is Eurostat’s experimental online job advertisements by ISCO-08 three-digit occupation, 2019–2024,","is Eurostat’s experimental online job advertisements by ISCO-08 three-digit occupation, 2019–2024 (Eurostat, 2026),"),
("The Canadian series is the Job Vacancy and Wage Survey by NOC 2021 unit group, 2015–2026,","The Canadian series is the Job Vacancy and Wage Survey by NOC 2021 unit group, 2015–2026 (Statistics Canada, 2026),"),
("US sector postings from Indeed Hiring Lab are used descriptively.","US sector postings from Indeed Hiring Lab (2026) are used descriptively."),
("over O*NET task statements, of a language-model score","over O*NET task statements (National Center for O*NET Development, 2026), of a language-model score"),
("Business Trends and Outlook Survey correlates","Business Trends and Outlook Survey (U.S. Census Bureau, 2026) correlates"),
]
hit={a:0 for a,_ in NARR}
for p in body:
    if not p.text.strip(): continue
    m=markup(p); m2=m
    for a,b in NARR:
        if a in m2: m2=m2.replace(a,b,1); hit[a]+=1
    if m2!=m: settext(p,m2)
missing=[a[:50] for a,c in hit.items() if c==0]; print("narrative rewrites not matched:",missing)
# ---- 2. parenthetical citations -> APA
def fmt_item(item):
    m=re.match(r"^([A-Z][A-Za-z\-’\.]*(?:(?:,| and|, and| &) [A-Z][A-Za-z\-’\.]*)*(?: et al\.)?) (\d{4}[a-z]?(?:, \d{4}[a-z]?)*)$",item.strip())
    if not m: return None
    au,yr=m.groups()
    if au.endswith("et al."): names=[au.replace(" et al.","")]; etal=True
    else: names=[x for x in re.split(r",\s*(?:and\s+)?|\s+and\s+|\s+&\s+",au) if x]; etal=False
    if etal or len(names)>=3: a=names[0]+" et al."
    elif len(names)==2: a=names[0]+" & "+names[1]
    else: a=names[0]
    return a+", "+yr
def conv(text):
    def g(mo):
        inner=mo.group(1); items=[x.strip() for x in inner.split(";")]
        out=[fmt_item(x) for x in items]
        if any(o is None for o in out): return mo.group(0)
        out=sorted(out,key=lambda s:s.lower()); return "("+"; ".join(out)+")"
    return re.sub(r"\(([^()]*?\d{4}[a-z]?)\)",g,text)
chg=0
for p in body:
    if not p.text.strip(): continue
    m=markup(p); m2=conv(m)
    if m2!=m: settext(p,m2); chg+=1
print("paragraphs with parenthetical citations converted:",chg)
for p in body:
    if "(Eurostat, 2025)" in p.text:
        settext(p,markup(p).replace("(Eurostat, 2025)","(Eurostat, 2026)"))
# ---- 3. reference list (APA 7)
REFS=[
([("Acemoglu, D., Autor, D., Hazell, J., & Restrepo, P. (2022). Artificial intelligence and jobs: Evidence from online vacancies. ",0),("Journal of Labor Economics",1),(", ",0),("40",1),("(S1), S293–S340.",0)]),
([("Acemoglu, D., & Restrepo, P. (2018). The race between man and machine: Implications of technology for growth, factor shares, and employment. ",0),("American Economic Review",1),(", ",0),("108",1),("(6), 1488–1542.",0)]),
([("Acemoglu, D., & Restrepo, P. (2019). Automation and new tasks: How technology displaces and reinstates labor. ",0),("Journal of Economic Perspectives",1),(", ",0),("33",1),("(2), 3–30.",0)]),
([("Aghion, P., Jones, B. F., & Jones, C. I. (2019). Artificial intelligence and economic growth. In A. Agrawal, J. Gans, & A. Goldfarb (Eds.), ",0),("The economics of artificial intelligence: An agenda",1),(". University of Chicago Press.",0)]),
([("Agrawal, A., Gans, J., & Goldfarb, A. (2018). ",0),("Prediction machines: The simple economics of artificial intelligence",1),(". Harvard Business Review Press.",0)]),
([("Aral, S., & Weill, P. (2007). IT assets, organizational capabilities, and firm performance: How resource allocations and organizational differences explain performance variation. ",0),("Organization Science",1),(", ",0),("18",1),("(5), 763–780.",0)]),
([("Attewell, P. (1992). Technology diffusion and organizational learning: The case of business computing. ",0),("Organization Science",1),(", ",0),("3",1),("(1), 1–19.",0)]),
([("Baumol, W. J. (1967). Macroeconomics of unbalanced growth: The anatomy of urban crisis. ",0),("American Economic Review",1),(", ",0),("57",1),("(3), 415–426.",0)]),
([("Bresnahan, T. F., & Trajtenberg, M. (1995). General purpose technologies: “Engines of growth”? ",0),("Journal of Econometrics",1),(", ",0),("65",1),("(1), 83–108.",0)]),
([("Brynjolfsson, E., & Hitt, L. M. (2000). Beyond computation: Information technology, organizational transformation and business performance. ",0),("Journal of Economic Perspectives",1),(", ",0),("14",1),("(4), 23–48.",0)]),
([("Brynjolfsson, E., Li, D., & Raymond, L. (2025). Generative AI at work. ",0),("The Quarterly Journal of Economics",1),(", ",0),("140",1),("(2), 889–942.",0)]),
([("Brynjolfsson, E., Rock, D., & Syverson, C. (2021). The productivity J-curve: How intangibles complement general purpose technologies. ",0),("American Economic Journal: Macroeconomics",1),(", ",0),("13",1),("(1), 333–372.",0)]),
([("Cooper, R. B., & Zmud, R. W. (1990). Information technology implementation research: A technological diffusion approach. ",0),("Management Science",1),(", ",0),("36",1),("(2), 123–139.",0)]),
([("David, P. A. (1990). The dynamo and the computer: An historical perspective on the modern productivity paradox. ",0),("American Economic Review",1),(", ",0),("80",1),("(2), 355–361.",0)]),
([("Demirci, O., Hannane, J., & Zhu, X. (2025). Who is AI replacing? The impact of generative AI on online freelancing platforms. ",0),("Management Science",1),(", ",0),("71",1),("(10), 8097–8108.",0)]),
([("Eloundou, T., Manning, S., Mishkin, P., & Rock, D. (2023). ",0),("GPTs are GPTs: An early look at the labor market impact potential of large language models",1),(" (arXiv:2303.10130). arXiv.",0)]),
([("Eurostat. (2026). ",0),("ICT usage in enterprises; online job advertisements by occupation (experimental); Labour force survey",1),(" [Data sets]. Retrieved October 4, 2026, from the Eurostat dissemination database.",0)]),
([("Felten, E., Raj, M., & Seamans, R. (2021). Occupational, industry, and geographic exposure to artificial intelligence. ",0),("Strategic Management Journal",1),(", ",0),("42",1),("(12), 2195–2217.",0)]),
([("Fichman, R. G., & Kemerer, C. F. (1999). The illusory diffusion of innovation: An examination of assimilation gaps. ",0),("Information Systems Research",1),(", ",0),("10",1),("(3), 255–275.",0)]),
([("Hui, X., Reshef, O., & Zhou, L. (2024). The short-term effects of generative artificial intelligence on employment: Evidence from an online labor market. ",0),("Organization Science",1),(", ",0),("35",1),("(6), 1977–1989.",0)]),
([("Indeed Hiring Lab. (2026). ",0),("Job postings index, by sector, United States",1),(" [Data set]. Retrieved October 4, 2026. Creative Commons Attribution 4.0.",0)]),
([("Kim, S., Jin, G. Z., & Lee, E. (2026). ",0),("Does generative AI crowd out human creators? Evidence from Pixiv",1),(" (NBER Working Paper No. 34733). National Bureau of Economic Research.",0)]),
([("Melville, N., Kraemer, K., & Gurbaxani, V. (2004). Information technology and organizational performance: An integrative model of IT business value. ",0),("MIS Quarterly",1),(", ",0),("28",1),("(2), 283–322.",0)]),
([("National Center for O*NET Development. (2026). ",0),("O*NET database",1),(" (Version 29.1) [Data set]; ESCO crosswalk. Retrieved October 4, 2026. Creative Commons Attribution 4.0.",0)]),
([("Statistics Canada. (2026). ",0),("Job vacancies by occupation",1),(" (Table 14-10-0444; NOC 2021) [Data set]. Retrieved October 4, 2026. Statistics Canada Open Licence.",0)]),
([("U.S. Census Bureau. (2026). ",0),("Business Trends and Outlook Survey (BTOS), sector estimates",1),(" [Data set]. Retrieved October 4, 2026.",0)]),
([("Webb, M. (2020). ",0),("The impact of artificial intelligence on the labor market",1),(" [Working paper]. SSRN. https://ssrn.com/abstract=3482150",0)]),
([("Zhou, E., & Lee, D. (2024). Generative artificial intelligence, human creativity, and art. ",0),("PNAS Nexus",1),(", ",0),("3",1),("(3), Article pgae052.",0)]),
([("Zhu, K., & Kraemer, K. L. (2005). Post-adoption variations in usage and value of e-business by organizations: Cross-country evidence from the retail industry. ",0),("Information Systems Research",1),(", ",0),("16",1),("(1), 61–84.",0)]),
]
refps=ps()[ri+1:]; model=refps[0]
prev=model
for k,parts in enumerate(REFS):
    e=copy.deepcopy(model._p); prev._p.addnext(e) if k else model._p.addprevious(e)
    from docx.text.paragraph import Paragraph
    q=Paragraph(e,model._parent); rp=base_rpr(q)
    for r in list(q.runs): r._r.getparent().remove(r._r)
    for t,it in parts:
        r=q.add_run(t); old=r._r.find(qn("w:rPr"))
        if old is not None: r._r.remove(old)
        nr=copy.deepcopy(rp); ie=nr.find(qn("w:i"))
        if ie is not None: nr.remove(ie)
        r._r.insert(0,nr); r.italic=True if it else None
        r.font.size=Pt(12)
    pf=q.paragraph_format; pf.left_indent=Inches(0.5); pf.first_line_indent=Inches(-0.5); pf.line_spacing=1.0; pf.space_after=Pt(6); q.alignment=WD_ALIGN_PARAGRAPH.LEFT
    prev=q
for p in refps: p._p.getparent().remove(p._p)
# ---- 4. title page, headings, alignment, spacing
P=ps()
# remove meta lines, merge title
for p in list(P[:5]):
    if p.text.startswith(("Manuscript for double-blind","Author names and affiliations")): p._p.getparent().remove(p._p)
P=ps()
title=P[0]; sub=P[1]; settext(title,title.text.strip()+" "+sub.text.strip()); sub._p.getparent().remove(sub._p)
for r in title.runs: r.font.size=Pt(16); r.bold=True
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
hd1=re.compile(r"^(\d+)\.\s+(.+)$"); hd2=re.compile(r"^(\d+\.\d+)\s+(.+)$")
UNN={"Abstract","Data, Code, and Use of Generative AI","References"}
APP=re.compile(r"^Appendix [A-Z]\.")
def style_head(p,text,center=True):
    settext(p,text)
    for r in p.runs: r.bold=True; r.italic=False; r.font.size=Pt(12)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing=1.0; p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(6)
for p in ps():
    t=p.text.strip()
    if not t or len(t)>110: continue
    m1=hd1.match(t); m2=hd2.match(t)
    if m1 and not m2 and not t.endswith("."): style_head(p,f"{m1.group(1)} {m1.group(2).upper()}")
    elif m2 and not t.endswith("."): style_head(p,t)
    elif t in UNN or APP.match(t): style_head(p,t.upper() if t!="Abstract" else "Abstract")
# page break before the first major section
for p in ps():
    if p.text.startswith("1 INTRODUCTION"): p.paragraph_format.page_break_before=True; break
# alignment: left only; body double spaced; tables single
ri=[k for k,p in enumerate(ps()) if p.text=="REFERENCES"][-1]
for k,p in enumerate(ps()):
    if p.alignment==WD_ALIGN_PARAGRAPH.JUSTIFY: p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if k<ri and len(p.text)>180 and not p.text.startswith(("Notes.","Source.")): p.paragraph_format.line_spacing=2.0
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.line_spacing=1.0
                if pp.alignment==WD_ALIGN_PARAGRAPH.JUSTIFY: pp.alignment=WD_ALIGN_PARAGRAPH.LEFT
d.save(out); print("saved",out)
