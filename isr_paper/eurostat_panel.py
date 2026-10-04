"""Exploratory sector x country panel: AI technology use (Eurostat isoc_eb_ain2, enterprises 10+) on sector visual / language exposure.
Exposure construction: O*NET-SOC -> ESCO/ISCO-08 (O*NET crosswalk) -> ISCO major group (mean over mapped O*NET-SOC);
sector exposure = EU27 2021 employment-share weighted mean over ISCO major groups (Eurostat lfsa_eisn2).
Outputs: data/eurostat/sector_exposure.csv, results printed + data/eurostat/panel_results.json"""
import json, csv, random, collections, numpy as np, pandas as pd
from openpyxl import load_workbook
import visual_intensity as VI

def jload(f):
    d=json.load(open(f)); ids=d["id"]; size=d["size"]
    idx={k:d["dimension"][k]["category"]["index"] for k in ids}; inv={k:{i:c for c,i in idx[k].items()} for k in ids}
    rows=[]
    for pos,v in d["value"].items():
        p=int(pos); key={}
        for k,s in reversed(list(zip(ids,size))): key[k]=inv[k][p%s]; p//=s
        key["v"]=v; rows.append(key)
    return pd.DataFrame(rows)

# ---- occupation-group exposure
xw=[r for r in load_workbook("data/ext/ESCO_to_ONET-SOC.xlsx",read_only=True).worksheets[0].iter_rows(min_row=5,values_only=True)]
lang={r["O*NET-SOC Code"]:float(r["dv_rating_beta"]) for r in csv.DictReader(open("data/ext/occ_level.csv"))}
V=collections.defaultdict(list); L=collections.defaultdict(list)
for esco,_,soc,_ in xw:
    if not esco or not soc: continue
    g=str(esco)[0]; g="OC%s"%g
    if soc in VI.onet: V[g].append(VI.onet[soc])
    if soc in lang: L[g].append(lang[soc])
Vg={g:np.mean(v) for g,v in V.items()}; Lg={g:np.mean(v) for g,v in L.items()}
print("ISCO major group  visual  language(beta)  n_V")
for g in sorted(Vg): print(" ",g,round(Vg[g],3),round(Lg.get(g,np.nan),3),len(V[g]))

# ---- sector exposure from EU27 employment structure, 2021
e=jload("data/eurostat/lfsa_eisn2.json")
e=e[(e.geo=="EU27_2020")&(e.time=="2021")&(e.age=="Y_GE15")&(e.sex=="T")&(e.isco08.isin(Vg))]
SECT=["C","D","E","F","G","H","I","J","L","M","N"]
rows=[]
for s in SECT:
    x=e[e.nace_r2==s].set_index("isco08").v.dropna(); x=x[x.index.isin(Vg)&x.index.isin(Lg)]
    w=x/x.sum(); rows.append((s,float((w*pd.Series(Vg)[w.index]).sum()),float((w*pd.Series(Lg)[w.index]).sum())))
ex=pd.DataFrame(rows,columns=["nace","vis","lang"]); ex["zvis"]=(ex.vis-ex.vis.mean())/ex.vis.std(); ex["zlang"]=(ex.lang-ex.lang.mean())/ex.lang.std()
ex.to_csv("data/eurostat/sector_exposure.csv",index=False); print(ex.round(3).to_string(index=False)); print("corr(vis,lang) across sectors:",round(ex.vis.corr(ex.lang),3))

# ---- AI use panel
a=jload("data/eurostat/isoc_eb_ain2.json"); a=a[(a.unit=="PC_ENT")&(a.size_emp=="GE10")&(a.nace_r2.isin(SECT))&(a.geo!="EU27_2020")&(~a.geo.isin(["EA","EA19","EA20"]))]
a=a.merge(ex[["nace","zvis","zlang"]],left_on="nace_r2",right_on="nace")

def fit(ind,use_lang=True,perm=None):
    d=a[a.indic_is==ind].dropna(subset=["v"]); d=d[d.v>0].copy(); d["y"]=np.log(d.v)
    d=d[d.geo.isin(d.groupby("geo").nace_r2.nunique()[lambda s:s>=8].index)]
    d["t"]=d.time.astype(int); d["sc"]=d.geo+d.nace_r2; d["ct"]=d.geo+d.time
    if perm: d["zvis"]=d.nace_r2.map(perm[0]); d["zlang"]=d.nace_r2.map(perm[1])
    cols=[]; X=[]
    for t in (2023,2024,2025):
        X.append(((d.t==t)*d.zvis).values); cols.append(f"vis_{t}")
        if use_lang: X.append(((d.t==t)*d.zlang).values); cols.append(f"lang_{t}")
    X=np.column_stack(X); FE=pd.get_dummies(d[["sc","ct"]],drop_first=False).astype(float).values
    Z=np.column_stack([X,FE]); beta=np.linalg.lstsq(Z,d.y.values,rcond=None)[0][:X.shape[1]]
    return dict(zip(cols,beta)),len(d),d.geo.nunique()

res={}
for ind,name in (("E_AI_TIR","image recognition/processing"),("E_AI_TNLG","natural language generation"),("E_AI_TML","machine learning"),("E_AI_TTM","text mining")):
    b,n,k=fit(ind); b2,_,_=fit(ind,use_lang=False)
    # randomization inference: permute (vis,lang) pairs across the 11 sectors
    rnd=random.Random(1); sec=list(ex.nace); obs=b["vis_2025"]; cnt=0; R=300
    for _ in range(R):
        pm=sec[:]; rnd.shuffle(pm); mp={s:p for s,p in zip(sec,pm)}
        pv={s:float(ex.set_index("nace").zvis[mp[s]]) for s in sec}; pl={s:float(ex.set_index("nace").zlang[mp[s]]) for s in sec}
        bb,_,_=fit(ind,perm=(pv,pl)); cnt+=abs(bb["vis_2025"])>=abs(obs)
    res[ind]=dict(name=name,n=n,countries=k,with_lang=b,vis_only=b2,perm_p_vis2025=(cnt+1)/(R+1))
    print(f"\n{name} ({ind}): N={n}, countries={k}\n  vis+lang:",{k_:round(v,3) for k_,v in b.items()},"\n  vis only:",{k_:round(v,3) for k_,v in b2.items()},"\n  permutation p (|vis_2025|, 300 draws):",round((cnt+1)/(R+1),3))
json.dump(res,open("data/eurostat/panel_results.json","w"),indent=1)
