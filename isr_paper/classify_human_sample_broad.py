import os
"""Apply the broad 'visual materials' rubric (prompt_broad.txt) to the 200 human-coded sample tasks and evaluate against coder A (design-weighted)."""
import json, os, sys, numpy as np, pandas as pd, concurrent.futures as cf
import classify_onet as C, agreement_creation as A
from openpyxl import load_workbook
C.PROMPT=open(sys.argv[1] if len(sys.argv)>1 else "prompt_broad.txt").read(); tag=sys.argv[2] if len(sys.argv)>2 else "broad_v1"
key=A.key.copy(); key["tid"]=key["id"]
ws=load_workbook("data/human_creation_sample_filled_coderA.xlsx",read_only=True)["coding"]; rows=list(ws.iter_rows(values_only=True))
h=pd.DataFrame(rows[1:],columns=rows[0]); h["h"]=pd.to_numeric(h.score_0_1_2,errors="coerce"); d=key.merge(h[["row","h"]],on="row"); d=d[d.h.notna()].copy(); d["h"]=d.h.astype(int)
bs=[d.iloc[i:i+20] for i in range(0,len(d),20)]; k=os.environ["OPENROUTER_API_KEY"]; lab={}
with cf.ThreadPoolExecutor(3) as ex:
    for r,_ in ex.map(lambda b:C.call(os.environ["LLM_MODEL"],[{"id":int(x.row),"text":f"[{x.occupation}] {x.text}"} for x in b.itertuples()],k),bs): lab.update(r)
d["b"]=d.row.map(lab); d.to_csv(f"data/human_sample_{tag}.csv",index=False)
print(tag,"n",len(d)); A.report("coder A vs broad LLM",d.h,d.b,d.weight)
sp=lambda a,b:pd.Series(a).rank().corr(pd.Series(b).rank()); print("Spearman(human,broad)=%.2f ; (human,narrow)=%.2f"%(sp(d.h,d.b),sp(d.h,d.s)))
w=d.weight; print("design-weighted share scored >=1: human %.1f%%, broad LLM %.1f%%, narrow LLM %.1f%%"%(100*(w*(d.h>=1)).sum()/w.sum(),100*(w*(d.b>=1)).sum()/w.sum(),100*(w*(d.s>=1)).sum()/w.sum()))
print(pd.crosstab(d.b,d.h,rownames=["broad LLM"],colnames=["human"]))
