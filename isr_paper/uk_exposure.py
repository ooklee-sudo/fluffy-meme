"""Build UK SOC-2020 (4-digit) occupational exposures from O*NET-based measures through ISCO-08.
Chain: O*NET-SOC --(ESCO/ISCO crosswalk published by O*NET; ESCO code prefix = ISCO-08 4-digit)--> ISCO-08 4-digit
       ISCO-08 4-digit <--(ONS SOC2020 coding index, 'ISCO-08 code based on SOC2020', weights = number of index entries)-- UK SOC 2020 unit group.
Exposure(UK SOC4) = sum_isco w(soc4,isco) * mean_{O*NET-SOC in isco} exposure."""
import csv, collections, numpy as np, pandas as pd
from openpyxl import load_workbook
def isco_weights():
    ws=load_workbook("data/ons/soc2020_coding_index.xlsx",read_only=True)["SOC2020 coding index"]
    it=ws.iter_rows(values_only=True); h=list(next(it)); i_soc=h.index("SOC_2020"); i_isco=h.index("ISCO-08 code based on SOC2020")
    cnt=collections.defaultdict(collections.Counter)
    for r in it:
        s,i=r[i_soc],r[i_isco]
        if s is None or i is None: continue
        try: s=f"{int(s):04d}"; i=f"{int(str(i)[:4]):04d}"
        except: continue
        cnt[s][i]+=1
    return {s:{i:c/sum(cn.values()) for i,c in cn.items()} for s,cn in cnt.items()}
def onet_to_isco():
    ws=load_workbook("data/ext/ESCO_to_ONET-SOC.xlsx",read_only=True).worksheets[0]
    m=collections.defaultdict(set)
    for esco,_,soc,_ in ws.iter_rows(min_row=5,values_only=True):
        if esco and soc: m[str(soc)].add(str(esco)[:4])
    return m
def uk_exposure(onet_exposure:dict):
    """onet_exposure: {O*NET-SOC code: value}. Returns Series indexed by UK SOC4 and the share of ISCO weight covered."""
    o2i=onet_to_isco(); byisco=collections.defaultdict(list)
    for soc,v in onet_exposure.items():
        for i in o2i.get(soc,()): byisco[i].append(v)
    E={i:float(np.mean(v)) for i,v in byisco.items()}
    inv=collections.defaultdict(set)
    for soc,iscos in o2i.items():
        for i in iscos: inv[i].add(soc)
    fr={i:len([c for c in inv[i] if c in onet_exposure])/len(inv[i]) for i in inv}      # share of mapped O*NET occupations that have an exposure value
    W=isco_weights(); out={}; cov={}; covo={}
    for s,w in W.items():
        tot=sum(wt for i,wt in w.items() if i in E)
        if tot>0:
            out[s]=sum(wt*E[i] for i,wt in w.items() if i in E)/tot; cov[s]=tot; covo[s]=sum(wt*fr.get(i,0) for i,wt in w.items())
    uk_exposure.onet_cov=pd.Series(covo)
    return pd.Series(out),pd.Series(cov)
if __name__=="__main__":
    lang={r["O*NET-SOC Code"]:float(r["dv_rating_beta"]) for r in csv.DictReader(open("data/ext/occ_level.csv"))}
    L,cov=uk_exposure(lang); print("UK SOC4 with language exposure:",len(L),"| mean ISCO weight covered %.2f"%cov.mean())
    print(L.loc[["2142","3411","3417","2141","3416","2494","3412","2492","2491","5111","6131"]].round(3).to_dict() if all(k in L.index for k in ["2142","3411","3417","2141"]) else L.head())
    L.to_csv("data/ons/uk_soc4_langexp.csv",header=["langexp"])
