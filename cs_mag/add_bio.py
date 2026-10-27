"""Append the author biography (AuthorBioHead/AuthorBio styles of the magazine template) with a floating photo."""
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import qn
d = Document("verify_the_record_CS_Mag_original.docx")
ST = lambda i: next(s for s in d.styles if s.style_id == i)
BIO = (" is a professor of department of Information System at Hanyang University, Seoul, Korea. He received a PhD in Management of "
       "Information Systems from Claremont Graduate University, USA. He was a faculty member at the University of Queensland, Australia, "
       "University of Nevada at Las Vegas, USA and North Carolina A&T State University, USA. His research interests are information system "
       "economics, information system maintenance, large language models, and critical social theory. Contact him at ooklee@hanyang.ac.kr.")
d.add_paragraph("ABOUT THE AUTHOR", style=ST("AuthorBioHead"))
p = d.add_paragraph(style=ST("AuthorBio"))
r = p.add_run(); r.add_picture("../nmi_screening/tai_template/author_photo.jpg", width=Inches(0.9), height=Inches(1.08))
inl = r._r.find(".//" + qn("wp:inline")); ext = inl.find(qn("wp:extent"))
anchor = parse_xml('<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" distT="0" distB="0" distL="0" distR="114300" '
                   'simplePos="0" relativeHeight="2" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
                   '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
                   f'<wp:extent cx="{ext.get("cx")}" cy="{ext.get("cy")}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapSquare wrapText="right"/></wp:anchor>')
anchor.append(inl.find(qn("wp:docPr"))); anchor.append(parse_xml('<wp:cNvGraphicFramePr xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"/>')); anchor.append(inl.find(qn("a:graphic")))
inl.getparent().replace(inl, anchor)
b = p.add_run("OOK LEE"); b.bold = True
p.add_run(BIO)
d.save("verify_the_record_CS_Mag.docx"); print("ok")
