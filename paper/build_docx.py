"""Build a clean Word file from the markdown source: Times New Roman, academic three-line tables, centered figures, page numbers."""
import subprocess, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, out = sys.argv[1], sys.argv[2]
subprocess.run(["pandoc", src, "-o", out, "--resource-path=."], check=True)
d = Document(out)
FONT = "Times New Roman"
def setfont(style, size, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    style.font.name = FONT; style.font.size = Pt(size); style.font.color.rgb = color
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), FONT)
for st in d.styles:
    name = st.name
    if not hasattr(st, "font") or st.type not in (1, 2): continue
    if name in ("Title",): setfont(st, 17, bold=True)
    elif name in ("Subtitle",): setfont(st, 11, italic=True)
    elif name == "Heading 1": setfont(st, 13, bold=True)
    elif name == "Heading 2": setfont(st, 11.5, bold=True, italic=False)
    elif name in ("Table Caption", "Image Caption", "Caption"): setfont(st, 9.5, italic=False)
    elif name in ("Body Text", "First Paragraph", "Compact", "Normal", "Block Text"): setfont(st, 11)
    elif name in ("Heading 3", "Heading 4"): setfont(st, 11, bold=True)
for st in d.styles:
    if st.name in ("Body Text", "First Paragraph", "Compact") and st.type == 1:
        pf = st.paragraph_format; pf.space_after = Pt(6); pf.line_spacing = 1.15
for sec in d.sections:
    sec.left_margin = sec.right_margin = Inches(1); sec.top_margin = sec.bottom_margin = Inches(1)
    # page number in footer
    p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); 
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if t: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._r.append(e)
        else: e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)
def border(cell_el, **edges):
    tcPr = cell_el.get_or_add_tcPr(); b = tcPr.find(qn("w:tcBorders"))
    if b is None: b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for edge, sz in edges.items():
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), "000000"); b.append(e)
for t in d.tables:
    t.autofit = True
    n = len(t.rows)
    for i, row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1); p.paragraph_format.line_spacing = 1.0
                for r in p.runs: r.font.size = Pt(8.5); r.font.name = FONT
            if i == 0: border(c._tc, top=10, bottom=6)
            if i == n - 1: border(c._tc, bottom=10)
            if i == 0:
                for p in c.paragraphs:
                    for r in p.runs: r.font.bold = True
for p in d.paragraphs:
    if p.style.name in ("Image Caption", "Captioned Figure"): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for shp in d.inline_shapes:
    w, h = shp.width, shp.height; target = Inches(6.3)
    if w > target: shp.height = int(h * target / w); shp.width = target
d.save(out); print("saved", out, "| tables:", len(d.tables), "| figures:", len(d.inline_shapes))
