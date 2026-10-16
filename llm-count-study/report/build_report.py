"""Assemble the English results draft (markdown -> docx). Run from this folder: python build_report.py"""
import re, subprocess
import pandas as pd

SHORT = {"openai/gpt-4o-mini": "GPT-4o-mini", "meta-llama/llama-3.3-70b-instruct": "Llama-3.3-70B",
         "anthropic/claude-haiku-5.5": "Haiku 5.5"}
DESIGN = {"A": "A Direct", "B": "B Daily", "C": "C Describe", "D": "D Hybrid"}


def dist_table(csv, ref):
    df = pd.read_csv(csv)
    rows = ["| Model | Prompt | Design | Mean | Var/mean | Zero | Alpha |", "|---|---|---|---|---|---|---|",
            f"| Reference | | | {ref[0]} | {ref[1]} | {ref[2]} | {ref[3]} |"]
    for m in SHORT:
        for an in (False, True):
            for c in "ABCD":
                r = df[(df.model == m) & (df.anchor == an) & (df.cond == c)].iloc[0]
                f = lambda v, e, nd=2: f"{v:.{nd}f}" + ("\\*" if e else "")
                rows.append(f"| {SHORT[m]} | {'anchored' if an else 'plain'} | {DESIGN[c]} | {f(r['mean'], r.mean_excl)} | "
                            f"{f(r.var_mean, r.var_mean_excl)} | {f(r.zero, r.zero_excl)} | {f(r.alpha, r.alpha_excl)} |")
    return "\n".join(rows)


def hill_table():
    d = pd.read_csv("hillstrom_table.csv")
    d = d[d.method == "C_prob"]
    rows = ["| Model | Prompt | Outcome | No email | Men's | Women's | Lift men | Lift women | RR men | RR women |",
            "|---|---|---|---|---|---|---|---|---|---|",
            "| Real (randomized) | | visit | 0.106 | 0.183 | 0.151 | +0.077 | +0.045 | 1.72 | 1.43 |",
            "| Real (randomized) | | purchase | 0.0057 | 0.0125 | 0.0088 | +0.0068 | +0.0031 | 2.19 | 1.54 |"]
    for o in ("visit", "purchase"):
        for m in SHORT:
            for an in (False, True):
                r = d[(d.model == m) & (d.anchor == an) & (d.outcome == o)].iloc[0]
                rr = lambda v, e: f"{v:.2f}" + ("\\*" if e else "")
                rows.append(f"| {SHORT[m]} | {'anchored' if an else 'plain'} | {o} | {r.ctrl:.3f} | {r.men:.3f} | {r.women:.3f} | "
                            f"{r.RD_men:+.3f} | {r.RD_women:+.3f} | {rr(r.RR_men, r.RR_men_excl)} | {rr(r.RR_women, r.RR_women_excl)} |")
    return "\n".join(rows)


def stats():
    r, g = pd.read_csv("retail_ci.csv"), pd.read_csv("grocery_ci.csv")
    out = {}
    for name, df, ceil_v in (("retail", r, 2.10), ("grocery", g, 6.61)):
        abc = df[df.cond.isin(list("ABC"))]
        out[name + "_abc_below_ref"] = int((abc.var_mean_hi < (2.49 if name == "retail" else 7.85)).sum())
        out[name + "_abc_n"] = len(abc)
        c = df[df.cond == "C"]
        out[name + "_c_cover_ceiling"] = int(((c.var_mean_lo <= ceil_v) & (c.var_mean_hi >= ceil_v)).sum())
        d = df[df.cond == "D"]
        out[name + "_d_cover_ref"] = int((~d.var_mean_excl).sum())
    return out


def references():
    src = open("../related_work_section.md").read().split("**References**")[1].strip().splitlines()
    refs = [l.strip() for l in src if l.strip()]
    refs += [
        "Chen, D. (2019). *Online Retail II* [Data set]. UCI Machine Learning Repository.",
        "dunnhumby. *The Complete Journey* [Data set]. Accessed through the completejourney_py package.",
        "Fader, P. S., & Hardie, B. G. S. (2001). Forecasting repeat sales at CDNOW: A case study. *Interfaces, 31*(3, Part 2), S94–S107.",
        "Hillstrom, K. (2008, March 20). The MineThatData e-mail analytics and data mining challenge. *MineThatData* (blog).",
    ]
    refs.sort(key=lambda s: re.sub(r"[^a-z]", "", s.lower()))
    return "\n\n".join(refs)


def related_work():
    t = open("../related_work_section.md").read().split("**References**")[0]
    return t.replace("**2. Related Work and Research Gap**", "## 2. Related Work and Research Gap").strip()


md = open("report_template.md").read()
S = stats()
md = (md.replace("{{RELATED}}", related_work()).replace("{{TABLE_RETAIL}}", dist_table("retail_ci.csv", (0.49, 2.49, 0.74, 3.05)))
        .replace("{{TABLE_GROCERY}}", dist_table("grocery_ci.csv", (4.84, 7.85, 0.19, 1.42)))
        .replace("{{TABLE_HILL}}", hill_table()).replace("{{REFS}}", references()))
open("report_en_v2.md", "w").write(md)
print(S)
OUT = "../research_report_en_v2.docx"
subprocess.run(["pandoc", "report_en_v2.md", "-o", OUT, "--resource-path=."], check=True)


def polish(path):
    """Compact bordered tables and readable fonts (pandoc's default tables are loose)."""
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = Document(path)
    for st in d.styles:
        if st.type == 1 and st.name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
            st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for t in d.tables:
        tblPr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "BFBFBF")
            b.append(e)
        tblPr.append(b)
        for ri, row in enumerate(t.rows):
            for c in row.cells:
                if ri == 0:
                    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "E7E6E2")
                    c._tc.get_or_add_tcPr().append(sh)
                for p in c.paragraphs:
                    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(8); r.font.name = "Calibri"; r.bold = True if ri == 0 else r.bold
    d.save(path)


polish(OUT)
print("built", OUT)
