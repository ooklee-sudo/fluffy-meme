"""Build the Word files from the Markdown sources and format tables and page numbers.
usage: python3 build_docx.py   (needs pandoc, python-docx)"""
import re, subprocess, statistics
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W = 9360                                                   # table width in DXA (6.5 inch)
NUM = re.compile(r"^[+\-−–]?\s*\$?[\d.,]+")                # numeric-looking cell

def add_page_numbers(doc):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for t, txt in (("begin", None), ("instr", " PAGE "), ("end", None)):
            r = p.add_run(); r.font.name = "Times New Roman"; r.font.size = Pt(12)
            if t == "instr": e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
            else: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
            r._r.append(e)

def border(tag, sz, color="000000"):
    b = OxmlElement(f"w:{tag}"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "0"); b.set(qn("w:color"), color); return b

def fix_table(tbl, font_pt):
    t = tbl._tbl; tblPr = t.tblPr
    for tag in ("w:tblW", "w:tblLayout", "w:tblBorders", "w:tblCellMar", "w:jc"):
        for e in tblPr.findall(qn(tag)): tblPr.remove(e)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(W)); w.set(qn("w:type"), "dxa"); tblPr.append(w)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tblPr.append(jc)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    bd = OxmlElement("w:tblBorders")
    bd.append(border("top", 8)); bd.append(border("bottom", 8)); bd.append(border("insideH", 4, "BFBFBF"))
    for side in ("left", "right", "insideV"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "nil"); bd.append(e)
    tblPr.append(bd)
    cm = OxmlElement("w:tblCellMar")
    for side, v in (("top", 40), ("left", 70), ("bottom", 40), ("right", 70)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); cm.append(e)
    tblPr.append(cm)
    rows = tbl.rows; ncol = len(rows[0].cells)
    texts = [[c.text.strip() for c in r.cells] for r in rows]
    weights = []
    for j in range(ncol):
        lens = [len(texts[i][j]) for i in range(len(rows))]
        mx, avg = max(lens), statistics.mean(lens)
        weights.append(min(48.0, max(7.0, 0.5 * mx + 0.5 * avg)))
    tot = sum(weights); widths = [int(W * x / tot) for x in weights]; widths[-1] += W - sum(widths)
    grid = t.find(qn("w:tblGrid"))
    if grid is not None: t.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for x in widths: g = OxmlElement("w:gridCol"); g.set(qn("w:w"), str(x)); grid.append(g)
    t.insert(list(t).index(tblPr) + 1, grid)
    numeric = [sum(bool(NUM.match(texts[i][j])) for i in range(1, len(rows))) >= 0.7 * max(1, len(rows) - 1) and j > 0 for j in range(ncol)]
    for i, row in enumerate(rows):
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        if i == 0: h = OxmlElement("w:tblHeader"); trPr.append(h)
        for j, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            for e in tcPr.findall(qn("w:tcW")): tcPr.remove(e)
            tw = OxmlElement("w:tcW"); tw.set(qn("w:w"), str(widths[j])); tw.set(qn("w:type"), "dxa"); tcPr.insert(0, tw)
            if i == 0:
                tb = OxmlElement("w:tcBorders"); tb.append(border("bottom", 8)); tcPr.append(tb)
            for p in cell.paragraphs:
                pf = p.paragraph_format; pf.first_line_indent = 0; pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (numeric[j] or (i == 0 and numeric[j])) else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(font_pt); r.font.name = "Times New Roman"
                    if i == 0: r.bold = True
    nxt = t.getnext()                                        # a little air between a table and the text below it
    if nxt is not None and nxt.tag == qn("w:p"):
        ppr = nxt.find(qn("w:pPr"))
        if ppr is None: ppr = OxmlElement("w:pPr"); nxt.insert(0, ppr)
        for e in ppr.findall(qn("w:spacing")): ppr.remove(e)
        sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "160"); ppr.append(sp)
    # keep the caption paragraph with the table
    prev = t.getprevious()
    if prev is not None and prev.tag == qn("w:p"):
        ppr = prev.find(qn("w:pPr"))
        if ppr is None: ppr = OxmlElement("w:pPr"); prev.insert(0, ppr)
        for e in ppr.findall(qn("w:keepNext")): ppr.remove(e)
        ppr.append(OxmlElement("w:keepNext"))

def build(md, out, ref, font_pt):
    subprocess.run(["pandoc", md, "-o", out, f"--reference-doc={ref}", "--resource-path=."], check=True)
    d = Document(out); add_page_numbers(d)
    for tb in d.tables: fix_table(tb, font_pt)
    for sec in d.sections:                                   # US Letter, 1-inch margins
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    text_w = Inches(6.5)
    for sh in d.inline_shapes:                               # figures span the full text width
        ratio = sh.height / sh.width; sh.width = text_w; sh.height = int(text_w * ratio)
    d.save(out); print(out, "| tables:", len(d.tables))

if __name__ == "__main__":
    build("manuscript_ms.md", "manuscript_ms.docx", "ref_misq.docx", 10)
    build("manuscript_misq.md", "manuscript_misq.docx", "ref_misq.docx", 10)
    build("online_appendix.md", "online_appendix_ms.docx", "ref_oa.docx", 9)
    build("ascollected_information_sheet.md", "ascollected_information_sheet.docx", "ref_oa.docx", 10)
