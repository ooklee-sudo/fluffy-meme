"""Build ISR_manuscript_draft_4.docx from draft 3 and the analysis outputs in results_real/.

usage: python build_draft4.py
Table 2 is filled from results_real/allocation_mechanism.txt so that the numbers cannot drift from the output.
"""
import copy, re
import docx
from docx.text.paragraph import Paragraph

d = docx.Document("ISR_manuscript_draft_3.docx")


# ---------------------------------------------------------------- helpers
def find(prefix, style=None):
    m = [p for p in d.paragraphs if p.text.startswith(prefix) and (style is None or p.style.name == style)]
    assert len(m) == 1, (prefix, len(m))
    return m[0]


def set_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def replace_in(p, old, new):
    assert old in p.text, old[:50]
    set_text(p, p.text.replace(old, new))


def add_after(ref, text, like):
    new = copy.deepcopy(like._p)
    ref._p.addnext(new)
    p = Paragraph(new, ref._parent)
    set_text(p, text)
    return p


M = lambda s: s.replace("-", "−")        # typographic minus

# ---------------------------------------------------------------- abstract
replace_in(find("Organizations are beginning"),
           "Binding authority made allocation more equal than chance.",
           "Binding authority made allocation more equal than chance, because holders favored agents with fewer "
           "tasks and less energy and gave themselves less, not more.")

# ---------------------------------------------------------------- 2.3 / 2.5 theory
body_like = find("The proposed interaction between authority")        # a plain Body Text paragraph
p17 = find("These accounts share a premise")
add_after(p17,
          "A cumulative-advantage account also bears on who stands. Agents that already hold resources or office "
          "bear a lower cost from a failed bid and gain a platform from a successful one, so resources and "
          "incumbency should raise the willingness to stand (Merton 1968; Arthur 1989). We examine this as an "
          "additional test in Section 4.4 and do not state it as a separate proposition.", body_like)

last2 = find("We also examine rights convertibility")
h = copy.deepcopy(find("2.4 Propositions")._p)
last2._p.addnext(h)
hp = Paragraph(h, last2._parent)
set_text(hp, "2.5 An account of how rights are exercised among agents (post hoc)")
t1 = add_after(hp,
    "The propositions above follow from accounts written for human actors. The results in Section 4 led us to a "
    "more specific account, which we state here so that it can be tested outside the data that suggested it. It "
    "was developed after the main results were examined and should be read as a hypothesis and not as a finding.",
    body_like)
t2 = add_after(t1,
    "In the human governance literature the gap between holding a decision right and using it is filled by "
    "norms, reputation, and the holder’s own interests. Among LLM agents the gap appears to be filled by two "
    "other things: the incentives that the system specification states explicitly, and the dispositions that the "
    "underlying model brings to the task. Four statements follow. First, incentive responsiveness: whether agents "
    "stand for office depends on whether the specification attaches a reward to it. Second, an equity default: "
    "when nothing in the prompt directs the holder to pursue its own gain, a holder with binding authority "
    "applies simple fairness heuristics, favoring agents who have received less and have less, and does not "
    "favor itself. Third, dispositional selection: who stands depends on the model’s disposition to take on "
    "roles, which varies with model family and capability, more than on the model’s success in the vote. "
    "Fourth, limited strategic depth: election rules and vote-for-access exchange matter only to the extent that "
    "agents reason about coalitions, and the agents in our setting, which answered directly without hidden "
    "reasoning, did so little.", body_like)
add_after(t2,
    "Each statement can fail. The first fails if rewards for office do not raise candidacy across model "
    "families. The second fails if holders favor themselves or past benefactors once the prompt makes their own "
    "interest salient. The third fails if differences between models appear mainly in success among candidates. "
    "The fourth fails if enabling reasoning makes election rules or vote-for-access exchange matter. Section 5.1 "
    "reports how far Study 1 supports each statement.", body_like)

# ---------------------------------------------------------------- 3.4 / 3.5 methods
p34 = find("We compare levels of each factor")
set_text(p34, p34.text.replace("As a benchmark for inequality",
    "As a check on the marginal tests, and to adjust for the other factors, we also fitted full-factorial models in "
    "which authority, rent, and population interact pairwise and the election rule enters as a main effect, with "
    "sum-to-zero coding so that each main-effect term is an average over the balanced design. Continuous outcomes "
    "used least squares with HC3 standard errors; shares of elections used quasi-binomial models and counts used "
    "quasi-Poisson models, both with robust standard errors. The five tests that correspond directly to "
    "propositions (rent on the share of elections with no candidate and on contested elections, authority on the "
    "Gini and on transfers, and election rule on re-election) are adjusted with the Holm method, and all other "
    "terms with the Benjamini–Hochberg method. Candidacy was modeled with a logit estimated by generalized "
    "estimating equations clustered on world. Success among candidates was modeled with a conditional logit "
    "within election, and the receipt of tasks with a conditional logit within day (Table 2). Intervals for these "
    "models come from bootstrap resampling of worlds (200 draws, 300 for the capability contrast, 100 for each "
    "rent column of Table 2, and 40 for its placebo column, whose interval is therefore rough). The "
    "interaction of authority with population was tested with a joint Wald test. As a benchmark for inequality"))

p35 = find("The design evolved over three stages")
set_text(p35, p35.text.replace(
    "and per-call costs are logged.",
    "and per-call costs are logged. The factorial models, the analyses of candidacy, winning, and allocation "
    "rules (Sections 4.2 and 4.4), the interaction of authority with population (Section 4.1), and the account in "
    "Section 2.5 were developed after the main results were examined and are exploratory."))

# ---------------------------------------------------------------- 4.1
p41 = find("Office rents changed the nature")
set_text(p41, p41.text +
    " The five tests tied to the propositions keep their conclusions in the factorial models after adjustment "
    "(rent, no candidate and contested elections: adjusted p < 0.0001; authority, Gini: p = 0.0007; authority, "
    "transfers: p < 0.0001; election rule, re-election: p = 0.29). One result was not anticipated and is "
    "exploratory. Binding authority moved candidacy in opposite directions depending on the population: the "
    "share of elections with no candidate fell from 35 to 20 percent in DeepSeek-only worlds, rose from 66 to 75 "
    "percent in Claude-only worlds, and was unchanged in mixed worlds (34 and 36 percent; joint test of the "
    "authority by composition interaction, χ²(2) = 17.2, p = 0.0002).")

# ---------------------------------------------------------------- 4.2
p42 = find("Binding authority did not concentrate access")
replace_in(p42, "(0.082, p < 0.001):", "(0.082, p < 0.001; Holm-adjusted p = 0.0007 in the factorial model):")
p3 = find("The proposed interaction between authority and rents")
replace_in(p3, "(difference in differences of 0.008; permutation p = 0.60)",
           "(difference in differences of 0.008; permutation p = 0.60; bootstrap interval from −0.017 to 0.033)")

# Table 2 numbers from the analysis output
txt = open("results_real/allocation_mechanism.txt").read()
secs = re.split(r"\n(?=== )", txt)
want = {"all": "Binding authority, office filled", "placebo": "Placebo", "r0": "Binding, rent=0", "r2": "Binding, rent=2"}
rows = {}
for key, head in want.items():
    s = [x for x in secs if x.startswith("== " + head)][0]
    for line in s.splitlines():
        m = re.match(r"^(cum_tasks_before|energy_before|is_holder|gifted_holder_before):.*?\s(-?\d+\.\d+)\s+(-?\d+\.\d+)\s+(-?\d+\.\d+)\s+(-?\d+\.\d+)\s*$", line)
        if m:
            rows[(key, m.group(1))] = (float(m.group(2)), float(m.group(4)), float(m.group(5)))
cell = lambda k, v: M(f"{rows[(k, v)][0]:.2f} [{rows[(k, v)][1]:.2f}, {rows[(k, v)][2]:.2f}]")
LAB = [("cum_tasks_before", "Tasks received so far (rotation)"),
       ("energy_before", "Energy at allocation (need)"),
       ("is_holder", "Agent is the holder (self-allocation)"),
       ("gifted_holder_before", "Energy given to the holder earlier (reciprocity)")]

m1 = add_after(p3,
    "Why was allocation under binding authority more equal than a lottery? Table 2 models which agents received "
    "tasks on the 2,025 days, in 70 worlds, on which the office was filled and authority was binding. Holders "
    f"favored agents who had received fewer tasks so far (coefficient {M('%.2f' % rows[('all','cum_tasks_before')][0])}, "
    f"95% interval {M('%.2f' % rows[('all','cum_tasks_before')][1])} to {M('%.2f' % rows[('all','cum_tasks_before')][2])}) "
    f"and agents with less energy ({M('%.2f' % rows[('all','energy_before')][0])}, {M('%.2f' % rows[('all','energy_before')][1])} "
    f"to {M('%.2f' % rows[('all','energy_before')][2])}). They also gave themselves less, not more: the coefficient for the "
    f"holder is {M('%.2f' % rows[('all','is_holder')][0])} ({M('%.2f' % rows[('all','is_holder')][1])} to "
    f"{M('%.2f' % rows[('all','is_holder')][2])}), and it is larger when office carried rents "
    f"({M('%.2f' % rows[('r2','is_holder')][0])}, {M('%.2f' % rows[('r2','is_holder')][1])} to {M('%.2f' % rows[('r2','is_holder')][2])}) "
    f"than when it did not ({M('%.2f' % rows[('r0','is_holder')][0])}, {M('%.2f' % rows[('r0','is_holder')][1])} to "
    f"{M('%.2f' % rows[('r0','is_holder')][2])}). Agents who had earlier transferred energy to the holder were not "
    f"favored in a way we can distinguish from zero ({M('%.2f' % rows[('all','gifted_holder_before')][0])}, "
    f"{M('%.2f' % rows[('all','gifted_holder_before')][1])} to {M('%.2f' % rows[('all','gifted_holder_before')][2])}). "
    "On days when allocation was a random draw, which serves as a placebo, the coefficients for tasks and energy "
    f"were close to zero and the coefficient for the holder was {M('%.2f' % rows[('placebo','is_holder')][0])} "
    f"({M('%.2f' % rows[('placebo','is_holder')][1])} to {M('%.2f' % rows[('placebo','is_holder')][2])}), which we take as "
    "the noise level of the method. The equalization under binding authority thus reflects the heuristics "
    "holders applied, rotation and need, and not self-dealing or the repayment of favors.", body_like)

cap = add_after(m1, "Table 2. Rules used when allocating tasks (conditional logit coefficients; 95% bootstrap "
                    "intervals from resampling worlds in brackets)", find("Table 1. Main effects"))
src = d.tables[0]._tbl
tbl = copy.deepcopy(src)
for tr in tbl.findall(docx.oxml.ns.qn("w:tr"))[2:]:
    tbl.remove(tr)
grid = tbl.find(docx.oxml.ns.qn("w:tblGrid"))
for gc in list(grid):
    grid.remove(gc)
widths = [2688, 1650, 1650, 1650, 1650]
for w in widths:
    gc = copy.deepcopy(src.find(docx.oxml.ns.qn("w:tblGrid"))[0]); gc.set(docx.oxml.ns.qn("w:w"), str(w)); grid.append(gc)
tbl.find(docx.oxml.ns.qn("w:tblPr")).find(docx.oxml.ns.qn("w:tblW")).set(docx.oxml.ns.qn("w:w"), str(sum(widths)))
trs = tbl.findall(docx.oxml.ns.qn("w:tr"))
hdr_tpl, row_tpl = trs[0], trs[1]
tbl.remove(row_tpl)


def fill_row(tr_tpl, vals):
    tr = copy.deepcopy(tr_tpl)
    tcs = tr.findall(docx.oxml.ns.qn("w:tc"))
    for tc in tcs[len(vals):]:
        tr.remove(tc)
    for tc, w, v in zip(tr.findall(docx.oxml.ns.qn("w:tc")), widths, vals):
        tc.find(docx.oxml.ns.qn("w:tcPr")).find(docx.oxml.ns.qn("w:tcW")).set(docx.oxml.ns.qn("w:w"), str(w))
        ts = tc.findall(".//" + docx.oxml.ns.qn("w:t"))
        ts[0].text = v
        for t in ts[1:]:
            t.text = ""
    return tr


tbl.remove(hdr_tpl)
tbl.append(fill_row(hdr_tpl, ["Predictor", "Binding, office filled", "Rent 0", "Rent 2", "Random draw (placebo)"]))
body_tr = copy.deepcopy(row_tpl)
for var, lab in LAB:
    tbl.append(fill_row(body_tr, [lab, cell("all", var), cell("r0", var), cell("r2", var), cell("placebo", var)]))
cap._p.addnext(tbl)

note = add_after(cap, "", body_like)
tbl.addnext(note._p)
set_text(note,
    "Note. Each day’s task recipients are modeled as a choice among the six agents. Negative coefficients for "
    "tasks received and for energy mean that agents with fewer tasks and less energy were favored. The first three "
    "columns use days with a filled office and binding authority (70 worlds); the last uses days on which tasks "
    "were drawn at random (advisory authority or a vacant office; 139 worlds).")
add_after(note,
    "These estimates rest on a subset of the data. Office is filled only when someone stands, which is itself an "
    "outcome of the design, and splitting by rent leaves 34 and 36 worlds. We treat the pattern as exploratory.",
    body_like)

# ---------------------------------------------------------------- 4.4
p44a = find("Model composition shaped which agents stood")
set_text(p44a,
    "Model composition shaped which agents stood and which won. In the 48 mixed Claude/DeepSeek worlds, agents "
    "powered by Claude stood in 105 elections and won 44, while DeepSeek agents stood 462 times and won 269. "
    "DeepSeek therefore held 86 percent of the terms filled, against its 82 percent share of candidacies. The raw "
    "win rate per candidacy was 58 percent for DeepSeek and 42 percent for Claude, but that gap mostly reflects "
    "who stood in which elections: in an election with only Claude candidates or only DeepSeek candidates, model "
    "family cannot matter. The informative comparison is the 62 contested elections in which both families stood. "
    "There Claude candidates won 26, against 25.6 expected if the winner were drawn uniformly from the "
    "candidates (conditional logit coefficient for Claude 0.03, 95% bootstrap interval −0.68 to 0.58, "
    "resampling worlds). The family difference is therefore one of candidacy and not of success once standing.")
p44b = find("The capability contrast within one model family")
set_text(p44b,
    "The capability contrast within one model family points the same way for candidacy and, more modestly, for "
    "success. The stronger DeepSeek V4 Pro stood 246 times and won 121; the weaker V4 Flash stood 161 times and "
    "won 56. In the 73 contested elections in which both stood, Pro won 47, against 38.4 expected under a uniform "
    "draw among candidates. The conditional logit coefficient for Pro is 0.59 (95% bootstrap interval 0.12 to "
    "1.16, resampling worlds), which allows for dependence among the elections of a world more carefully than a "
    "normal approximation would. In a candidacy model clustered on world, Pro had 1.7 times the odds of standing "
    "(coefficient 0.54, p < 0.001). P5 is thus supported as regards selection into candidacy and, more modestly, "
    "as regards success once standing.")
add_after(p44b,
    "Beyond model family, candidacy followed the pattern that cumulative-advantage accounts imply (full design, "
    "logit clustered on world; all p < 0.001). Agents who held office at the previous election had 2.9 times the "
    "odds of standing again (coefficient 1.06, 95% interval 0.87 to 1.25), and agents with more energy than their "
    "peers were more likely to stand: each step down the energy ranking lowered the odds of standing by about 7 "
    "percent (coefficient −0.07, interval −0.11 to −0.04). Rents raised the odds of standing by a factor "
    "of 1.7 (coefficient 0.54, interval 0.35 to 0.74), and Claude agents had about a quarter of the odds of "
    "DeepSeek agents (coefficient −1.44, interval −1.64 to −1.25). Resources and incumbency therefore feed "
    "candidacy, although, as Section 4.2 shows, holding office did not translate into capture of access.",
    body_like)

# ---------------------------------------------------------------- 5.1 / 5.2 / 5.3
p51 = find("The data do not support the expectation")
replace_in(p51, "but inequality of access did not.",
           "but inequality of access did not. Where holders had binding authority, they used it, on the evidence in "
           "Table 2, as a rotation device that favored agents with fewer tasks and less energy and gave themselves "
           "less, not more.")
p57 = find("The observation that agents often held office")
add_after(p57,
    "Section 2.5 stated four hypotheses, and Study 1 bears on each. Incentive responsiveness is supported: rents "
    "lowered the share of elections with no candidate from 55 to 34 percent and raised the odds of standing by a "
    "factor of 1.7. The equity default is supported: under binding authority, holders favored agents with fewer "
    "tasks and less energy and gave themselves less (Table 2). The self-allocation result runs against the "
    "expectation that rents tempt holders, since holders gave themselves less when office paid, and we note the "
    "direction without interpreting the difference between rent conditions further. Dispositional selection is "
    "supported for model family, where Claude agents stood far less often but won about as often as their share "
    "of contested elections implied, and modestly supported for capability, where V4 Pro both stood more often "
    "and won more often than V4 Flash. Limited strategic depth is consistent with the null results for election "
    "rules and for vote-for-access exchange, but Study 1 did not vary reasoning, so it cannot show that depth is "
    "the explanation. All four statements were formed after the data were seen and remain to be tested on new "
    "data.", body_like)

p52 = find("Three implications follow")
replace_in(p52, "This should not be generalized to settings",
           "The mechanism evidence suggests that this protection came from the fairness heuristics the holders "
           "applied and not from authority as such, so it may not survive prompts or models that make a holder’s "
           "own interest salient. It should not be generalized to settings")

p53 = find("The environment is small")
replace_in(p53,
    "Our permutation tests are marginal and uncorrected for multiple comparisons, and the three seeds per cell "
    "limit what can be said about interactions.",
    "The permutation tests in Table 1 are marginal and uncorrected for multiple comparisons. The factorial models "
    "with Holm and Benjamini–Hochberg adjustment address this for the tests tied to the propositions, but with "
    "six seeds per cell the design can detect an authority-by-rent interaction in the Gini only if it exceeds "
    "about 0.04.")
replace_in(p53,
    "The selection of candidacy over winning in Section 4.4 rests on a null-model comparison that treats "
    "elections as independent.",
    "The allocation analysis in Section 4.2 uses only the 70 worlds and 2,025 days in which the office was filled "
    "under binding authority, the filling of office is itself an outcome of candidacy, and the placebo coefficient "
    "for self-allocation was not exactly zero, so we read the self-allocation result as a difference of degree "
    "against a noise floor. The account in Section 2.5 and the analyses of candidacy, winning, and allocation "
    "rules were developed after the main results and have not been tested on new data. A preregistered "
    "replication with larger populations, varied prompt framing, and reasoning enabled is the necessary next step.")

d.save("ISR_manuscript_draft_4.docx")
print("saved")
