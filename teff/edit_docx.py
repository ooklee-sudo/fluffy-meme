"""Small helpers for revising the MDPI-template manuscript in place.

Mini-markup used in inserted text:
  *x*      italic            _{x}   subscript          ^{x}   superscript
  **x**    bold
New or changed text is coloured blue when `mark=True` (a revision-marked
version); with `mark=False` it is plain (a clean version).
"""
import copy
import re

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph

BLUE = RGBColor(0x00, 0x00, 0xFF)
TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{.+?\}|\^\{.+?\})")


class Editor:
    def __init__(self, path, mark=True):
        self.doc = docx.Document(path)
        self.mark = mark

    # --- lookup ---------------------------------------------------------
    def para(self, startswith, nth=0):
        hits = [p for p in self.doc.paragraphs
                if p.text.strip().startswith(startswith)]
        if len(hits) <= nth:
            raise KeyError(f"no paragraph starting with {startswith!r}")
        return hits[nth]

    def table_after(self, para):
        e = para._p.getnext()
        while e is not None and e.tag != qn("w:tbl"):
            e = e.getnext()
        return e

    # --- run construction -------------------------------------------------
    def _add_runs(self, p, text, base_bold=False, color=True):
        for tok in TOKEN.split(text):
            if not tok:
                continue
            bold = base_bold
            ital = sub = sup = False
            if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
                tok, bold = tok[2:-2], True
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                tok, ital = tok[1:-1], True
            elif tok.startswith("_{"):
                tok, sub = tok[2:-1], True
            elif tok.startswith("^{"):
                tok, sup = tok[2:-1], True
            r = p.add_run(tok)
            r.bold = bold or None
            r.italic = ital or None
            r.font.subscript = sub or None
            r.font.superscript = sup or None
            if self.mark and color:
                r.font.color.rgb = BLUE
        return p

    # --- paragraphs -------------------------------------------------------
    def insert_after(self, anchor, text, style=None, like=None, bold_lead=None):
        """Insert a paragraph after `anchor` (Paragraph or table element)."""
        el = anchor._p if isinstance(anchor, Paragraph) else anchor
        new_p = OxmlElement("w:p")
        el.addnext(new_p)
        p = Paragraph(new_p, self.doc._body)
        if like is not None:
            p.style = like.style
            if like.paragraph_format.alignment is not None:
                p.paragraph_format.alignment = like.paragraph_format.alignment
        if style:
            p.style = self.doc.styles[style]
        if bold_lead:
            self._add_runs(p, bold_lead + " ", base_bold=True)
        self._add_runs(p, text)
        return p

    def set_text(self, p, text, color=True):
        """Replace the whole paragraph text (formatting from mini-markup)."""
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        self._add_runs(p, text, color=color)
        return p

    def sub(self, p, old, new):
        """Replace `old` by `new` inside one run of paragraph p."""
        for r in p.runs:
            if old in r.text:
                before, after = r.text.split(old, 1)
                r.text = before
                new_runs = []
                tmp = OxmlElement("w:p")
                holder = Paragraph(tmp, p._parent)
                self._add_runs(holder, new)
                for nr in holder.runs:
                    new_runs.append(copy.deepcopy(nr._r))
                anchor = r._r
                for nr in new_runs:
                    anchor.addnext(nr)
                    anchor = nr
                if after:
                    tail = copy.deepcopy(r._r)
                    for t in tail.findall(qn("w:t")):
                        t.text = after
                        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    anchor.addnext(tail)
                return True
        raise KeyError(f"{old[:40]!r} not found inside a single run")

    def append(self, p, text):
        self._add_runs(p, text)
        return p

    def delete(self, p):
        p._p.getparent().remove(p._p)

    # --- tables -------------------------------------------------------------
    def make_table(self, anchor, rows, widths, font_sz=16, header_rows=1,
                   first_left=True):
        """Build a three-line table in the manuscript's style after `anchor`."""
        el = anchor._p if isinstance(anchor, Paragraph) else anchor
        tbl = OxmlElement("w:tbl")
        pr = OxmlElement("w:tblPr")
        for tag, attrs in (("w:tblW", {"w:w": str(sum(widths)), "w:type": "dxa"}),
                           ("w:tblInd", {"w:w": "0", "w:type": "dxa"}),
                           ("w:tblLayout", {"w:type": "fixed"})):
            e = OxmlElement(tag)
            for k, v in attrs.items():
                e.set(qn(k), v)
            pr.append(e)
        tbl.append(pr)
        grid = OxmlElement("w:tblGrid")
        for w in widths:
            g = OxmlElement("w:gridCol")
            g.set(qn("w:w"), str(w))
            grid.append(g)
        tbl.append(grid)
        n = len(rows)
        for ri, row in enumerate(rows):
            tr = OxmlElement("w:tr")
            trpr = OxmlElement("w:trPr")
            trpr.append(OxmlElement("w:cantSplit"))
            if ri < header_rows:
                trpr.append(OxmlElement("w:tblHeader"))
            tr.append(trpr)
            for ci, txt in enumerate(row):
                tc = OxmlElement("w:tc")
                tcpr = OxmlElement("w:tcPr")
                tw = OxmlElement("w:tcW")
                tw.set(qn("w:w"), str(widths[ci]))
                tw.set(qn("w:type"), "dxa")
                tcpr.append(tw)
                bd = OxmlElement("w:tcBorders")
                edges = []
                if ri == 0:
                    edges.append(("w:top", "8"))
                if ri == header_rows - 1:
                    edges.append(("w:bottom", "4"))
                if ri == n - 1:
                    edges.append(("w:bottom", "8"))
                for edge, sz in edges:
                    b = OxmlElement(edge)
                    b.set(qn("w:val"), "single")
                    b.set(qn("w:sz"), sz)
                    b.set(qn("w:space"), "0")
                    b.set(qn("w:color"), "000000")
                    bd.append(b)
                tcpr.append(bd)
                va = OxmlElement("w:vAlign")
                va.set(qn("w:val"), "center")
                tcpr.append(va)
                tc.append(tcpr)
                p = OxmlElement("w:p")
                tc.append(p)
                tr.append(tc)
                para = Paragraph(p, self.doc._body)
                para.style = self.doc.styles["MDPI_4.2_table_body"]
                para.alignment = 0 if (ci == 0 and first_left) else 1
                self._add_runs(para, str(txt), base_bold=(ri < header_rows))
                for r in para.runs:
                    r.font.size = docx.shared.Pt(font_sz / 2)
            tbl.append(tr)
        el.addnext(tbl)
        return tbl

    def save(self, path):
        self.doc.save(path)
