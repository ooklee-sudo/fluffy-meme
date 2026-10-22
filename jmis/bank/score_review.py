"""Score the reviewers' sheets against the key (Section 5.1).

Usage: python score_review.py key_and_map.csv reviewer1.xlsx reviewer2.xlsx
Reports, per reviewer: answered rows, agreement with the key; between reviewers: raw agreement and
Cohen's kappa; and the list of rows where a reviewer disagrees with the key or flagged a comment.
Rows both reviewers get wrong, or where they disagree with the key, are the items to inspect, revise or remove.
"""
import csv
import sys
from collections import Counter

from openpyxl import load_workbook


def read(path):
    ws = load_workbook(path, read_only=True).active
    out = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if r[0]:
            out[r[0]] = ((r[5] or "").strip().upper() or None, (r[6] or "").strip())
    return out


def kappa(x, y):
    n = len(x); po = sum(a == b for a, b in zip(x, y)) / n
    cx, cy = Counter(x), Counter(y)
    pe = sum(cx[k] * cy[k] for k in set(cx) | set(cy)) / n ** 2
    return (po - pe) / (1 - pe) if pe < 1 else float("nan")


def main():
    key = {r["row"]: r for r in csv.DictReader(open(sys.argv[1]))}
    revs = [read(p) for p in sys.argv[2:4]]
    for k, rv in enumerate(revs, 1):
        done = {r: a for r, (a, _) in rv.items() if a in ("A", "B", "C")}
        agree = sum(done[r] == key[r]["gold_letter"] for r in done)
        print(f"reviewer {k}: {len(done)}/{len(key)} answered, agrees with key on {agree} ({agree / max(len(done), 1):.1%})")
    both = [r for r in key if all(rv.get(r, (None,))[0] in ("A", "B", "C") for rv in revs)]
    x, y = ([rv[r][0] for r in both] for rv in revs)
    print(f"between reviewers: {len(both)} rows, raw agreement {sum(a == b for a, b in zip(x, y)) / max(len(both), 1):.1%}, Cohen's kappa {kappa(x, y):.3f}")
    print("\nrows to inspect (a reviewer disagrees with the key, or left a comment):")
    for r in key:
        ans = [rv.get(r, (None, ""))[0] for rv in revs]; com = [rv.get(r, (None, ""))[1] for rv in revs]
        if any(a != key[r]["gold_letter"] for a in ans) or any(com):
            print(f"  {r} {key[r]['item_id']} key={key[r]['gold_letter']} answers={ans} comments={[c for c in com if c]}")


if __name__ == "__main__":
    main()
