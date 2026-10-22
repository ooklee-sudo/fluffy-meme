"""Blind review packet for the two human reviewers (Section 5.1).

Makes, from bank.json:
  review/reviewer_sheet.xlsx   what each reviewer receives: shuffled anonymous rows, question and three
                               options, empty "Your answer (A/B/C)" and "Comments" columns. No key, no
                               domain, no tier.
  review/key_and_map.csv       authors only: row number -> item id, gold letter, domain, tier.
Usage: python review_packet.py BANKDIR OUTDIR [--seed N]
"""
import argparse
import csv
import json
import os
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("bankdir"); ap.add_argument("outdir"); ap.add_argument("--seed", type=int, default=7)
    a = ap.parse_args()
    bank = json.load(open(os.path.join(a.bankdir, "bank.json")))
    order = list(range(len(bank)))
    random.Random(a.seed).shuffle(order)
    os.makedirs(a.outdir, exist_ok=True)

    wb = Workbook(); ws = wb.active; ws.title = "Review"
    ws.append(["Row", "Question", "A", "B", "C", "Your answer (A/B/C)", "Comments (unclear, ambiguous, or more than one defensible answer)"])
    for c in ws[1]:
        c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
    rows = []
    for n, i in enumerate(order, 1):
        it = bank[i]; row = f"r{n:03d}"
        ws.append([row, it["stem"], it["options"]["A"], it["options"]["B"], it["options"]["C"], None, None])
        gold = next(L for L, r in it["roles"].items() if r == "gold")
        rows.append([row, it["item_id"], gold, it["domain"], it["tier"]])
    for col, w in zip("ABCDEFG", (7, 90, 16, 16, 16, 14, 40)):
        ws.column_dimensions[col].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"F2:F{len(bank) + 1}")
    ws.freeze_panes = "A2"
    wb.save(os.path.join(a.outdir, "reviewer_sheet.xlsx"))
    with open(os.path.join(a.outdir, "key_and_map.csv"), "w", newline="") as f:
        w = csv.writer(f); w.writerow(["row", "item_id", "gold_letter", "domain", "tier"]); w.writerows(rows)
    print(len(bank), "rows written to", a.outdir)


if __name__ == "__main__":
    main()
