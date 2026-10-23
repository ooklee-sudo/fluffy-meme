"""Converts main.tex to a Word file with pandoc after rewriting the IEEE-specific parts (citations, cross-references, theorem blocks)."""
import re, subprocess, os
HERE = os.path.dirname(os.path.abspath(__file__))
s = open(os.path.join(HERE, "main.tex")).read()

# bibliography: numbered by order of first citation, as IEEEtran does
bib = dict(re.findall(r"\\bibitem\{([^}]+)\}\s*(.*?)(?=\n\n\\bibitem|\n\\end\{thebibliography\})", s, flags=re.S))
body = s[: s.index("\\begin{thebibliography}")]
order = []
for m in re.finditer(r"\\cite\{([^}]+)\}", body):
    for k in m.group(1).split(","):
        k = k.strip()
        if k not in order: order.append(k)
num = {k: i + 1 for i, k in enumerate(order)}
body = re.sub(r"\\cite\{([^}]+)\}", lambda m: "[" + ", ".join(str(num[k.strip()]) for k in m.group(1).split(",")) + "]", body)
def clean(t):
    t = re.sub(r"\\emph\{([^}]*)\}", r"\\emph{\1}", t); t = t.replace("~", " ").replace("``", "\u201c").replace("''", "\u201d").replace("--", "\u2013")
    return re.sub(r"\s+", " ", t).strip()
refs = "\n\n".join(f"[{num[k]}] {clean(bib[k])}" for k in order)

# cross-references (IEEE numbering)
REF = {"tab:channels": "I", "tab:main": "II", "tab:mech": "III", "tab:reasoning": "IV", "fig:decomp": "1", "fig:reasoning": "2",
       "sec:limits": "VII", "sec:reasoning": "VI-E", "sec:model": "III", "prop:value": "1", "prop:sign": "3", "prop:bayes": "6"}
body = re.sub(r"\\ref\{([^}]+)\}", lambda m: REF[m.group(1)], body)
body = body.replace("Table~", "Table ").replace("Fig.~", "Fig. ").replace("Proposition~", "Proposition ").replace("Section~", "Section ").replace("Appendix~", "Appendix ")

# theorem blocks -> bold run-in headings
body = re.sub(r"\\setcounter\{proposition\}\{\d+\}\n", "", body)
nums = iter([1, 3, 6])
body = re.sub(r"\\begin\{proposition\}\[([^\]]*)\]\\label\{[^}]*\}", lambda m: f"\n\n\\textbf{{Proposition {next(nums)} ({m.group(1)}).}} ", body)
body = body.replace("\\end{proposition}", "\n\n")

# IEEE-specific commands
body = re.sub(r"\\IEEEPARstart\{(.)\}\{([^}]*)\}", r"\1\2", body)
body = re.sub(r"\\markboth\{[^}]*\}\{[^}]*\}\n", "", body)
body = re.sub(r"\\thanks\{[^}]*\}", "", body)
body = re.sub(r"\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}", lambda m: "\n\n\\emph{Index Terms}\u2014" + m.group(1).strip() + "\n\n", body, flags=re.S)
body = body.replace("\\begin{table*}", "\\begin{table}").replace("\\end{table*}", "\\end{table}")
body = re.sub(r"\\begin\{table\}\[t\]", r"\\begin{table}", body)
body = re.sub(r"\\begin\{figure\}\[t\]", r"\\begin{figure}", body)
body = re.sub(r"\\includegraphics\[[^\]]*\]\{figs/(fig_[a-z]+)\.pdf\}", lambda m: "\\includegraphics{" + os.path.join(HERE, "..", "..", "submission", "figures", {"fig_decomp": "Fig1_decomposition", "fig_reasoning": "Fig2_reasoning"}[m.group(1)]) + ".png}", body)
body = body.replace("\\maketitle", "")
body = body.replace("\\documentclass[journal]{IEEEtran}", "\\documentclass{article}")
body = re.sub(r"\\usepackage\{(cite|microtype)\}\n", "", body)
body = body.replace("\\begin{document}", "\\begin{document}\n\\maketitle")
body = body.replace("\\section*{Acknowledgments}", "\\section*{Acknowledgments}")
out = body + "\n\\section*{References}\n\n" + refs + "\n\n\\end{document}\n"
open(os.path.join(HERE, "main_word.tex"), "w").write(out)
subprocess.run(["pandoc", os.path.join(HERE, "main_word.tex"), "-f", "latex", "-o", os.path.join(HERE, "manuscript.docx"), "--resource-path", HERE], check=True)
print("ok", len(order), "references")

# ---- Word formatting
import docx
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d = docx.Document(os.path.join(HERE, "manuscript.docx"))
for sec in d.sections:
    sec.left_margin = sec.right_margin = Inches(1); sec.top_margin = sec.bottom_margin = Inches(1)
for st in d.styles:
    if st.type == 1:
        st.font.name = "Times New Roman"
        rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Times New Roman")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
for name, size, bold in (("Normal", 10, False), ("Body Text", 10, False), ("First Paragraph", 10, False), ("Compact", 10, False), ("Title", 16, True), ("Heading 1", 12, True),
                         ("Heading 2", 11, True), ("Heading 3", 10, True), ("Abstract", 10, True), ("Table Caption", 9, False), ("Image Caption", 9, False), ("Block Text", 9, False)):
    try:
        sty = next(x for x in d.styles if x.name == name)
        f = sty.font; f.size = Pt(size); f.bold = bold; f.color.rgb = docx.shared.RGBColor(0, 0, 0)
        col = sty.element.get_or_add_rPr().find(qn("w:color"))
        for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
            if col is not None and col.get(qn(a)) is not None: del col.attrib[qn(a)]
    except StopIteration:
        pass
def border(tbl):
    pr = tbl._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "000000"); b.append(e)
    pr.append(b)
for t in d.tables:
    border(t)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(8)
d.save(os.path.join(HERE, "manuscript.docx"))
print("formatted")
