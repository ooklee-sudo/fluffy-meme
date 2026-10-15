"""ANALYSIS_PLAN_HUMAN_VALIDATION.md. Run from study3/.
Step 1 (after coders 1 and 2 return):  python human_val/analysis_human_validation.py third-sheet c1.xlsx c2.xlsx
   -> human_val/validation_sheet_coder3.xlsx (only reviews with a disagreement on any outcome; text only)
Step 2:  python human_val/analysis_human_validation.py analyse c1.xlsx c2.xlsx [c3.xlsx]  -> human_val/results_human_validation.json"""
import sys, json, re, html, numpy as np, pandas as pd, openpyxl
from sklearn.metrics import cohen_kappa_score as ck
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
MODE = sys.argv[1]; K = pd.read_csv('human_val/validation_key_PRIVATE.csv')
OUTS = {'defect': lambda d: (d.defect == '1').astype(int), 'neg': lambda d: (d.overall == 'neg').astype(int), 'perfpos': lambda d: (d.performance == '1').astype(int)}
def read(path, itemcol, key=None):
    key = K if key is None else key
    rows = list(openpyxl.load_workbook(path, data_only=True)['Coding'].iter_rows(min_row=2, values_only=True))
    d = pd.DataFrame([(int(r[0]), *(str(x).strip().lower() if x is not None else '' for x in r[2:5])) for r in rows], columns=['item', 'overall', 'defect', 'performance'])
    d = d.merge(key[['rid', itemcol]].rename(columns={itemcol: 'item'}), on='item')
    for c, ok in (('overall', {'neg', 'neu', 'pos'}), ('defect', {'0', '1'}), ('performance', {'0', '-1', '1', '2'})):
        d[c] = d[c].str.replace(r'\.0$', '', regex=True); bad = ~d[c].isin(ok)
        if bad.any(): print(f'WARNING {path}: {bad.sum()} invalid/missing values in {c} (items {d.item[bad].tolist()[:10]})')
    return d.set_index('rid')
c1 = read(sys.argv[2], 'item_coder1'); c2 = read(sys.argv[3], 'item_coder2')
B = pd.DataFrame({f'{o}_{k}': f(c) for o, f in OUTS.items() for k, c in (('c1', c1), ('c2', c2))})
dis = [r for r in B.index if any(B.loc[r, f'{o}_c1'] != B.loc[r, f'{o}_c2'] for o in OUTS)]
if MODE == 'third-sheet':
    T = pd.read_json('sample_text.jsonl', lines=True).set_index('rid'); order = pd.Series(dis).sample(frac=1, random_state=3).tolist()
    wb = openpyxl.load_workbook(sys.argv[2]); wb.remove(wb['Coding']); ws = wb.create_sheet('Coding', 1)
    ws.append(['item', 'review text', 'Overall sentiment: neg / neu / pos', 'Defect reported: 1 = yes, 0 = no', 'Performance: 0 not mentioned, -1 negative, 1 positive, 2 neutral/mixed', 'notes'])
    for i, r in enumerate(order, 1): ws.append([i, html.unescape(re.sub(r'<br\s*/?>', '\n', str(T.loc[r, 'text'])))] + [None] * 4)
    for c, w in zip('ABCDEF', (6, 90, 18, 16, 26, 30)): ws.column_dimensions[c].width = w
    for row in ws.iter_rows(min_row=2): row[1].alignment = Alignment(wrap_text=True, vertical='top')
    for cell in ws[1]: cell.font = Font(bold=True)
    for col, opts in zip('CDE', ('"neg,neu,pos"', '"0,1"', '"0,-1,1,2"')):
        dv = DataValidation(type='list', formula1=opts); ws.add_data_validation(dv); dv.add(f'{col}2:{col}{len(order) + 1}')
    wb.save('human_val/validation_sheet_coder3.xlsx')
    pd.DataFrame({'rid': order, 'item_coder3': range(1, len(order) + 1)}).to_csv('human_val/validation_key3_PRIVATE.csv', index=False)
    print(len(order), 'reviews with a disagreement -> human_val/validation_sheet_coder3.xlsx'); sys.exit()
# ---- analyse
R = {'n': int(len(B)), 'n_disagreement': len(dis), 'kappa': {o: float(ck(B[f'{o}_c1'], B[f'{o}_c2'])) for o in OUTS}}
for o in OUTS: B[o] = np.where(B[f'{o}_c1'] == B[f'{o}_c2'], B[f'{o}_c1'], -1)
if len(sys.argv) > 4:
    c3 = read(sys.argv[4], 'item_coder3', pd.read_csv('human_val/validation_key3_PRIVATE.csv'))
    for o, f in OUTS.items():   # majority of three = the third coder's vote where coders 1 and 2 differ
        v3 = f(c3); m = (B[o] == -1) & B.index.isin(v3.index); B.loc[m, o] = v3.reindex(B.index[m]).values
unresolved = {o: int((B[o] == -1).sum()) for o in OUTS}; R['unresolved'] = unresolved
M = pd.read_csv('sample_meta.csv'); M['cell'] = 2 * M.treated + M.post; share = M.cell.value_counts(normalize=True)
COLS = ['overall', 'defect', 'price', 'durability', 'performance', 'fit', 'ease', 'design', 'service']
import glob
H = pd.concat([pd.read_csv(f, header=None, names=['rid'] + COLS) for f in sorted(glob.glob('out/H_*.csv'))]).drop_duplicates('rid').set_index('rid')
S = pd.read_csv('out_latent/S.csv').set_index('rid'); Z = pd.read_csv('out_latent/Z.csv').set_index('rid'); L = pd.read_csv('out_latent/L.csv').set_index('rid')
CL = {'defect': {'H': (H.defect.astype(int)), 'Z': Z.defect, 'L': L.defect},
      'neg': {'H': (H.overall == 'neg').astype(int), 'S': S.neg, 'Z': Z.neg},
      'perfpos': {'H': (H.performance == 1).astype(int), 'Z': Z.perfpos}}
LAT = json.load(open('results_latent.json'))['outcomes']
D = B.join(K.set_index('rid')[['cell']]); D['w'] = D.cell.map(share) / D.cell.map(D.cell.value_counts(normalize=True))
rng = np.random.default_rng(20261016)
def rates(d, x, y):
    w = d.w.values; Se = (w * (x == 1) * (y == 1)).sum() / (w * (y == 1)).sum(); Sp = (w * (x == 0) * (y == 0)).sum() / (w * (y == 0)).sum(); return Se, Sp, Se + Sp - 1
R['outcomes'] = {}
for o in OUTS:
    d = D[D[o] >= 0]; y = d[o].values; o_ = {'n_resolved': int(len(d)), 'prevalence_by_cell': d.groupby('cell')[o].mean().tolist()}
    p = d.groupby('cell')[o].mean(); o_['human_did'] = float((p[3] - p[2]) - (p[1] - p[0]))
    for k, s in CL[o].items():
        x = s.reindex(d.index).values.astype(int); Se, Sp, J = rates(d, x, y); bs = []
        for _ in range(2000):
            i = rng.integers(0, len(d), len(d)); bs.append(rates(d.iloc[i], x[i], y[i]))
        bs = np.array(bs); ci = np.nanpercentile(bs, [2.5, 97.5], 0)
        lat = LAT[o]['latent']; key = {'H': f'H_{o}', 'S': 'S_neg', 'Z': f'Z_{o}', 'L': 'L_defect'}[k]
        o_[k] = dict(Se=Se, Sp=Sp, J=J, J_ci=ci[:, 2].tolist(), latent_J=lat['J'].get(key), latent_Se=lat['Se'].get(key), latent_Sp=lat['Sp'].get(key))
    h = o_['H']; o_['validated'] = bool(h['J_ci'][0] <= h['latent_J'] <= h['J_ci'][1] and abs(h['latent_J'] - h['J']) <= 0.10)
    o_['corrected_did_human_J'] = LAT[o]['naive_fe']['b'] / h['J']
    R['outcomes'][o] = o_
json.dump(R, open('human_val/results_human_validation.json', 'w'), indent=1, default=float)
print(json.dumps({'kappa': R['kappa'], 'unresolved': unresolved}, indent=0))
for o, v in R['outcomes'].items():
    print(f"\n{o}: validated={v['validated']}  human DiD {v['human_did']:+.3f}  corrected DiD (FE / human J) {v['corrected_did_human_J']:+.4f}")
    for k in CL[o]: print(f"  {k}: human Se {v[k]['Se']:.2f} Sp {v[k]['Sp']:.2f} J {v[k]['J']:.2f} [{v[k]['J_ci'][0]:.2f},{v[k]['J_ci'][1]:.2f}]   latent J {v[k]['latent_J']:.2f}")
