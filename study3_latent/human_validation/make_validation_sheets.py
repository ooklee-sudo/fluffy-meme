"""Blind human-validation sheets for ANALYSIS_PLAN_HUMAN_VALIDATION.md.
Draws 50 reviews per DiD cell (seed 20261016) from the reviews coded by all four classifiers (the zero-shot subsample),
writes one sheet per coder (different random order, text only) and a private key. Run from study3/."""
import random, re, html, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
M = pd.read_csv('sample_meta.csv'); M['cell'] = 2 * M.treated + M.post
Z = pd.read_csv('out_latent/Z.csv'); T = pd.read_json('sample_text.jsonl', lines=True).set_index('rid')
pool = M[M.rid.isin(Z.rid)]
pick = pool.groupby('cell').sample(n=50, random_state=20261016)
key = pick[['rid', 'cell', 'treated', 'post']].reset_index(drop=True)
FIELDS = [('overall', 'Overall sentiment: neg / neu / pos', '"neg,neu,pos"'),
          ('defect', 'Defect reported: 1 = yes, 0 = no', '"0,1"'),
          ('performance', 'Performance: 0 not mentioned, -1 negative, 1 positive, 2 neutral/mixed', '"0,-1,1,2"')]
GUIDE_KO = ['[코더 안내]',
            '1. 각 리뷰의 글만 읽고 세 항목을 코딩하세요. 별점, 제품 정보, 다른 사람의 답은 보지 마세요.',
            '2. 검색하거나 다른 코더와 상의하지 마세요. 혼자 작업하세요.',
            '3. 항목 정의는 Codebook 시트를 따르세요. 애매하면 가장 가까운 값을 고르고 notes 칸에 짧게 적으세요.',
            '4. 셀을 클릭하면 선택 목록이 나옵니다. 빈 칸 없이 200건 모두 코딩해 주세요(약 1.5~2시간).',
            '5. 이 연구의 가설이나 목적은 알려 드리지 않습니다. 끝나면 파일을 그대로 돌려주세요.']
GUIDE_EN = ['[Instructions]', '1. Code each review from its text alone. Do not look up ratings, products, or other coders\' answers.',
            '2. Work alone; do not search or discuss items.', '3. Follow the Codebook sheet. If unsure, choose the closest value and add a short note.',
            '4. Use the drop-down lists. Please code all 200 reviews (about 1.5-2 hours).']
CODEBOOK = ['Codebook (identical to the one given to the LLM coders; three fields only)', '',
            'overall: overall sentiment of the reviewer toward the product: neg, neu, pos. Judge from the text only.',
            'defect: 1 if the reviewer reports a product failure, malfunction, defect, or damaged/incorrect/wrong item; else 0.',
            'performance: how well it works; does the job, power, effectiveness, noise level.',
            '   0 = not mentioned, -1 = mentioned negatively, 1 = mentioned positively, 2 = mentioned, neutral or mixed.']
def sheet(path, seed):
    order = key.sample(frac=1, random_state=seed).reset_index(drop=True)
    wb = Workbook(); ws = wb.active; ws.title = 'Coding'
    ws.append(['item', 'review text'] + [f[1] for f in FIELDS] + ['notes'])
    for i, r in enumerate(order.itertuples(), 1): ws.append([i, html.unescape(re.sub(r'<br\s*/?>', '\n', str(T.loc[r.rid, 'text'])))] + [None] * 4)   # readable line breaks; text otherwise as given to the LLM
    for c, w in zip('ABCDEF', (6, 90, 18, 16, 26, 30)): ws.column_dimensions[c].width = w
    for row in ws.iter_rows(min_row=2): row[1].alignment = Alignment(wrap_text=True, vertical='top')
    for cell in ws[1]: cell.font = Font(bold=True); cell.fill = PatternFill('solid', fgColor='DDEBF7'); cell.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = 'C2'
    for col, (_, _, opts) in zip('CDE', FIELDS):
        dv = DataValidation(type='list', formula1=opts, allow_blank=True); ws.add_data_validation(dv); dv.add(f'{col}2:{col}{len(order) + 1}')
    g = wb.create_sheet('Instructions', 0)
    for line in GUIDE_KO + [''] + GUIDE_EN: g.append([line])
    g.column_dimensions['A'].width = 120
    cb = wb.create_sheet('Codebook')
    for line in CODEBOOK: cb.append([line])
    cb.column_dimensions['A'].width = 120
    wb.save(path); return order.assign(item=range(1, len(order) + 1))
oa = sheet('human_val/validation_sheet_coder1.xlsx', 1); ob = sheet('human_val/validation_sheet_coder2.xlsx', 2)
key = key.merge(oa[['rid', 'item']].rename(columns={'item': 'item_coder1'}), on='rid').merge(ob[['rid', 'item']].rename(columns={'item': 'item_coder2'}), on='rid')
key.to_csv('human_val/validation_key_PRIVATE.csv', index=False)
print(len(key), key.cell.value_counts().sort_index().to_dict())
