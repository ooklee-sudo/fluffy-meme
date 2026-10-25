"""Minimal helpers on top of python-docx: Times New Roman 12, mini-markup for italics/sub/superscripts, tables, figures."""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TOK = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{.*?\}|\^\{.*?\})")


def new_doc(spacing=2.0, footer=True, running=None):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, a, Inches(1))
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.line_spacing = spacing; pf.space_after = Pt(0); pf.space_before = Pt(0)
    d._spacing = spacing
    if footer:
        p = s.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if t:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
            else:
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
            r._r.append(e)
        r.font.size = Pt(10)
    if running:
        hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run(running); hr.font.size = Pt(9); hr.italic = True
    return d


def runs(p, text, size=None, bold=False, italic=False):
    for part in TOK.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2]); r.bold = True; r.italic = italic
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = not italic; r.bold = bold
        elif part.startswith("_{"):
            r = p.add_run(part[2:-1]); r.font.subscript = True; r.bold = bold; r.italic = italic
        elif part.startswith("^{"):
            r = p.add_run(part[2:-1]); r.font.superscript = True; r.bold = bold; r.italic = italic
        else:
            r = p.add_run(part); r.bold = bold; r.italic = italic
        if size:
            r.font.size = Pt(size)
    return p


def P(d, text, indent=True, align=None, size=None, after=0, before=0, bold=False, italic=False, spacing=None, keep=False):
    p = d.add_paragraph()
    pf = p.paragraph_format
    if indent:
        pf.first_line_indent = Inches(0.5)
    if spacing is not None:
        pf.line_spacing = spacing
    pf.space_after = Pt(after); pf.space_before = Pt(before)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if keep:
        pf.keep_with_next = True
    runs(p, text, size=size, bold=bold, italic=italic)
    return p


def H1(d, text):
    p = P(d, text, indent=False, bold=True, before=12, after=6, spacing=1.15, keep=True)
    return p


def H2(d, text):
    return P(d, text, indent=False, bold=True, italic=True, before=8, after=4, spacing=1.15, keep=True)


def TITLE(d, text, size=16):
    return P(d, text, indent=False, align="center", bold=True, size=size, after=12, spacing=1.15)


def EQ(d, text, num=None):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5; pf.space_before = Pt(4); pf.space_after = Pt(4)
    pf.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    runs(p, text)
    if num:
        p.add_run("\t(" + str(num) + ")")
    return p


def BLOCK(d, label, text, italic_body=True):
    """Proposition / Lemma / Definition block."""
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.3); pf.right_indent = Inches(0.3)
    pf.space_before = Pt(6); pf.space_after = Pt(6); pf.line_spacing = 1.5
    r = p.add_run(label + " "); r.bold = True
    runs(p, text, italic=italic_body)
    return p


def PROOF(d, text, label="Proof."):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Inches(0.5); pf.space_after = Pt(6)
    r = p.add_run(label + " "); r.italic = True
    runs(p, text)
    return p


def BUL(d, items, size=None):
    for t in items:
        p = d.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = d._spacing if size is None else 1.15
        runs(p, t, size=size)


def _border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom"):
        if edge in kw:
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(kw[edge])); e.set(qn("w:color"), "000000")
            b.append(e)
    tcPr.append(b)


def TABLE(d, caption, rows, widths=None, note=None, size=10, align=None):
    cp = P(d, caption, indent=False, bold=False, before=10, after=4, spacing=1.15, keep=True)
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
            if j > 0 and (align is None or align[j] == "c"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(p, str(val), size=size, bold=(i == 0))
            if i == 0:
                _border(c, top=12, bottom=6)
            if i == len(rows) - 1:
                _border(c, bottom=12)
            if widths:
                c.width = Inches(widths[j])
    if note:
        P(d, note, indent=False, size=9, before=3, after=8, spacing=1.0)
    else:
        P(d, "", indent=False, spacing=1.0)
    return t


def FIG(d, path, caption, width=6.2):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(8)
    p.add_run().add_picture(path, width=Inches(width))
    P(d, caption, indent=False, size=10, before=2, after=10, spacing=1.0)


def PB(d):
    d.add_page_break()


def REF(d, text):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.5); pf.first_line_indent = Inches(-0.5)
    pf.line_spacing = d._spacing; pf.space_after = Pt(0)
    runs(p, text)
