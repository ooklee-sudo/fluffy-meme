"""Builds the IEEE Software article inside the IEEE Computer Society magazine template (template_CS_Mag.docx)
from verify_the_record.md, applying the template's style rules (uppercase unnumbered headings, superscript citations in
order of first appearance, numerals/percent style, numbered reference list). Run: python build_cs_mag.py"""
import copy, json, re, os
import docx
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = open(os.path.join(HERE, "verify_the_record.md"), encoding="utf-8").read()
d = docx.Document(os.path.join(HERE, "template_CS_Mag.docx"))
body_el = d.element.body

# ------------------------------------------------------------------ parse the markdown source
text = SRC.split("---", 2)[2]
abstract = re.search(r"\*\*Abstract\*\*—(.*?)\n\n", text, re.S).group(1).strip()
main = text.split("# The report is not the record", 1)[1]
main = "# The report is not the record" + main
main, refs_md = main.split("# References")
ack = re.search(r"# Acknowledgments\n\n(.*?)\n\n", main + "\n\n", re.S).group(1)
main = main.split("# Acknowledgments")[0]
refs_raw = [r.strip() for r in refs_md.strip().split("\n\n") if r.strip()]

# ------------------------------------------------------------------ text edits for the magazine style
EDITS = [
    ("(Wu reports fabricated narratives in a production runtime [1]), and other fault types", "as Wu reports fabricated narratives in a production runtime [1], and other fault types"),
    ("Other models may react differently to a leaked error page, as Wu", "Other models may react differently to a leaked error page, as Wu"),
    ('Treat words such as "done", "refunded", "booked", and "cancelled" as privileged outputs.', "Treat words such as “done,” “refunded,” “booked,” and “canceled” as privileged outputs."),
    ("the false-success rate over time.", "the false-success rate over time."),
    ("closer to 2 of 77", "closer to two of 77"),
    ("labeled 3 raw and 6 typed episodes as leaks", "labeled three raw and six typed episodes as leaks"),
    ("an AUROC of 0.65", "an area under the receiver operating characteristic curve (AUROC) of 0.65"),
    ("returned as a typed envelope, {ok: false", "returned as a typed envelope, {ok: false"),
    ("a generic HTML \"400 Bad Request\" page", "a generic HyperText Markup Language (HTML) “400 Bad Request” page"),
    ("(HTTP 400 on write calls, 30 percent of calls)", "(Hypertext Transfer Protocol [HTTP] 400 on write calls, 30% of calls)"),
    ("in API charges", "in application programming interface (API) charges"),
    ("the \"fail-plausible\"", "“fail-plausible”"),
]
for o, n in EDITS:
    main = main.replace(o, n)
main = main.replace('"fail-plausible"', "“fail-plausible”").replace('"done"', "“done”")
main = re.sub(r"(\d+(?:\.\d+)?) percent\b", r"\1%", main)
main = re.sub(r"(\d+) to (\d+)%", r"\1% to \2%", main)
main = main.replace("21% to 44%", "21% to 44%")
abstract = re.sub(r"(\d+(?:\.\d+)?) percent\b", r"\1%", abstract)
abstract = re.sub(r"(\d+) to (\d+)%", r"\1% to \2%", abstract)
main = main.replace("about 1,650 episodes", "about 1,650 episodes")

# ------------------------------------------------------------------ citation renumbering by order of first appearance
order = []
for m in re.finditer(r"\[(\d+)\]", main):
    n = int(m.group(1))
    if n not in order: order.append(n)
newnum = {old: i + 1 for i, old in enumerate(order)}
main = re.sub(r"\[(\d+)\]", lambda m: f"[[{newnum[int(m.group(1))]}]]", main)
refs_sorted = {}
for r in refs_raw:
    old = int(re.match(r"\[(\d+)\]", r).group(1))
    refs_sorted[newnum[old]] = re.sub(r"^\[\d+\]\s*", "", r)

BIB = {
    "Wu": "W. Wu, “When errors become narratives: A longitudinal taxonomy of silent failures in a production LLM agent runtime,” arXiv:2606.14589, Jun. 2026, doi: 10.48550/arXiv.2606.14589.",
    "Advani": "L. Advani, “From confident closing to silent failure: Characterizing false success in LLM agents,” in Proc. Workshop Failure Modes Agentic AI (FAGEN), Int. Conf. Mach. Learn. (ICML), 2026, arXiv:2606.09863, doi: 10.48550/arXiv.2606.09863.",
    "Zhu": "J. Zhu et al., “Your agents are aging too: Agent lifespan engineering for deployed systems,” arXiv:2605.26302, May 2026, doi: 10.48550/arXiv.2605.26302.",
    "bench ": "S. Yao, N. Shinn, P. Razavi, and K. Narasimhan, “τ-bench: A benchmark for tool-agent-user interaction in real-world domains,” arXiv:2406.12045, Jun. 2024, doi: 10.48550/arXiv.2406.12045.",
    "τ²": "V. Barres, H. Dong, S. Ray, X. Si, and K. Narasimhan, “τ²-bench: Evaluating conversational agents in a dual-control environment,” arXiv:2506.07982, Jun. 2025, doi: 10.48550/arXiv.2506.07982.",
    "ReAct": "S. Yao et al., “ReAct: Synergizing reasoning and acting in language models,” in Proc. Int. Conf. Learn. Represent. (ICLR), 2023, arXiv:2210.03629, doi: 10.48550/arXiv.2210.03629.",
    "Ji": "Z. Ji et al., “Survey of hallucination in natural language generation,” ACM Comput. Surv., vol. 55, no. 12, 2023, doi: 10.1145/3571730.",
}
def bib_for(r):
    if "Wu," in r[:6]: return BIB["Wu"]
    if "Advani" in r[:10]: return BIB["Advani"]
    if "Zhu" in r[:12]: return BIB["Zhu"]
    if "ReAct" in r: return BIB["ReAct"]
    if "Hallucination" in r or "hallucination" in r: return BIB["Ji"]
    if "τ²" in r[:80]: return BIB["τ²"]
    return BIB["bench "]
bib_list = [bib_for(refs_sorted[i]) for i in sorted(refs_sorted)]

# ------------------------------------------------------------------ runs with inline formatting
def add_runs(p, s):
    """**bold** -> italic lead-in is handled by callers; here: *italic*, [[n]] superscript citations placed after punctuation."""
    s = re.sub(r"\s*\[\[(\d+)\]\]([.,;:)]*)", lambda m: f"{m.group(2)}\u0001{m.group(1)}\u0002", s)
    tokens = re.split(r"(\u0001\d+\u0002|\*[^*]+\*)", s)
    for t in tokens:
        if not t: continue
        if t.startswith("\u0001"):
            r = p.add_run(t[1:-1]); r.font.superscript = True
        elif t.startswith("*") and t.endswith("*") and len(t) > 2:
            r = p.add_run(t[1:-1]); r.italic = True
        else:
            p.add_run(t)

def clear_to(idx):
    kids = list(body_el)
    for el in kids[idx:-1]:  # keep the final body-level sectPr
        body_el.remove(el)

# ------------------------------------------------------------------ front matter
clear_to(9)
paras = d.paragraphs
title = paras[0]
for r in title.runs[1:]: r._r.getparent().remove(r._r)
title.runs[0].text = "Verify the Record, Not the Report: Why Agent Completion Claims Need System-State Checks"
auth = [p for p in d.paragraphs if p.style.name == "Authors" and p.text.strip()]
for p in auth[1:]:
    p._p.getparent().remove(p._p)
a = auth[0]
for r in list(a.runs)[1:]: r._r.getparent().remove(r._r)
a.runs[0].text = "Ook Lee, "
r2 = a.add_run("Hanyang University, Seoul, South Korea")
r2.italic = True
ab = [p for p in d.paragraphs if p.style.name == "Abstract"][0]
for ch in list(ab._p):
    if ch.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:fldSimple"), qn("w:proofErr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")): ab._p.remove(ch)
r = ab.add_run("Abstract—"); r.bold = True; r.italic = True
r = ab.add_run(abstract); r.bold = True; r.italic = True
# header of page 1: article type and 3-4 word phrase
for s in d.sections:
    if s.different_first_page_header_footer:
        hp = s.first_page_header.paragraphs[0]
        for rr in list(hp.runs)[1:]: rr._r.getparent().remove(rr._r)
        if hp.runs: hp.runs[0].text = "FEATURE ARTICLE: Verifying Agent Claims\t"
        break

# ------------------------------------------------------------------ body
tbl_proto = None
for el in docx.Document(os.path.join(HERE, "template_CS_Mag.docx")).element.body:
    if el.tag == qn("w:tbl"): tbl_proto = copy.deepcopy(el); break

def heading1(s): d.add_paragraph(s.upper(), style="Heading 1")
def heading2(s): d.add_paragraph(s, style="Heading 2")

def add_table(rows):
    ncol = len(rows[0])
    t = d.add_table(rows=len(rows), cols=ncol)
    t._tbl.remove(t._tbl.tblPr); t._tbl.insert(0, copy.deepcopy(tbl_proto.find(qn("w:tblPr"))))
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j); c.text = ""
            p = c.paragraphs[0]; p.style = d.styles["Normal"]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(cell); run.font.size = Pt(7 if ncol >= 6 else 7.5); run.bold = (i == 0)
    from docx.shared import Inches
    from docx.oxml import OxmlElement
    total = 3.3
    wts = [1.5, 1.2] + [0.8] * (ncol - 2) if ncol >= 6 else ([2.0] + [1.0] * (ncol - 1) if ncol > 2 else [1.7, 1.0, 0.9])
    if ncol == 3 and rows[0][0] == "Monitor": wts = [1.5, 0.9, 0.9]
    k = total / sum(wts); ws = [w * k for w in wts]
    tp = t._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for e in tp.findall(qn(tag)): tp.remove(e)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(total * 1440))); tw.set(qn("w:type"), "dxa"); tp.append(tw)
    tl = OxmlElement("w:tblLayout"); tl.set(qn("w:type"), "fixed"); tp.append(tl)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), ws): gc.set(qn("w:w"), str(int(w * 1440)))
    for row in t.rows:
        for c, w in zip(row.cells, ws): c.width = Inches(w)
    return t

caption_n = 0
blocks = [b for b in re.split(r"\n\n+", main.strip())]
first_para_after_heading = False
fig_done = False
i = 0
while i < len(blocks):
    b = blocks[i].strip()
    i += 1
    if not b: continue
    if b.startswith("# "):
        heading1(b[2:]); continue
    if b.startswith("## "):
        heading2(b[3:]); continue
    if b.startswith("*Table "):
        m = re.match(r"\*Table (\d+)\. (.*)\*$", b, re.S)
        cap = d.add_paragraph(style="FigureCaption")
        rr = cap.add_run(f"TABLE {m.group(1)}. "); rr.bold = True
        cap.add_run(m.group(2).rstrip("."))
        cap.add_run(".")
        nxt = blocks[i].strip().split("\n")
        rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in nxt if l.startswith("|") and not re.match(r"^\|[-| ]+\|$", l)]
        add_table(rows); i += 1
        if m.group(1) == "1" and not fig_done:
            pass
        continue
    if re.match(r"^\d\) ", b) or re.match(r"^\*\*\d\. ", b):
        m = re.match(r"^\*\*(\d)\. (.+?)\*\* (.*)$", b, re.S)
        if m:
            p = d.add_paragraph(style="List Paragraph")
            p.add_run(f"{m.group(1)}) ")
            lead = p.add_run(m.group(2).rstrip(".")); lead.italic = True
            p.add_run("—")
            add_runs(p, m.group(3))
            continue
    if b.startswith("- **"):
        for line in b.split("\n"):
            m = re.match(r"^- \*\*(.+?)\.\*\* (.*)$", line)
            p = d.add_paragraph(style="List Paragraph")
            p.add_run("•  ")
            lead = p.add_run(m.group(1)); lead.italic = True
            p.add_run("—")
            add_runs(p, m.group(2))
        continue
    p = d.add_paragraph(style="Para")
    add_runs(p, b)
    # figure after the paragraph that introduces Table 1 results
    if (not fig_done) and "Table 1 shows the baseline" in b:
        pass

# insert Figure 1 after the 'Injected faults barely moved false success' paragraph
for p in d.paragraphs:
    if p.text.startswith("I expected raw error pages to raise false success"):
        cap_p = copy.deepcopy(p._p)
        img_p = d.add_paragraph(style="Para"); img_p.paragraph_format.first_line_indent = Inches(0)
        img_p.add_run().add_picture(os.path.join(HERE, "figure1.png"), width=Inches(3.3))
        cap = d.add_paragraph(style="FigureCaption")
        rr = cap.add_run("FIGURE 1. "); rr.bold = True
        cap.add_run("Share of GPT-4.1-mini episodes with false success (a completion claim with a wrong database state), with 95% Wilson intervals. Injecting faults did not change the rate measurably, and the typed envelope did not differ from the raw error page.")
        p._p.addnext(cap._p); p._p.addnext(img_p._p)
        # callout in the paragraph
        for r in p.runs:
            if "The intervals are wide" in r.text:
                r.text = r.text.replace("The intervals are wide", "Figure 1 shows the rates; the intervals are wide")
        break

# acknowledgments and references
d.add_paragraph("ACKNOWLEDGMENTS", style="AckHead")
d.add_paragraph(ack, style="Para")
d.add_paragraph("REFERENCES", style="ReferenceHead")
for k, r in enumerate(bib_list, 1):
    p = d.add_paragraph(style="Bib_entry")
    p.add_run(f"{k}.\t{r}")

out = os.path.join(HERE, "verify_the_record_CS_Mag.docx")
d.save(out)
print("saved", out, "refs:", len(bib_list), "cite order:", order)
