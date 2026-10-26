import docx
from docx.shared import Pt, Inches
d = docx.Document()
for s in d.sections:
    s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(1)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11.5)
def p(t="", bold=False, after=8):
    para = d.add_paragraph(); r = para.add_run(t); r.bold = bold
    para.paragraph_format.space_after = Pt(after); return para
p("[Date]", after=12)
p("Editor-in-Chief\nIEEE Software\nIEEE Computer Society", after=12)
p("Re: Submission of a feature article, “Verify the Record, Not the Report: Why Agent Completion Claims Need System-State Checks”", bold=True, after=12)
p("Dear Editor,")
p("I am pleased to submit the enclosed article for consideration as a feature article in IEEE Software. It is approximately 2,900 words including tables, figure caption, and references.")
p("Teams are now giving language-model agents write access to ticketing, ordering, and booking systems, and the agents report their own results in fluent prose. The article asks a practical question: which control actually narrows the gap between what an agent says it did and what the system of record shows? I ran about 1,650 episodes of a public customer-service benchmark (τ-bench) with two agent models, injected connector faults into 30% of write calls, and compared a raw HTML error page with a typed error envelope. False success was common even without injected faults, and neither the faults nor the typed envelope changed it measurably, because agents usually told the customer when a write had failed. The control that did work was a check of the agent’s claim against the database state. A transcript-reading monitor, whether rule-based or a second language model, missed most false successes.")
p("I believe the article suits IEEE Software’s readers, who are practicing engineers and architects deciding how to put agents into production workflows. It ends with a short, ordered list of controls to build first, and it states plainly where the evidence is limited: one benchmark, one failure type, and one injection scheme. The code and per-episode data will be made available at [repository link].")
p("Disclosures:", bold=True, after=4)
for t in [
    "Originality. The article is not under consideration elsewhere. An earlier, substantially different version, without the experiment reported here, was submitted to IEEE IT Professional and was declined at the editorial stage. The present manuscript was rewritten around new experimental evidence, and I verified every cited source against the original publication.",
    "Use of AI tools. In line with IEEE policy, I disclose that I used Claude (Anthropic), accessed through Claude Code, to write the experiment code, assist with data analysis, and help draft and edit the text. I reviewed the outputs, checked them against the data, and take full responsibility for the content. The disclosure also appears in the Acknowledgments.",
    "Conflicts of interest and funding. [AUTHOR TO COMPLETE]",
]:
    para = d.add_paragraph(t, style="List Bullet"); para.paragraph_format.space_after = Pt(4)
p("", after=4)
p("Thank you for considering the article. I would be glad to revise it in response to the reviewers’ comments.")
p("Sincerely,", after=24)
p("Ook Lee\nHanyang University, Seoul, South Korea\nooklee@hanyang.ac.kr")
d.save("cover_letter_IEEE_Software.docx")
