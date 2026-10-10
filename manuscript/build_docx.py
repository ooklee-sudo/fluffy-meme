"""Convert an elsarticle-style manuscript (.tex) into a Word (.docx) file via pandoc.

Citations (natbib author-year) are expanded to text, section/table/proposition numbers are resolved and
fixed as text, simple math in tables becomes plain text with subscripts, math elsewhere becomes native Word equations.

Usage:  python build_docx.py dss_manuscript.tex dss_manuscript.docx
Requires: pandoc, python-docx.
"""
import re
import subprocess
import sys
import tempfile
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CM = {'benaroch1999case': ('Benaroch and Kauffman', '1999'), 'benaroch2002managing': ('Benaroch', '2002'),
      'biderman2023pythia': ('Biderman et al.', '2023'), 'broadie1997continuity': ('Broadie et al.', '1997'),
      'cobbe2021gsm8k': ('Cobbe et al.', '2021'), 'dixit1994investment': ('Dixit and Pindyck', '1994'),
      'houlsby2019parameter': ('Houlsby et al.', '2019'), 'hu2022lora': ('Hu et al.', '2022'),
      'ilharco2023editing': ('Ilharco et al.', '2023'), 'kirkpatrick2017overcoming': ('Kirkpatrick et al.', '2017'),
      'mcdonald1986value': ('McDonald and Siegel', '1986'), 'merton1976option': ('Merton', '1976'),
      'sculley2015hidden': ('Sculley et al.', '2015')}


def find_close(s, i):
    d = 0
    for k in range(i, len(s)):
        if s[k] == '{' and (k == 0 or s[k - 1] != '\\'):
            d += 1
        elif s[k] == '}' and (k == 0 or s[k - 1] != '\\'):
            d -= 1
            if d == 0:
                return k
    raise ValueError('unbalanced braces')


def ay(keys):
    return '; '.join(f'{CM[k.strip()][0]}, {CM[k.strip()][1]}' for k in keys.split(','))


def nm(keys):
    ks = [k.strip() for k in keys.split(',')]
    return ' and '.join(f'{CM[k][0]} ({CM[k][1]})' for k in ks)


def to_pandoc_tex(src):
    fm0, fm1 = src.index('\\begin{frontmatter}'), src.index('\\end{frontmatter}')
    front = src[fm0:fm1]
    title = re.search(r'\\title\{(.*?)\}\n', front, re.S).group(1).strip()
    abstract = re.search(r'\\begin\{abstract\}(.*?)\\end\{abstract\}', front, re.S).group(1).strip()
    kw = re.search(r'\\begin\{keyword\}(.*?)\\end\{keyword\}', front, re.S).group(1)
    kw = re.sub(r'\s*\\sep\s*', '; ', kw).strip()
    am = re.search(r'\\author\[[^\]]*\]\{', front)
    author = front[am.end():find_close(front, am.end() - 1)]
    author = re.sub(r'\\corref\{[^}]*\}', '', author).strip()
    address = re.search(r'\\address\[[^\]]*\]\{(.*?)\}\n', front, re.S).group(1).strip()
    cort = re.search(r'\\cortext\[[^\]]*\]\{(.*?)\}\n', front, re.S)
    cort = cort.group(1).strip() if cort else ''
    body = src[fm1 + len('\\end{frontmatter}'):]
    body = body[:body.index('\\end{document}')]

    # bibliography -> References section
    bs, bt, be = body.index('\\bibliographystyle'), body.index('\\begin{thebibliography}'), body.index('\\end{thebibliography}')
    items = re.findall(r'\\bibitem\[[^\]]*\]\{[^}]*\}\s*(.*?)(?=\\bibitem|\Z)', body[bt:be], re.S)
    refs = '\\section*{References}\n\n' + '\n\n'.join(' '.join(x.split()) for x in items) + '\n\n'
    body = body[:bs] + refs + body[be + len('\\end{thebibliography}'):]

    # label numbering
    labels, sec, tab, prop, appendix = {}, 0, 0, 0, False
    for m in re.finditer(r'\\section\{|\\begin\{table\}|\\begin\{proposition\}|\\appendix', body):
        tok, pos = m.group(), m.start()
        if tok == '\\appendix':
            appendix = True
        elif tok == '\\section{':
            sec += 1
            c = find_close(body, pos + len('\\section'))
            lab = re.match(r'\\label\{([^}]+)\}', body[c + 1:])
            if lab:
                labels[lab.group(1)] = 'A' if appendix else str(sec)
        elif tok == '\\begin{table}':
            tab += 1
            e = body.index('\\end{table}', pos)
            for l in re.findall(r'\\label\{([^}]+)\}', body[pos:e]):
                labels[l] = str(tab)
        else:
            prop += 1
            mm = re.match(r'\\begin\{proposition\}(\[[^\]]*\])?\s*\\label\{([^}]+)\}', body[pos:])
            if mm:
                labels[mm.group(2)] = str(prop)

    # section headings with numbers
    out, i, sec, appendix = [], 0, 0, False
    while True:
        m = re.search(r'\\section\{|\\appendix', body[i:])
        if not m:
            out.append(body[i:])
            break
        st = i + m.start()
        out.append(body[i:st])
        if m.group() == '\\appendix':
            appendix = True
            i = st + len('\\appendix')
            continue
        c = find_close(body, st + len('\\section'))
        ttl = body[st + len('\\section{'):c]
        lab = re.match(r'\\label\{[^}]+\}', body[c + 1:])
        if appendix:
            out.append('\\section*{Appendix A. ' + ttl + '}')
        else:
            sec += 1
            out.append('\\section*{' + f'{sec}. ' + ttl + '}')
        i = c + 1 + (lab.end() if lab else 0)
    body = ''.join(out)

    # propositions
    pc = [0]

    def prop_sub(m):
        pc[0] += 1
        t = m.group(1)[1:-1] if m.group(1) else ''
        return f'\\textbf{{Proposition {pc[0]}' + (f' ({t})' if t else '') + '.} '
    body = re.sub(r'\\begin\{proposition\}(\[[^\]]*\])?\s*(?:\\label\{[^}]+\})?', prop_sub, body)
    body = body.replace('\\end{proposition}', '')
    body = re.sub(r'\\label\{[^}]+\}', '', body)

    # \paragraph -> bold run-in
    while True:
        i = body.find('\\paragraph{')
        if i < 0:
            break
        c = find_close(body, i + len('\\paragraph'))
        body = body[:i] + '\\textbf{' + body[i + len('\\paragraph{'):c] + '}' + body[c + 1:]

    # citations
    body = re.sub(r'\\citep\[([^\]]*);\]\[\]\{([^}]+)\}', lambda m: f'({m.group(1)}; {ay(m.group(2))})', body)
    body = re.sub(r'\\citep\{([^}]+)\}', lambda m: f'({ay(m.group(1))})', body)
    body = re.sub(r'\\citet\{([^}]+)\}', lambda m: nm(m.group(1)), body)
    body = re.sub(r'\\citealt\{([^}]+)\}', lambda m: ay(m.group(1)), body)

    # cross references
    body = re.sub(r'~\\ref\{([^}]+)\}', lambda m: ' ' + labels.get(m.group(1), '??' + m.group(1)), body)
    body = re.sub(r'\\ref\{([^}]+)\}', lambda m: labels.get(m.group(1), '??' + m.group(1)), body)

    # plain-text symbols; table cleanup
    body = body.replace('$\\to$', '→').replace('$^\\dagger$', '†').replace('$-$', '−')
    body = body.replace('\\smallskip\\footnotesize ', '')

    def clean_math(m):
        x = m.group(1).strip()
        x = x.replace('\\tau', 'τ').replace('\\to', '→').replace('\\,', ' ').replace('\\ ', ' ')
        x = re.sub(r'\\mathrm\{([^}]*)\}', r'\1', x)
        x = re.sub(r'\^\*', '*', x)
        x = re.sub(r'_\{([^}]*)\}', r'\\textsubscript{\1}', x)
        x = re.sub(r'_([A-Za-z0-9])', r'\\textsubscript{\1}', x)
        x = x.replace('=', ' = ').replace('  ', ' ')
        return x

    def fix_tab(m):
        t = re.sub(r'\$([^$]*)\$', clean_math, m.group(0))
        return t.replace('\\Gamma', 'Γ').replace('\\theta', 'θ').replace('  = ', ' = ')
    body = re.sub(r'\\begin\{tabular\}.*?\\end\{tabular\}', fix_tab, body, flags=re.S)

    head = ('\\documentclass[12pt]{article}\n\\usepackage{amsmath,amssymb,booktabs}\n'
            '\\newcommand{\\E}{\\mathbb{E}}\n\\begin{document}\n'
            f'\\title{{{title}}}\n\\author{{{author}\\\\{address}' + (f'\\\\{cort}' if cort else '') + '}\n\\date{}\n\\maketitle\n'
            f'\\begin{{abstract}}\n{abstract}\n\\end{{abstract}}\n\\noindent\\textbf{{Keywords:}} {kw}\n\n')
    return head + body + '\n\\end{document}\n', title


def setfont(style, size=None, bold=None, color=True, name='Times New Roman'):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color:
        f.color.rgb = RGBColor(0, 0, 0)


def style_docx(path_in, path_out, title):
    d = Document(path_in)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, m, Inches(1))
    sec.header_distance = Inches(0.5)
    sec.footer_distance = Inches(0.5)
    sec.gutter = Inches(0)
    for st in d.styles:
        if st.type == 1:
            n = st.name.lower()
            if n == 'title':
                setfont(st, 16, True)
            elif n.startswith('heading 1'):
                setfont(st, 13, True)
            elif n.startswith('heading'):
                setfont(st, 12, True)
            elif n in ('table caption', 'caption'):
                setfont(st, 11, False)
            else:
                setfont(st, 12)
        elif st.type == 2:
            setfont(st, 12, color=False)
    for name in ('Body Text', 'First Paragraph', 'Compact', 'Abstract'):
        try:
            pf = d.styles[name].paragraph_format
            pf.line_spacing, pf.space_after = 1.5, Pt(6)
            if name != 'Compact':
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        except KeyError:
            pass
    for name in ('Title', 'Author', 'Date'):
        try:
            d.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except KeyError:
            pass
    for st in d.styles:
        if st.type == 1 and st.name.lower() == 'heading 1':
            st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(14), Pt(6)

    for t in d.tables:
        n = len(t.columns)
        total = 6.5
        first = min(2.4, total - 0.62 * (n - 1)) if n >= 5 else 2.4
        ws = [first] + [(total - first) / (n - 1)] * (n - 1)
        tbl = t._tbl
        pr = tbl.tblPr
        lay = pr.find(qn('w:tblLayout'))
        if lay is None:
            lay = OxmlElement('w:tblLayout')
            pr.append(lay)
        lay.set(qn('w:type'), 'fixed')
        tw = pr.find(qn('w:tblW'))
        if tw is None:
            tw = OxmlElement('w:tblW')
            pr.append(tw)
        tw.set(qn('w:type'), 'dxa')
        tw.set(qn('w:w'), str(int(total * 1440)))
        for gc, w in zip(tbl.find(qn('w:tblGrid')).findall(qn('w:gridCol')), ws):
            gc.set(qn('w:w'), str(int(w * 1440)))
        for r in t.rows:
            for c, w in zip(r.cells, ws):
                c.width = Inches(w)
        b = pr.find(qn('w:tblBorders'))
        if b is not None:
            pr.remove(b)
        b = OxmlElement('w:tblBorders')
        for edge in ('top', 'bottom'):
            e = OxmlElement(f'w:{edge}')
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), '12')
            e.set(qn('w:color'), '000000')
            b.append(e)
        pr.append(b)
        for c in t.rows[0].cells:
            tcPr = c._tc.get_or_add_tcPr()
            tb = OxmlElement('w:tcBorders')
            e = OxmlElement('w:bottom')
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), '6')
            e.set(qn('w:color'), '000000')
            tb.append(e)
            tcw = tcPr.find(qn('w:tcW'))
            (tcw.addnext(tb) if tcw is not None else tcPr.insert(0, tb))
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        run.font.size = Pt(9 if n >= 8 else 10)
                        run.font.name = 'Times New Roman'
    d.core_properties.title = title
    d.core_properties.author = ''
    d.save(path_out)


def main(tex_path, docx_path):
    src = open(tex_path, encoding='utf8').read()
    ptex, title = to_pandoc_tex(src)
    with tempfile.TemporaryDirectory() as tmp:
        tp, raw = os.path.join(tmp, 'm.tex'), os.path.join(tmp, 'raw.docx')
        open(tp, 'w', encoding='utf8').write(ptex)
        r = subprocess.run(['pandoc', '-f', 'latex', '-t', 'docx', tp, '-o', raw], capture_output=True, text=True)
        if r.returncode:
            raise SystemExit(r.stderr)
        if '??' in ptex:
            print('warning: unresolved reference', re.findall(r'\?\?\S+', ptex)[:3])
        style_docx(raw, docx_path, title)
        if os.environ.get('KEEP_PANDOC_TEX'):
            open(os.environ['KEEP_PANDOC_TEX'], 'w', encoding='utf8').write(ptex)
    print('wrote', docx_path)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
