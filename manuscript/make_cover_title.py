"""Build the cover letter and the title page (Word) for the IJOC submission.

Usage:  python make_cover_title.py
Writes: cover_letter_ijoc.docx, title_page_ijoc.docx (in this directory).
Bracketed items are for the author to complete or confirm before submission.
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

TITLE = ("When Is Adapter Transfer Worth Its Cost? A Decision Model for Maintaining Fine-Tuned "
         "Language-Model Adapters Across Foundation-Model Upgrades")
RUNTITLE = "When Is Adapter Transfer Worth Its Cost?"
AUTHOR = "Ook Lee"
AFFIL = "Department of Information System, Hanyang University, Seoul, Republic of Korea"
EMAIL = "[corresponding-author e-mail]"
KEYWORDS = ("AI asset management; foundation model upgrades; parameter-efficient fine-tuning; "
            "maintenance policy; optimal stopping; adapter transfer")
REPO = "https://github.com/ooklee-sudo/fluffy-meme"
DATE = "October 11, 2026"


def base_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, m, Inches(1))
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Times New Roman")
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.line_spacing = 1.15
    return d


def para(d, text="", bold=False, italic=False, align=None, size=None, after=None, keep=None):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def labeled(d, label, text):
    p = d.add_paragraph()
    r = p.add_run(label + " ")
    r.bold = True
    p.add_run(text)
    return p


def cover_letter():
    d = base_doc()
    para(d, DATE, after=14)
    para(d, "[Name of the Editor-in-Chief or Area Editor]", after=0)
    para(d, "INFORMS Journal on Computing", after=14)
    p = para(d, "", after=10)
    r = p.add_run("Re: Submission of “" + TITLE + "”")
    r.bold = True
    para(d, "Dear Editor,")
    para(d, "I submit the enclosed manuscript for consideration in the INFORMS Journal on Computing.")
    para(d, "Organizations that build on foundation models they do not control face a recurring computational "
            "decision. After every provider upgrade they must choose between retraining the low-rank adapters that "
            "hold their task knowledge, continuing to use them, and applying a cheaper transfer procedure. The paper "
            "studies this decision as a computing problem and makes four contributions that fit the scope of the journal.")
    para(d, "First, we propose Anchored Calibration Transfer (ACT), a closed-form, label-free correction of low-rank "
            "adapters, and we give its algorithm and complexity. Second, we show by profiling that a straightforward "
            "implementation re-runs the network once for every layer, so that its cost grows with the square of the "
            "depth, and we give a cached implementation that leaves the output unchanged. On the same trained adapters "
            "the two versions select identical anchoring strengths and give identical held-out loss, and the cached "
            "version is 1.4 to 2.4 times faster at 70 to 410 million parameters. Third, we embed the method in an optimal "
            "stopping model that yields a closed-form retraining threshold and a break-even recovery that any transfer "
            "method must exceed to be worth its cost. Fourth, we report measurements on open model pairs: copying the old "
            "adapter recovers 96 percent of the retraining benefit after a small upgrade but only 32 to 36 percent after "
            "the largest, ACT’s gain over copying rises with the size of the upgrade, and the break-even rule shows "
            "that ACT pays for itself after sufficiently large upgrades at 70 and 160 million parameters but not at 410 million.")
    para(d, "The paper states the limits of its evidence. The experiments use models of 70 to 410 million parameters, a "
            "single task, and one or two seeds per configuration, and an instruction-tuned pair on which ACT gives no gain "
            "remains unexplained. We report these results as they are, including the cases in which the method does not "
            "pay for itself.")
    labeled(d, "Software and data.", "The code, simulation seeds, and result files are available at " + REPO +
            " [Author to confirm that the repository is public or deposited in line with the IJOC Software and Data "
            "Disclosure Policy, and replace the link by a permanent one if required].")
    labeled(d, "Use of generative AI.", "During the preparation of this work I used Claude (Anthropic) to assist with "
            "writing and running the simulation and experiment code, analyzing results, and drafting and language-editing "
            "the manuscript. I reviewed and edited the content as needed and take full responsibility for it. "
            "[Author to confirm that this statement is accurate and consistent with the journal’s policy.]")
    labeled(d, "Originality.", "[Author to confirm: the manuscript is original, has not been published, and is not under "
            "consideration elsewhere. State any conflict of interest, or that there is none.]")
    labeled(d, "Suggested reviewers.", "[Optional; to be completed by the author. Names are not supplied here.]")
    para(d, "Thank you for considering the manuscript.", after=14)
    para(d, "Sincerely,", after=22)
    para(d, AUTHOR, bold=True, after=0)
    para(d, "Professor", after=0)
    para(d, AFFIL, after=0)
    para(d, EMAIL, after=0)
    d.core_properties.title = "Cover letter"
    d.core_properties.author = ""
    d.save("cover_letter_ijoc.docx")


def title_page():
    d = base_doc()
    para(d, "Title Page", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, after=14)
    labeled(d, "Title:", TITLE)
    labeled(d, "Running title:", RUNTITLE)
    labeled(d, "Author:", AUTHOR)
    labeled(d, "Affiliation:", AFFIL)
    labeled(d, "Corresponding author:", AUTHOR + ", " + EMAIL)
    labeled(d, "ORCID:", "[Author to add, if any]")
    labeled(d, "Manuscript type:", "Original research article")
    labeled(d, "Journal:", "INFORMS Journal on Computing")
    labeled(d, "Key words:", KEYWORDS)
    labeled(d, "Length:", "Abstract of about 280 words; main text of about 7,000 words including tables and captions "
                          "(counted from the Word version; recount from the final file).")
    labeled(d, "Funding:", "[Author to confirm: state any funding, or that the research received no specific funding.]")
    labeled(d, "Author contributions:", "[Author to confirm. Single-author manuscript.]")
    labeled(d, "Conflict of interest:", "[Author to confirm: the author declares no known competing interests, or state them.]")
    labeled(d, "Acknowledgments and use of generative AI:", "During the preparation of this work the author used Claude "
            "(Anthropic) to assist with writing and running the simulation and experiment code, analyzing results, and "
            "drafting and language-editing the manuscript. The author reviewed and edited the content as needed and takes "
            "full responsibility for the content of the publication. [Author to confirm that this statement is accurate.]")
    labeled(d, "Software and data availability:", "The code, simulation seeds, and result files are available at " + REPO +
            ". [Author to confirm that the repository is accessible to reviewers and to deposit a permanent copy as the "
            "IJOC Software and Data Disclosure Policy requires.]")
    para(d, "")
    para(d, "Note: if the journal reviews manuscripts anonymously, the manuscript file must not contain the author’s "
            "name or affiliation; in that case this page carries them and the manuscript carries none.",
         italic=True, size=10)
    d.core_properties.title = "Title page"
    d.core_properties.author = ""
    d.save("title_page_ijoc.docx")


if __name__ == "__main__":
    cover_letter()
    title_page()
    print("wrote cover_letter_ijoc.docx, title_page_ijoc.docx")
