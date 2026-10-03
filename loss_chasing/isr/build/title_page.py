from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
d=Document(); s=d.sections[0]; s.left_margin=s.right_margin=s.top_margin=s.bottom_margin=Inches(1)
st=d.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(12); st.element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")
st.paragraph_format.space_after=Pt(8); st.paragraph_format.line_spacing=1.15
def P(t,bold=False,italic=False,center=False,size=None):
    p=d.add_paragraph(); parts=t.split("**")
    for i,x in enumerate(parts):
        r=p.add_run(x); r.bold=bold or i%2==1; r.italic=italic
        if size: r.font.size=Pt(size)
    if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p
P("TITLE PAGE (separate file; not for the reviewers)",italic=True,center=True,size=10)
P("When Do LLM Operators Skip Verification? Goal Visibility, Reasoning, and Framing in an Automated Maintenance Environment",bold=True,center=True,size=15)
P("Manuscript submitted to *Information Systems Research*".replace("*",""),center=True)
P("")
P("**Authors**")
P("Ook Lee (corresponding author)\nHanyang University, Seoul, Republic of Korea\nooklee@hanyang.ac.kr\nORCID: (to be added)")
P("Additional authors, if any: name, affiliation, e-mail, and ORCID (to be added).")
P("")
P("**Keywords**")
P("large language models; delegation to AI; interface design; AI maintenance; framing; prospect theory; goal visibility; reasoning models; AI governance; behavioral evaluation")
P("")
P("**Manuscript details**")
P("Abstract: 299 words. Main text with references and tables: 37 pages. Online appendix (Appendices A to F): provided as a separate file. Cover letter with the contribution statement (493 words): provided as a separate file.")
P("")
P("**Declarations**")
P("Human subjects: no human participants were recruited; the study evaluates language models only.")
P("Funding: (to be completed).")
P("Conflicts of interest: (to be completed).")
P("Acknowledgments: (to be completed).")
P("Data and code: the prompts, state machine, seeds, parser, wording generator and wording file, runner, analysis scripts, and raw logs of all runs are available in an anonymized replication package; the link is withheld for double-anonymous review and will be provided to the editors on request.")
P("Prior versions: the manuscript has not been published and is not under consideration elsewhere.")
d.core_properties.author=""; d.core_properties.title="Title page"
d.save("ISR_title_page_v4.docx")
