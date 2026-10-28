"""Cover letter for the ISR submission."""
import json, datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent; ROOT = HERE.parent
INJ = json.load(open(ROOT / "inject" / "inject_results.json")); CODE = json.load(open(ROOT / "inject_code" / "code_results.json"))
n_runs = sum(v["n_runs"] for v in INJ.values()) + sum(v["n_runs"] for v in CODE.values())
gpt = abs(INJ["gpt61terra"]["success"]["blame_minus_control"]["diff"]) * 100
TITLE = "When Should the Firm Take the Task Back? Error Dynamics and the Economics of Monitoring Autonomous AI Agents"

d = Document()
for s in d.sections: s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(0.8)
cp = d.core_properties; cp.author = ""; cp.comments = ""
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.0
def P(t, bold=False, italic=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic; return p

P("[Date]")
P("Editor-in-Chief\nInformation Systems Research\nINFORMS")
P(f"Re: Submission of the manuscript “{TITLE}”", bold=True)
P("Dear Editor-in-Chief,")
P(f"I am pleased to submit the manuscript “{TITLE}” for consideration as a research article in Information Systems "
  "Research. The manuscript is original, has not been published, and is not under consideration at any other journal.")
P("Firms are beginning to delegate multistep work, such as customer service and software maintenance, to autonomous AI agents "
  "while keeping a person who can take the task back. Every such deployment needs a rule that says when control should be "
  "withdrawn, and today these rules are set largely by intuition. The manuscript asks when a firm should take a task back from an "
  "agent and what its monitoring system should track to make that decision well.")
P("The paper combines an economic model with large-scale evidence. The model treats delegation as an optimal stopping problem and "
  "distinguishes three reasons why errors matter: they may trigger further errors (contagion), tasks may differ in difficulty "
  "(heterogeneity), and an error may derail a task so that its output is wrong (derailment). Each channel calls for a different "
  "monitoring statistic, and under derailment the optimal rule is a weighted count of errors with a closed-form threshold. The value "
  "of monitoring is greatest for tasks at the margin of delegation. The evidence comes from 80,036 public trajectories of a "
  f"software-engineering agent and from two randomized fault-injection experiments, in customer service and in the repair of real "
  f"software bugs, with five agents, including two recent commercial models, and {n_runs:,} runs. Most of the error clustering "
  "visible in observational data reflects heterogeneity rather than contagion. Transient failures rarely harm recent models, but "
  f"errors that misreport the state of the environment derail them, lowering task success by up to {gpt:.0f} percentage points. "
  "What an error says matters more than how many errors occur.")
P("Contribution statement", bold=True)
P("The manuscript makes four contributions.")
def ITEM(label, text):
    p = d.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
    r = p.add_run(label + " "); r.bold = True; p.add_run(text)
ITEM("(1) Theoretical contribution.", "It formalizes the revocation of a delegation to an autonomous AI agent as an optimal "
     "stopping problem and characterizes the optimal hand-off rule under three distinct error mechanisms: contagion, task "
     "heterogeneity, and derailment. It shows that the statistic a monitoring system should track differs across mechanisms "
     "(a recency-weighted intensity, an error rate, or a weighted error count), derives the optimal threshold in closed form "
     "under derailment, and proves that the value of monitoring is greatest for tasks at the margin of delegation.")
ITEM("(2) Empirical contribution.", "It provides causal evidence on how errors affect autonomous agents. Randomized fault "
     f"injection in two task domains, with five agents and {n_runs:,} runs, separates contagion from persistence and shows that "
     "the damage done by an error depends on what the error says: transient failures rarely harm recent models, whereas errors that "
     f"misreport the state of the environment lower task success by up to {gpt:.0f} percentage points. A large observational "
     "corpus further shows that ignoring task heterogeneity overstates contagion by a factor of three to four.")
ITEM("(3) Methodological contribution.", "It introduces a design that firms and researchers can use to measure the parameters "
     "that govern oversight: a random-effects self-exciting model that separates heterogeneity from contagion in agent logs, and "
     "a low-cost fault-injection protocol that estimates the probability that each kind of error derails a task. All code and "
     "run-level data are provided in a replication package.")
ITEM("(4) Practical contribution.", "It gives managers implementable design rules: count agent errors by type rather than "
     "in total, calibrate the hand-off threshold for each task domain, and target monitoring investment at tasks near the margin "
     "of delegation. It also shows that the accuracy of the data agents read is a determinant of agent performance, which makes "
     "data governance part of AI governance.")
P("These contributions speak directly to three conversations in Information Systems Research: human–AI delegation, which the "
  "paper extends from the decision to delegate to the decision to revoke a delegation; IS control, which it connects to the "
  "economics of information; and information quality, which it links to the governance of agentic AI.")
P("All data used are public or were generated by the author. The code, run-level results, and analysis scripts will be made "
  "available in a public replication package. The author used a generative AI assistant to help write analysis code, run the "
  "experiments, and edit the text, as disclosed on the title page; the author designed the study, verified all analyses, and takes "
  "full responsibility for the content. The author has no conflicts of interest to declare and received no external funding for "
  "this research.")
P("Thank you for considering the manuscript. I look forward to hearing from you.")
P("Sincerely,")
P("Ook Lee\nDepartment of Information Systems\nHanyang University, Seoul, Republic of Korea\nooklee@hanyang.ac.kr")
d.save(HERE / "ISR_Cover_Letter.docx"); print("saved")
