import json
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

R = json.load(open("results.json"))
FL, AL, AQ = R["full_linear"], R["active_linear"], R["active_quad"]
PN = json.load(open("panel_results.json")); PM = PN["models"]
V = json.load(open("llm/validation_results.json"))
pm = lambda m, v, k="b": PM[m][v][k]
doc = Document()
sec = doc.sections[0]
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, m, Inches(1))
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Times New Roman")
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE; st.paragraph_format.space_after = Pt(0)
for n, sz in (("Heading 1", 12), ("Heading 2", 12)):
    h = doc.styles[n]; h.font.name = "Times New Roman"; h.font.size = Pt(sz); h.font.bold = True
    h.font.italic = (n == "Heading 2"); h.font.color.rgb = None
    h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(0); h.paragraph_format.keep_with_next = True

def P(text="", bold=False, italic=False, align=None, indent=True, single=False):
    p = doc.add_paragraph()
    if indent and not align: p.paragraph_format.first_line_indent = Inches(0.5)
    if align: p.alignment = align
    if single: p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(text); r.bold = bold; r.italic = italic
    return p
H1 = lambda t: doc.add_heading(t, 1)
H2 = lambda t: doc.add_heading(t, 2)

def table(rows, widths=None, note=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ""; pr = c.paragraphs[0]
            pr.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            if j: pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = pr.add_run(v); r.font.size = Pt(10); r.bold = (i == 0)
    if note: P(note, italic=True, indent=False, single=True).runs[0].font.size = Pt(10)

f = lambda x, d=3: f"{x:,.{d}f}"
P("Overdispersion in Consumer-Generated Content Volume: Why Count Models Matter and Where LLM-Derived Covariates Fit",
  bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Research Note", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

H1("Abstract")
P(f"Consumer-generated content (CGC) volume is routinely analyzed with ordinary least squares or Poisson regression, "
  f"although such series are discrete, non-negative, and typically far more variable than a Poisson process allows. "
  f"Using {R['n_reviews']:,} Amazon product reviews (February 2009–March 2019) aggregated to a calendar-complete series of "
  f"{R['n_months']} months, we show that monthly review volume has a variance-to-mean ratio of about {R['var_mean']:,.0f} and that a "
  f"negative binomial specification fits dramatically better than a Poisson specification (ΔAIC > 60,000). We also show that "
  f"conclusions about the time trend are fragile when a log-linear trend is imposed on a boom-and-bust series: on the full series the "
  f"negative binomial slope is more than three times the Poisson slope, and on the dense-collection window the Poisson slope is significant while the negative binomial slope is not. Under a quadratic trend on the "
  f"active window, the dispersion parameter falls from {AL['alpha']:.2f} to {AQ['alpha']:.2f}, indicating that much "
  f"apparent overdispersion reflects misspecified mean dynamics. A product-month panel of {PN['n_products']} products with product and month fixed effects "
  f"reduces the Poisson Pearson dispersion statistic to about {PM['M1 FE only']['pearson_disp']:.0f} and shows that review volume is positively associated with a product's prior review stock "
  f"(elasticity ≈ {pm('M3 +quality','lcum'):.2f}) but not with its prior average rating. Finally, we implement an LLM coding pipeline on a random sample of 400 reviews, "
  f"evaluate it against star ratings and against a second LLM coder (κ for overall sentiment = {V['A_vs_B_kappa']['overall']:.2f}; aspect-mention κ from {min(V['A_vs_B_kappa'][k] for k in V['A_vs_B_kappa'] if k.endswith('_mention')):.2f} to {max(V['A_vs_B_kappa'][k] for k in V['A_vs_B_kappa'] if k.endswith('_mention')):.2f}), "
  f"and provide a blind human-coding protocol; human validation is not yet completed and is reported as pending. We provide code and codebook to reproduce all results.",
  indent=False, single=True)
P("Keywords: consumer-generated content, online reviews, count data, negative binomial regression, overdispersion, large language models",
  indent=False, single=True)

H1("Introduction")
P("Online reviews are a primary source of market intelligence and a common object of information systems research "
  "(Chevalier and Mayzlin 2006; Duan et al. 2008; Godes and Mayzlin 2004). Many studies analyze review volume—the number of reviews "
  "posted per product or per period—as a dependent variable or as a mediator of sales. Volume is a count, and counts of "
  "this kind are generally overdispersed relative to the Poisson benchmark, which favors negative binomial or related "
  "mixed-Poisson models (Cameron and Trivedi 2013; Hilbe 2011; Lawless 1987). Yet the diagnostics that establish "
  "whether a count model is adequate are seldom reported, and the practical consequences of getting them wrong are rarely shown.")
P("This note makes four contributions. First, using a public review corpus, we document the magnitude of overdispersion in monthly review "
  "volume and the resulting gap between Poisson and negative binomial fit. Second, we show that two routine practices—ignoring dispersion when testing a trend and imposing a log-linear trend on a rise-and-fall series—"
  "can inflate or reverse substantive conclusions about trend, and we identify a likely data-collection artifact that is easily mistaken for demand dynamics. "
  "Third, we estimate product-month panel models that separate product heterogeneity and common time shocks from within-product dynamics. Fourth, we implement an LLM-based coding pipeline for review text "
  "(Berger et al. 2020; Brown et al. 2020), report its agreement with star ratings and with a second LLM coder, and supply the blind protocol for human validation. "
  "The LLM measures are validated on a sample and are not used as covariates in the panel models; scaling them to all reviews is future work.")

H1("Data")
P(f"We use the Datafiniti Amazon product consumer reviews data, originally released by Datafiniti in May 2019 and redistributed on Zenodo "
  f"under a CC BY 4.0 license (Rajan 2026). The file contains {R['n_reviews']:,} reviews dated February 2009 to March 2019 for "
  f"{R['brand']['Amazon']['products'] + R['brand']['AmazonBasics']['products']} products. Reviews are not confined to the AmazonBasics line: "
  f"{R['brand']['Amazon']['n']:,} reviews ({R['brand']['Amazon']['n'] / R['n_reviews']:.0%}) belong to {R['brand']['Amazon']['products']} products "
  f"listed under the brand \"Amazon\" (including Fire tablets and other Amazon devices), and {R['brand']['AmazonBasics']['n']:,} reviews "
  f"({R['brand']['AmazonBasics']['n'] / R['n_reviews']:.0%}) to {R['brand']['AmazonBasics']['products']} AmazonBasics products. "
  f"We therefore treat the data as a pooled sample of Amazon-owned products rather than a single brand.")
P(f"Review dates were converted to year–month. Timestamps in the file use more than one ISO 8601 layout; default format inference either fails or, if errors are coerced, "
  f"silently discards 62 records, so we parse each record under the full ISO 8601 specification. We construct a calendar-complete "
  f"series of {R['n_months']} months, of which {R['n_zero_months']} contain no reviews and {R['n_nonempty']} contain at least one. "
  f"We retain the empty months, because a count model estimated only on non-empty months conditions on positive counts. "
  f"We also report results for the window in which collection is dense (June 2014–March 2019, 58 months, {sum(v for k, v in R['by_year'].items() if int(k) >= 2014) - 0:,} reviews counting all of 2014 onward)."
  .replace(f"{sum(v for k, v in R['by_year'].items() if int(k) >= 2014) - 0:,} reviews counting all of 2014 onward", "28,296 reviews"))
P("Review volume in this file is a measure of reviews that were collected, not necessarily of reviews written. Almost all reviews are "
  "dated 2015–2017, and counts fall to single digits from June 2018, although the file was assembled in 2019. This pattern is "
  "consistent with a crawl whose coverage changed over time. We return to this point in the Discussion.", indent=True)

P(f"Panel construction. For the panel we keep the {PN['n_products']} products with at least 50 reviews, which account for most reviews. For each product we build a month series from its first review through "
  f"March 2019 including months without reviews, and keep product-months after the product's first review (so lagged covariates exist) and in calendar months in which at least one review was posted by any product "
  f"(month effects are not identified otherwise). The panel has {PN['n_obs']:,} product-months and {PN['n_reviews']:,} reviews; {PN['zero_share']:.0%} of product-months are zero. "
  f"Because the product distribution is highly skewed (two battery products account for 43% of all reviews), standard errors are clustered by product.")
P("LLM sample. We drew a simple random sample of 400 reviews (seed 20261013) among reviews of 20 to 1,500 characters. "
  "Reviews were coded blind to star rating and product by two LLM coders of different capability tiers from the same model family "
  "(larger: coder A; smaller: coder B; exact model versions are to be inserted by the authors before submission), using the same written codebook (Appendix A). Each coder labeled overall sentiment (negative, neutral, positive), "
  "whether a product defect or failure is reported, and seven aspects (price, durability, battery, performance, ease of use, design, service) as not mentioned, negative, neutral/mixed, or positive. "
  "Coding was done in eight batches of 50 reviews per coder with the review text treated as data, and the outputs are stored with the code.")

H1("Method")
P("Let Y_t be the number of reviews dated in month t. The Poisson model specifies Y_t ~ Poisson(μ_t) with log μ_t = β_0 + β_1 t. The negative binomial (NB2) "
  "model specifies Var(Y_t) = μ_t + α μ_t², with α > 0 the dispersion parameter, and the same mean function. The Poisson model is the "
  "α = 0 special case, so we test it by a likelihood-ratio statistic whose null distribution is a 50:50 mixture of χ²(0) and χ²(1) because α lies on the boundary "
  "(Cameron and Trivedi 2013). We also report the Pearson dispersion statistic of the Poisson fit and compare models by AIC. "
  "Because the series rises and then falls, we add a quadratic term (log μ_t = β_0 + β_1 t + β_2 t²) as a robustness specification "
  "estimated on the active window. All models are estimated by maximum likelihood (code provided), with t measured in years and centered.")

P("Panel model. For product i in month t we estimate E[Y_it] = exp(α_i + γ_t + x_it'β), with product effects α_i, calendar-month effects γ_t, and covariates x_it known at t−1: the log of one plus the product's cumulative reviews "
  "(a social-proof or popularity stock) and either its cumulative average rating (centered at 4.5) or its cumulative share of reviews rated 1–2 stars. We estimate Poisson pseudo-maximum-likelihood with product-clustered standard errors, "
  "which is consistent under overdispersion if the mean is correctly specified, and an NB2 version with the dispersion parameter fixed at a moment estimate as a check. A linear product-age term is not separately identified given product and month effects and is omitted.")
P("Validation of LLM coding. We assess (i) convergent validity of the overall-sentiment label against star ratings (Spearman correlation and recall of 1–2 and 5-star reviews), "
  "(ii) inter-coder reliability between the two LLM coders by Cohen's κ for each field, and (iii) agreement with human coders, for which we provide a blind coding sheet of 150 sampled reviews and the evaluation code. "
  "Agreement between two LLMs does not establish validity because their errors may be correlated.")

H1("Results")
P(f"Overdispersion. The variance-to-mean ratio of monthly volume is {R['var_mean']:,.0f} on the full series and {R['var_mean_active']:,.0f} on the active window. "
  f"Under the linear-trend Poisson model the Pearson dispersion statistic is {FL['pearson_disp']:,.0f}, far above the value of 1 implied by the model. "
  f"The NB model fits far better (full series: AIC {FL['nb_aic']:,.0f} versus {FL['pois_aic']:,.0f}; likelihood-ratio statistic {FL['lr']:,.0f}, p < 0.001), "
  f"with α̂ = {FL['alpha']:.2f}. The Poisson assumption is clearly untenable.")
P(f"Trend. Table 1 reports the estimates. On the calendar-complete series the Poisson slope is {f(FL['pois_b_month'])} per month, whereas the NB slope is {f(FL['nb_b_month'])} per month "
  f"(SE {f(FL['nb_se_month'])}, p < 0.001). The NB slope is therefore about {FL['nb_b_month'] / FL['pois_b_month']:.1f} times larger than the Poisson slope; "
  f"allowing for overdispersion does not shrink the trend here. A plausible reason is that the NB likelihood down-weights the very large counts of 2016–2017 relative to the Poisson likelihood, so the "
  f"many near-zero early months have more influence on the slope. Neither log-linear fit describes a series that peaks in {R['peak_month'][:4]} and decays "
  f"(Figure 1); a single slope is not a meaningful summary of it.")
P(f"On the active window, the NB log-linear slope is statistically indistinguishable from zero ({f(AL['nb_b_month'])} per month, p = {AL['nb_p'][1]:.2f}) while the Poisson slope is "
  f"{f(AL['pois_b_month'])}, significant only because its standard errors ignore overdispersion. Adding a quadratic term reduces AIC from {AL['nb_aic']:,.0f} to {AQ['nb_aic']:,.0f} and "
  f"lowers the dispersion parameter from {AL['alpha']:.2f} to {AQ['alpha']:.2f}; the quadratic coefficient is strongly negative ({f(AQ['nb_b'][2])}, p < 0.001), "
  f"implying a peak in late 2016. Remaining overdispersion is concentrated in a few months: volume reached {R['peak']:,} in January 2017, "
  f"far above the fitted curve (Figure 1).")
table([["", "Poisson", "Negative binomial", "Poisson", "Negative binomial"],
       ["Window", "Full (122 mo.)", "Full (122 mo.)", "Active (58 mo.)", "Active (58 mo.)"],
       ["Trend, log-linear (per month)", f(FL["pois_b_month"]), f(FL["nb_b_month"]), f(AL["pois_b_month"]), f(AL["nb_b_month"])],
       ["  Standard error", "—", f(FL["nb_se_month"]), "—", f(AL["nb_se_month"])],
       ["Dispersion α", "—", f(FL["alpha"], 2), "—", f(AL["alpha"], 2)],
       ["AIC", f"{FL['pois_aic']:,.0f}", f"{FL['nb_aic']:,.0f}", f"{AL['pois_aic']:,.0f}", f"{AL['nb_aic']:,.0f}"],
       ["Quadratic NB: α / AIC", "", "", "", f"{AQ['alpha']:.2f} / {AQ['nb_aic']:,.0f}"]],
      note="Table 1. Poisson and negative binomial trend models for monthly review volume. Poisson standard errors are omitted because they assume equidispersion and are misleading here.")
doc.add_picture("figure1.png", width=Inches(6.0)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
P("Figure 1. Monthly review volume (n = 28,332) and negative binomial fit with a quadratic trend on the active window.", italic=True, indent=False, single=True)
P("Text covariates. As a descriptive check on the usefulness of text-derived covariates, we computed dictionary-based rates of positive and negative terms per 1,000 words by month "
  "(word lists in analysis.py; months with at least one review in the active window). Their correlations with monthly volume are weak (positive-term rate r = 0.23; negative-term rate r = 0.09), and an NB model with both rates and a quadratic trend "
  "did not converge to a stable solution, so we report no coefficients. We read this as evidence that dictionary measures add little here, not as evidence about LLM measures.")

H2("Product-month panel")
P(f"Table 2 reports the panel estimates. With only product and month effects, the Poisson Pearson dispersion statistic is {PM['M1 FE only']['pearson_disp']:.1f}, compared with {FL['pearson_disp']:,.0f} for the aggregate linear-trend model. "
  f"Dispersion remains well above one, so cluster-robust or negative binomial inference is still needed, but most of the aggregate overdispersion is attributable to differences across products and months. "
  f"The cumulative review stock is positively associated with next-month volume: the elasticity is {pm('M2 +social proof','lcum'):.2f} (SE {pm('M2 +social proof','lcum','se'):.2f}) without a rating control and "
  f"{pm('M3 +quality','lcum'):.2f} (SE {pm('M3 +quality','lcum','se'):.2f}) with it, and {pm('M3-NB2','lcum'):.2f} (SE {pm('M3-NB2','lcum','se'):.2f}) in the NB2 check. "
  f"The prior average rating ({pm('M3 +quality','avg_prev_c'):.2f}, p = {pm('M3 +quality','avg_prev_c','p'):.2f}) and prior negative-review share ({pm('M4 +neg share','negshare_prev'):.2f}, p = {pm('M4 +neg share','negshare_prev','p'):.2f}) are not significantly related to volume. "
  f"Dropping product effects raises the stock elasticity to {pm('M5 between (no product FE)','lcum'):.2f} and gives no brand-group difference (devices versus AmazonBasics: {pm('M5 between (no product FE)','device'):.2f}, p = {pm('M5 between (no product FE)','device','p'):.2f}). "
  f"These are associations: the stock of reviews rises mechanically with a product's age and popularity, so we do not interpret the elasticity as a causal social-proof effect.")
def c(m, v): return f"{pm(m,v):.3f} ({pm(m,v,'se'):.3f})"
table([["", "(1) FE only", "(2) + stock", "(3) + avg. rating", "(4) + neg. share", "(5) NB2"],
       ["log(1 + prior reviews)", "", c("M2 +social proof", "lcum"), c("M3 +quality", "lcum"), c("M4 +neg share", "lcum"), c("M3-NB2", "lcum")],
       ["Prior avg. rating (− 4.5)", "", "", c("M3 +quality", "avg_prev_c"), "", c("M3-NB2", "avg_prev_c")],
       ["Prior neg. share (1–2★)", "", "", "", c("M4 +neg share", "negshare_prev"), ""],
       ["Product / month effects", "Yes / Yes", "Yes / Yes", "Yes / Yes", "Yes / Yes", "Yes / Yes"],
       ["Pearson dispersion (Poisson)", f"{PM['M1 FE only']['pearson_disp']:.1f}", f"{PM['M2 +social proof']['pearson_disp']:.1f}", f"{PM['M3 +quality']['pearson_disp']:.1f}", f"{PM['M4 +neg share']['pearson_disp']:.1f}", "—"]],
      note=f"Table 2. Product-month panel ({PN['n_products']} products, {PN['n_obs']:,} product-months). Coefficients with product-clustered standard errors in parentheses; columns 1–4 are Poisson PML, column 5 is NB2 with α fixed at {PM['M3-NB2']['alpha_used']:.2f} (moment estimate), so its standard errors understate uncertainty about α.")
H2("LLM coding: convergent validity and reliability")
sa, sb = V["A_vs_stars"], V["B_vs_stars"]; kk = V["A_vs_B_kappa"]
P(f"Both coders labeled all 400 reviews. Coder A labeled {sa['pos_share']:.0%} of reviews positive and {sa['neg_share']:.0%} negative (coder B: {sb['pos_share']:.0%} and {sb['neg_share']:.0%}); "
  f"the sample is dominated by 5-star reviews ({sa['n_five']} of 400) and has only {sa['n_low_star']} reviews rated 1–2 stars. "
  f"Overall sentiment correlates with star rating at Spearman ρ = {sa['spearman']:.2f} (coder A) and {sb['spearman']:.2f} (coder B). "
  f"Coder A labeled {sa['pos_recall_five']:.0%} of 5-star reviews positive and {sa['neg_recall_low_star']:.0%} of 1–2-star reviews negative; of the reviews it labeled negative, {sa['neg_precision']:.0%} had 1–2 stars. A positive/not-positive accuracy of {V['A_vs_stars']['acc_pos_neg']:.0%} is below the {sa['majority_baseline_acc']:.0%} accuracy "
  f"of always predicting positive, so we do not report accuracy as evidence of validity.")
P(f"Inter-coder reliability is high for overall sentiment (κ = {kk['overall']:.2f}), price (κ = {kk['price']:.2f}), service (κ = {kk['service']:.2f}), design (κ = {kk['design']:.2f}), ease of use (κ = {kk['ease']:.2f}), and defect reports (κ = {kk['defect']:.2f}), "
  f"and lower for durability (κ = {kk['durability']:.2f}), battery (κ = {kk['battery']:.2f}), and especially performance (κ = {kk['performance']:.2f}), which is also the most prevalent aspect ({V['prevalence_A']['performance']:.0%} of reviews) and the broadest category in the codebook. "
  f"Aspect prevalence under coder A ranges from {V['prevalence_A']['service']:.0%} (service) to {V['prevalence_A']['performance']:.0%} (performance), and {V['defect_A']:.0%} of reviews report a defect. "
  f"Human validation is pending: the blind sheet (human_coding_sheet.xlsx) and evaluate.py are provided, and the human-agreement columns are to be completed before submission.")
table([["Field", "κ (A vs. B)", "Raw agreement", "Prevalence (A)"]] +
      [[n, f"{kk[k]:.2f}", f"{V['A_vs_B_agree'][k]:.0%}", ("—" if k in ("overall",) else (f"{V['defect_A']:.0%}" if k == "defect" else f"{V['prevalence_A'][k]:.0%}"))]
       for n, k in (("Overall sentiment", "overall"), ("Defect reported", "defect"), ("Price", "price"), ("Durability", "durability"), ("Battery", "battery"),
                    ("Performance", "performance"), ("Ease of use", "ease"), ("Design", "design"), ("Service", "service"))],
      note="Table 3. Agreement between the two LLM coders on 400 reviews. Aspect κ is computed on the four-category code. Human agreement: pending.")

H1("Discussion")
H2("Implications for research")
P("First, overdispersion testing should be reported, not assumed. In this corpus the Poisson model is rejected by every criterion, and its reported significance "
  "is an artifact of ignoring dispersion. Second, the NB correction does not by itself make a trend estimate interpretable. A log-linear trend imposed on a boom-and-bust series yields a statistically significant but substantively "
  "meaningless slope, and the size of the estimated dispersion depends heavily on how well the mean is specified (α̂ falls from 2.45 to 0.42 once curvature is allowed). "
  "Researchers should therefore model the mean dynamics, or use product fixed effects and release-time controls in panel data, before interpreting α. "
  "Third, sample construction choices—dropping empty months, silent timestamp parsing failures, and pooling distinct brands—change results in ways that are rarely documented.")
P("Fourth, aggregate overdispersion is partly heterogeneity. Product and month effects reduce the Poisson dispersion statistic from the hundreds to the teens, so research designs with product-level panels should be preferred to aggregate series, "
  "and review stock should be treated as a state variable rather than an exogenous trend.")
H2("Implications for LLM-based measurement")
P("In this sample, two LLM coders agree closely on overall sentiment, price, service, design, ease of use, and defect reports, and agree less on durability, battery, and performance. "
  "Researchers should therefore (i) document the codebook, model versions, and prompts, (ii) report reliability per aspect rather than a single accuracy figure, and (iii) treat low-reliability aspects as noisy measures when they enter count models, "
  "since measurement error in covariates attenuates their coefficients. Star ratings are an imperfect benchmark, and about a third of the reviews that coder A labeled negative carry 3–5 stars; whether these are coder errors or genuine complaints in well-rated reviews can only be settled by human coding. Because the LLM coding here uses a sample, the LLM-derived variables are not entered into the panel models; "
  "doing so requires coding all reviews and propagating measurement error, which we leave to future work.")
H2("Limitations")
P("The data come from one retailer, are pooled across Amazon-owned products, and appear to reflect a time-varying crawl. The collapse in counts after 2017 "
  "is unlikely to reflect a real decline in reviewing, so the trends we estimate describe the collected sample and not consumer behavior. The panel covers 28 products, is dominated by a few, and its estimates are associations. "
  "The LLM validation uses a 400-review random sample dominated by 5-star reviews, so negative-class statistics rest on few observations; both coders are LLMs whose errors may be correlated; "
  "and agreement with human coders has not yet been measured. Coding was performed with specific model versions that may not be available in the future, so we archive all outputs.")
H1("Conclusion")
P("Review volume is an overdispersed count series for which Poisson models are inadequate. Negative binomial models are the appropriate starting point, but our results show that "
  "their value depends on a correctly specified mean function and a carefully constructed sample. Product-month panels explain much of the aggregate dispersion, and LLM coders can be made reliable for several aspects but not all. We encourage future work to report dispersion diagnostics, "
  "document sample construction (including empty periods and data-collection coverage), and validate LLM-derived measures against human coders before using them as covariates.")

H1("Data and Code Availability")
P("The data are available at https://doi.org/10.5281/zenodo.18758751. The scripts analysis.py (aggregate models and Figure 1), panel.py (panel models), and llm/ (sample, codebook, batches, coder outputs, blind human coding sheet, and evaluate.py) reproduce every number.", indent=False)

H1("Appendix A. Coding codebook")
for line in open("llm/codebook.md").read().splitlines()[2:]:
    if line.strip(): P(line.lstrip("- ").replace("`", ""), indent=False, single=True)

H1("References")
refs = [
 "Berger, J., Humphreys, A., Ludwig, S., Moe, W. W., Netzer, O., and Schweidel, D. A. 2020. “Uniting the Tribes: Using Text for Marketing Insight,” Journal of Marketing (84:1), pp. 1–25.",
 "Brown, T. B., et al. 2020. “Language Models Are Few-Shot Learners,” in Advances in Neural Information Processing Systems (33), pp. 1877–1901.",
 "Cameron, A. C., and Trivedi, P. K. 2013. Regression Analysis of Count Data (2nd ed.), Cambridge, UK: Cambridge University Press.",
 "Chevalier, J. A., and Mayzlin, D. 2006. “The Effect of Word of Mouth on Sales: Online Book Reviews,” Journal of Marketing Research (43:3), pp. 345–354.",
 "Duan, W., Gu, B., and Whinston, A. B. 2008. “Do Online Reviews Matter? An Empirical Investigation of Panel Data,” Decision Support Systems (45:4), pp. 1007–1016.",
 "Godes, D., and Mayzlin, D. 2004. “Using Online Conversations to Study Word-of-Mouth Communication,” Marketing Science (23:4), pp. 545–560.",
 "Hilbe, J. M. 2011. Negative Binomial Regression (2nd ed.), Cambridge, UK: Cambridge University Press.",
 "Lawless, J. F. 1987. “Negative Binomial and Mixed Poisson Regression,” Canadian Journal of Statistics (15:3), pp. 209–225.",
 "Rajan G, S. 2026. “Datafiniti Amazon Product Consumer Reviews Dataset (May 2019)” [Data set], Zenodo, https://doi.org/10.5281/zenodo.18758751 (original data released by Datafiniti, 2019).",
]
for r in refs:
    p = P(r, indent=False); p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5)

# footer page numbers
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
for typ, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if typ: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), typ)
    else: e = OxmlElement("w:instrText"); e.text = txt
    r._r.append(e)
doc.save("MISQ_Research_Note_CGC_Count_Models_revised.docx")
