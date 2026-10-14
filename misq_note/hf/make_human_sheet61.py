import random, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
K=pd.read_csv('qual/wire_cases_KEY.csv',dtype={'id':str}); a=pd.read_csv('raw/truth_archive.csv',dtype={'id':str}).set_index('id').content.fillna('')
ids=list(K.id); random.seed(61); order=ids[:]; random.shuffle(order)
key=pd.DataFrame({'item':range(1,len(order)+1),'id':order}).merge(K[['id','wid','cls']].rename(columns={'cls':'model1'}),on='id'); key.to_csv('human_coding_key_61.csv',index=False)
code="""CODEBOOK v2 (read before coding; code each post from its text only)
Each item is a Truth Social post by Donald Trump (2026) about Iran, the Strait of Hormuz, oil, or related negotiations. Choose exactly ONE class. Use only the text shown: do not look anything up, do not look at prices, work alone, do not discuss items with the other coder.

ESC   = the post RAISES or KEEPS UP the risk of military conflict or oil-supply disruption. This includes: threats and warnings; announcements or celebration of attacks, strikes, blockade, seizures or destruction; ultimatums; DEMANDS or CONDITIONS addressed to Iran that are backed by pressure (e.g. 'Iran must..., or ...'), and criticism that the other side is failing to agree, when no pause, concession or progress is announced.
DEESC = the post LOWERS the risk: it ANNOUNCES or clearly signals a pause, postponement or cancellation of strikes, a ceasefire, an agreement or a deal, progress in talks, the reopening of the Strait of Hormuz or lifting of the blockade, sanctions relief, or other relief. A post only describing terms or demands, without announcing a pause, agreement or progress, is NOT DEESC.
MIXED = a clear announcement of de-escalation (pause, progress, deal) AND a clear threat or escalation in the same post (e.g. 'we will hold off, but if they do not sign we will strike').
OTHER = not about the military or diplomatic course of the conflict or about oil supply: domestic politics, attacks on the media or critics without a stance on the conflict, boasts without a stance on what happens next, polls, unrelated topics. Re-posts and links ('RT: https://...' or a bare link) with no stated stance in the visible text: OTHER. If the visible text quotes a headline with a clear stance, code the stance of that headline.

Confidence: 1 = clear, 2 = somewhat unclear.
Column E is for optional notes (e.g., why the item was hard)."""
for coder in 'AB':
    wb=Workbook(); ws=wb.active; ws.title='Codebook'; ws['A1']=code; ws['A1'].alignment=Alignment(wrap_text=True,vertical='top'); ws.column_dimensions['A'].width=130; ws.row_dimensions[1].height=330
    s=wb.create_sheet('Coding'); s.append(['item','post text','class (ESC/DEESC/MIXED/OTHER)','confidence (1/2)','notes'])
    for c in s[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='DDDDDD')
    rows=list(zip(key.item,key.id))
    if coder=='B': random.Random(6161).shuffle(rows)
    for n,i in rows: s.append([int(n),a[i],None,None,None])
    for r in range(2,len(rows)+2): s.cell(r,2).alignment=Alignment(wrap_text=True,vertical='top')
    s.column_dimensions['A'].width=7; s.column_dimensions['B'].width=110; s.column_dimensions['C'].width=22; s.column_dimensions['D'].width=14; s.column_dimensions['E'].width=40; s.freeze_panes='A2'
    dv=DataValidation(type='list',formula1='"ESC,DEESC,MIXED,OTHER"',allow_blank=True); s.add_data_validation(dv); dv.add(f'C2:C{len(rows)+1}')
    dv2=DataValidation(type='list',formula1='"1,2"',allow_blank=True); s.add_data_validation(dv2); dv2.add(f'D2:D{len(rows)+1}')
    wb.save(f'human_coding_sheet61_coder{coder}.xlsx')
print(len(key),key.model1.value_counts().to_dict())
