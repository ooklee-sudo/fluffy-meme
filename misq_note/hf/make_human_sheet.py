import json,random,numpy as np,pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
random.seed(20261014)
allp=json.load(open("posts_all_for_coder2.json")); txt={x["id"]:x["text"] for x in allp}
samp={}
for x in json.load(open("posts_coded.json")): samp[x["id"]]="W"
for x in json.load(open("posts_ext_to_code.json")): samp[x["id"]]=x["sample"]
c1={x["id"]:x["cls"] for x in json.load(open("posts_coded.json"))}; c1.update({x["id"]:x["cls"] for x in json.load(open("posts_ext_coder1.json"))})
c2={x["id"]:x["cls"] for x in json.load(open("posts_all_coder2.json"))}
ids=list(txt); by={s:[i for i in ids if samp[i]==s] for s in["W","A","B"]}
quota={"W":67,"A":23,"B":60}; pick=[]
for s,q in quota.items(): pick+=random.sample(by[s],q)
random.shuffle(pick)
key=pd.DataFrame({"item":range(1,len(pick)+1),"id":pick,"sample":[samp[i] for i in pick],"model1":[c1[i] for i in pick],"model2":[c2[i] for i in pick]}); key.to_csv("human_coding_key.csv",index=False)
code="""CODEBOOK (read before coding; code each post from its text only)
Each item is a Truth Social post by Donald Trump (2024-2026) that matched keywords about Iran, oil, or oil-producing countries. Choose exactly ONE class:

ESC   = the post RAISES the risk of military conflict or of oil-supply disruption: threatens, warns of, announces or celebrates attacks/strikes/blockade/destruction (e.g. on Iran), ultimatums, new sanctions or secondary tariffs on oil buyers or oil-exporting countries, threats against such countries or their energy infrastructure.
DEESC = the post LOWERS that risk: announces or hints at holding off/cancelling/postponing strikes, ceasefire, deal/memorandum/agreement progress, talks, reopening of routes (e.g. Strait of Hormuz), peace, sanctions relief, higher oil production or lower oil prices through supply.
MIXED = clear escalation AND clear de-escalation in the same post (e.g. 'deal or we strike').
OTHER = not about the military/diplomatic course of conflicts involving oil-producing countries or oil supply (polls, domestic politics, media attacks, tariffs on non-oil goods, unrelated mentions, retweets/links of others' commentary without a clear stance).

Confidence: 1 = clear, 2 = somewhat unclear.
Rules: (1) use only the text shown; do not look anything up and do not look at market prices; (2) code the post's own stance (for a retweet or link, the stance of the quoted headline if clear, otherwise OTHER); (3) work alone, without discussing items with the other coder; (4) if the text is cut off at the end, code what you see.
Fill column C (class) and D (confidence); column E is for optional notes."""
for coder in["A","B"]:
    wb=Workbook(); ws=wb.active; ws.title="Codebook"; ws["A1"]=code; ws["A1"].alignment=Alignment(wrap_text=True,vertical="top"); ws.column_dimensions["A"].width=130; ws.row_dimensions[1].height=330
    s=wb.create_sheet("Coding"); s.append(["item","post text","class (ESC/DEESC/MIXED/OTHER)","confidence (1/2)","notes"])
    for c in s[1]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="DDDDDD")
    rows=list(zip(range(1,len(pick)+1),pick))
    if coder=="B": random.Random(7).shuffle(rows)       # coder B sees a different order
    for n,i in rows: s.append([n,txt[i],None,None,None])
    for r in range(2,len(rows)+2): s.cell(r,2).alignment=Alignment(wrap_text=True,vertical="top")
    s.column_dimensions["A"].width=7; s.column_dimensions["B"].width=110; s.column_dimensions["C"].width=22; s.column_dimensions["D"].width=14; s.column_dimensions["E"].width=40; s.freeze_panes="A2"
    dv=DataValidation(type="list",formula1='"ESC,DEESC,MIXED,OTHER"',allow_blank=True); s.add_data_validation(dv); dv.add(f"C2:C{len(rows)+1}")
    dv2=DataValidation(type="list",formula1='"1,2"',allow_blank=True); s.add_data_validation(dv2); dv2.add(f"D2:D{len(rows)+1}")
    wb.save(f"human_coding_sheet_coder{coder}.xlsx")
print(len(pick),key["sample"].value_counts().to_dict(),key.model1.value_counts().to_dict(),key.model2.value_counts().to_dict())
