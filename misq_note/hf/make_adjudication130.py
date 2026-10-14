import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
d=pd.read_csv('human130_disagreements.csv',dtype={'id':str}); a=pd.read_csv('raw/truth_archive.csv',dtype={'id':str}).set_index('id').content.fillna('')
wb=Workbook(); ws=wb.active; ws.title='Adjudication'
ws.append(['item','post text','class (ESC/DEESC/MIXED/OTHER)','notes'])
for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='DDDDDD')
for r in d.itertuples(): ws.append([int(r.item),a[r.id],None,None])
for r in range(2,len(d)+2): ws.cell(r,2).alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=7; ws.column_dimensions['B'].width=110; ws.column_dimensions['C'].width=22; ws.column_dimensions['D'].width=40; ws.freeze_panes='A2'
dv=DataValidation(type='list',formula1='"ESC,DEESC,MIXED,OTHER"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'C2:C{len(d)+1}')
ws2=wb.create_sheet('Codebook v2'); ws2['A1']="Use codebook v2 (same as the coders' sheet). ESC = raises/keeps up risk (threats, attacks, ultimatums, pressure demands, criticism of failure to agree without announced progress). DEESC = announces/clearly signals pause, ceasefire, agreement, progress, reopening, relief. MIXED = both clearly. OTHER = not about the course of the conflict or oil supply; link/re-post with no stated stance. Code from the text only; do not look anything up and do not look at prices."; ws2['A1'].alignment=Alignment(wrap_text=True,vertical='top'); ws2.column_dimensions['A'].width=120; ws2.row_dimensions[1].height=140
wb.save('adjudication_sheet_130.xlsx'); print(len(d))
