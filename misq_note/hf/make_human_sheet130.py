import random, json, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
pc=pd.DataFrame(json.load(open('posts_coded.json'))); K=pd.read_csv('qual/wire_cases_KEY.csv',dtype={'id':str}); a=pd.read_csv('raw/truth_archive.csv',dtype={'id':str}).set_index('id').content.fillna('')
rest=pc[~pc.id.isin(K.id)].copy(); ids=list(rest.id); random.seed(130); random.shuffle(ids)
key=pd.DataFrame({'item':range(1,len(ids)+1),'id':ids}).merge(rest[['id','cls']].rename(columns={'cls':'model1'}),on='id'); key.to_csv('human_coding_key_130.csv',index=False)
code=open('make_human_sheet61.py').read().split('code="""')[1].split('"""')[0]
for coder in 'AB':
    wb=Workbook(); ws=wb.active; ws.title='Codebook'; ws['A1']=code; ws['A1'].alignment=Alignment(wrap_text=True,vertical='top'); ws.column_dimensions['A'].width=130; ws.row_dimensions[1].height=330
    s=wb.create_sheet('Coding'); s.append(['item','post text','class (ESC/DEESC/MIXED/OTHER)','confidence (1/2)','notes'])
    for c in s[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='DDDDDD')
    rows=list(zip(key.item,key.id))
    if coder=='B': random.Random(1313).shuffle(rows)
    for n,i in rows: s.append([int(n),a[i],None,None,None])
    for r in range(2,len(rows)+2): s.cell(r,2).alignment=Alignment(wrap_text=True,vertical='top')
    s.column_dimensions['A'].width=7; s.column_dimensions['B'].width=110; s.column_dimensions['C'].width=22; s.column_dimensions['D'].width=14; s.column_dimensions['E'].width=40; s.freeze_panes='A2'
    dv=DataValidation(type='list',formula1='"ESC,DEESC,MIXED,OTHER"',allow_blank=True); s.add_data_validation(dv); dv.add(f'C2:C{len(rows)+1}')
    dv2=DataValidation(type='list',formula1='"1,2"',allow_blank=True); s.add_data_validation(dv2); dv2.add(f'D2:D{len(rows)+1}')
    wb.save(f'human_coding_sheet130_coder{coder}.xlsx')
print(len(key),key.model1.value_counts().to_dict())
