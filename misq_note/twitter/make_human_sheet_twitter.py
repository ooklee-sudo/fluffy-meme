import random, json, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
F=pd.read_csv('frame_features.csv',dtype={'id':str}); T=pd.read_csv('frame_oil_iran.csv',dtype={'id':str}).set_index('id').text
a=json.load(open('labels_coder1.json')); b=json.load(open('labels_coder2.json')); F['m1']=F.id.map(lambda i:a[i]['cls']); F['m2']=F.id.map(lambda i:b[i]['cls'])
K=F[F.keep].copy(); sig=K[(K.m1!='OTHER')|(K.m2!='OTHER')]; oth=K[(K.m1=='OTHER')&(K.m2=='OTHER')]
random.seed(2017); pick=list(sig.id)+random.sample(list(oth.id),45); random.shuffle(pick)
key=pd.DataFrame({'item':range(1,len(pick)+1),'id':pick}).merge(F[['id','m1','m2']],on='id'); key.to_csv('human_key_twitter.csv',index=False)
print(len(K),len(sig),len(oth),len(pick))
code="""CODEBOOK v2 (read before coding; code each tweet from its text only)
Each item is a tweet by Donald Trump (2017-2021) that mentioned Iran, oil, OPEC, Saudi Arabia, energy or related words. Choose exactly ONE class. Use only the text shown: do not look anything up, do not look at prices, work alone, do not discuss items with the other coder.

ESC   = the tweet RAISES or KEEPS UP the risk of military conflict or oil-supply disruption. This includes: threats and warnings; announcements or celebration of attacks, strikes, blockade, seizures or destruction; ultimatums; DEMANDS or CONDITIONS addressed to Iran backed by pressure; criticism that the other side is failing to agree, when no pause, concession or progress is announced; threatening, announcing or celebrating sanctions on oil exporters.
DEESC = the tweet LOWERS the risk: it ANNOUNCES or clearly signals a pause, postponement or cancellation of strikes, a ceasefire, an agreement or a deal, progress in talks, the reopening of routes, sanctions relief, waivers or exemptions for oil supply, or urges OPEC/Saudi Arabia to produce more or oil prices to fall. A tweet only describing terms or demands, without announcing a pause, agreement or progress, is NOT DEESC.
MIXED = a clear announcement of de-escalation (pause, progress, deal) AND a clear threat or escalation in the same tweet.
OTHER = not about the military or diplomatic course of a conflict or about oil-supply risk: energy or pipeline policy without bearing on conflict or supply risk, domestic politics, attacks on the media or critics without a stance on the conflict, boasts without a stance on what happens next, unrelated topics. A tweet that is only a link or has no stated stance: OTHER.

Confidence: 1 = clear, 2 = somewhat unclear. Column E is for optional notes."""
for coder in 'AB':
    wb=Workbook(); ws=wb.active; ws.title='Codebook'; ws['A1']=code; ws['A1'].alignment=Alignment(wrap_text=True,vertical='top'); ws.column_dimensions['A'].width=130; ws.row_dimensions[1].height=300
    s=wb.create_sheet('Coding'); s.append(['item','tweet text','class (ESC/DEESC/MIXED/OTHER)','confidence (1/2)','notes'])
    for c in s[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='DDDDDD')
    rows=list(zip(key.item,key.id))
    if coder=='B': random.Random(1717).shuffle(rows)
    for n,i in rows: s.append([int(n),T[i],None,None,None])
    for r in range(2,len(rows)+2): s.cell(r,2).alignment=Alignment(wrap_text=True,vertical='top')
    s.column_dimensions['A'].width=7; s.column_dimensions['B'].width=110; s.column_dimensions['C'].width=22; s.column_dimensions['D'].width=14; s.column_dimensions['E'].width=40; s.freeze_panes='A2'
    dv=DataValidation(type='list',formula1='"ESC,DEESC,MIXED,OTHER"',allow_blank=True); s.add_data_validation(dv); dv.add(f'C2:C{len(rows)+1}')
    dv2=DataValidation(type='list',formula1='"1,2"',allow_blank=True); s.add_data_validation(dv2); dv2.add(f'D2:D{len(rows)+1}')
    wb.save(f'human_coding_sheet_twitter_coder{coder}.xlsx')
