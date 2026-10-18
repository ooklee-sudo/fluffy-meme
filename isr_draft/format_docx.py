"""Post-process the pandoc output into a manuscript-style Word file (Times New Roman 12, double spaced)."""
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, dst = sys.argv[1], sys.argv[2]
d = Document(src)

for sec in d.sections:
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)

def set_font(style, size, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    f = style.font
    f.name = "Times New Roman"; f.size = Pt(size); f.color.rgb = color
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Times New Roman")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]

ST = {x.name.lower(): x for x in d.styles}
for name in ("normal", "body text", "first paragraph", "compact", "block text"):
    if name in ST:
        st = ST[name]; set_font(st, 12)
        pf = st.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE; pf.space_after = Pt(0); pf.space_before = Pt(0)
for name, size in (("title", 16), ("heading 1", 14), ("heading 2", 12), ("heading 3", 12)):
    if name in ST:
        st = ST[name]; set_font(st, size, bold=True, italic=(name == "heading 3"))
        st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        st.paragraph_format.keep_with_next = True
for name in ("image caption", "table caption", "caption"):
    if name in ST:
        st = ST[name]; set_font(st, 10)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

for p in d.paragraphs:
    if p.style.name.lower() == "title":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if p.style.name.lower() in ("body text", "first paragraph") and not p.text.startswith("Working draft"):
        p.paragraph_format.first_line_indent = Inches(0.5)
    # references: single spaced, hanging indent
    if p.text and p.runs and p.paragraph_format is not None:
        pass
    for r in p.runs:
        r.font.name = "Times New Roman"
        rpr = r._element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), "Times New Roman")

# reference list formatting
in_refs = False
for p in d.paragraphs:
    if p.style.name.lower().startswith("heading") and p.text.strip() == "References":
        in_refs = True; continue
    if in_refs and p.text.strip():
        p.paragraph_format.first_line_indent = Inches(-0.5); p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE; p.paragraph_format.space_after = Pt(6)

# tables: borders, small font, single spacing
for t in d.tables:
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for b in ("top", "bottom", "insideH"):
        e = OxmlElement(f"w:{b}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)
    t.autofit = False
    ncol = len(t.columns)
    widths = [Inches(1.6)] + [Inches(4.85 / max(ncol - 1, 1))] * (ncol - 1)
    if ncol == 10:
        widths = [Inches(x) for x in (1.45, 0.5, 0.5, 0.5, 0.68, 0.6, 0.5, 0.68, 0.55, 0.5)]
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = w
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.inches * 1440)))
    tw = tblPr.find(qn("w:tblW"))
    if tw is not None:
        tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(6.45 * 1440)))
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for ri, row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(8); r.font.name = "Times New Roman"; r.font.bold = (ri == 0)

# figures: fit to page width
for shp in d.inline_shapes:
    ratio = shp.height / shp.width
    shp.width = Inches(6.4); shp.height = int(Inches(6.4) * ratio)

# page numbers
for sec in d.sections:
    p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text
        run._r.append(e)
    run.font.name = "Times New Roman"; run.font.size = Pt(10)
d.save(dst)
