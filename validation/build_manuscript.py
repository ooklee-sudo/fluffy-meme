import copy, re, docx
from docx.shared import Pt, Inches, Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
SRC='/root/.claude/uploads/85836436-e025-583a-b00a-2fcdeefc7114/d7e2a7ed-Poisson-SNS_IEEE_manuscript.docx'
OUT='/home/user/fluffy-meme/Poisson-SNS_IEEE_Access_manuscript.docx'
d=docx.Document(SRC)
ps=d.paragraphs

# ---- text edits for IEEE Access ----
def set_text(p,t):
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    p.runs[0].text=t
# AI-use disclosure goes under "Acknowledgment" (IEEE Access requirement)
h=[p for p in ps if p.text=='Declaration of Generative AI Use'][0]
set_text(h,'Acknowledgment')
b=h._p.getnext()
bp=[p for p in ps if p._p is b][0]
set_text(bp,'Use of generative AI: '+bp.text)
# References: IEEE style fixes
fix=[('annu. ACM symp. User interface softw. technol.','Annu. ACM Symp. User Interface Softw. Technol.'),
     ('New York, NY, USA: ACM, 2023.','2023, Art. no. 2.'),
     ('int. Conf. Knowl. Discovery data mining','Int. Conf. Knowl. Discovery Data Mining'),
     ('int. Conf. Web search data mining','Int. Conf. Web Search Data Mining'),
     ('S​3: Social','S³: Social'),
     ('“S: Social','“S³: Social')]
for p in ps:
    if p.style.name=='Bibliography':
        for r in p.runs:
            for a,c in fix:
                if a in r.text: r.text=r.text.replace(a,c)
        t=p.text
        m=re.match(r'\[(\d+)\]',t)
def para_sub(p,pat,rep):
    full=p.text; new=re.sub(pat,rep,full)
    if new!=full:
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
        p.runs[0].text=new
for p in ps:
    if p.style.name=='Bibliography': para_sub(p,r'\.” ?arXiv:(\d+\.\d+), (\d{4})\.',r',” \2, arXiv:\1.')
# arXiv refs: IEEE puts the identifier after the title with year
for p in ps:
    if p.style.name=='Bibliography':
        for r in p.runs:
            r.text=re.sub(r'\.” ?arXiv:(\d+\.\d+), (\d{4})\.',r',” \2, arXiv:\1.',r.text)
        # runs may split; handle at paragraph-level if needed
for p in ps:
    if p.style.name=='Bibliography' and 'arXiv' in p.text and '.” arXiv' in p.text:
        print('UNFIXED',p.text[:60])

import sys; sys.path.insert(0,'/home/user/fluffy-meme/validation')
from study4_text import add_study4
add_study4(d, ps)
ps=d.paragraphs[:0] or ps
# ---- formatting: IEEE Access look (Times 10 pt, two columns) ----
st=d.styles
def sty(n,size=None,after=None,before=None,line=None,ind=None,first=None,bold=None):
    s=[x for x in st if x.name==n][0]
    if size: s.font.size=Pt(size)
    pf=s.paragraph_format
    if ind is not None: pf.left_indent=Inches(ind)
    if after is not None: pf.space_after=Pt(after)
    if before is not None: pf.space_before=Pt(before)
    if line: pf.line_spacing=line
    if first is not None: pf.first_line_indent=Inches(first)
    if bold is not None: s.font.bold=bold
sty('Title',24,after=8)
sty('Author',11,after=8)
sty('Abstract',9,after=4,bold=True)
for n in ('Body Text','First Paragraph'): sty(n,10,after=0,line=1.0,first=0.17)
sty('Compact',7.5,after=0,line=1.0,first=0)
sty('Heading 1',10,after=4,before=10)
sty('Heading 2',10,after=3,before=6)
sty('Table Caption',8,after=3,before=4)
sty('Image Caption',8,after=4,before=2)
sty('Bibliography',8,after=1,ind=0.33,first=-0.33)
# abstract + index terms: 9 pt bold-style per IEEE
for p in ps[2:4]:
    for r in p.runs: r.font.size=Pt(9)
    p.paragraph_format.first_line_indent=Inches(0)
    p.paragraph_format.space_after=Pt(4)
for r in ps[2].runs[1:]: r.font.bold=True
for r in ps[3].runs: r.font.bold=True
# paragraphs that are not body prose: no indent
for p in ps:
    if p.style.name in ('Heading 1','Heading 2','Bibliography','Image Caption','Table Caption','Captioned Figure'):
        if p.style.name!='Bibliography': p.paragraph_format.first_line_indent=Inches(0)
    if p.style.name=='First Paragraph' and p._p.getprevious() is not None and p._p.getprevious().tag==qn('w:tbl'):
        pass
# tables
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Inches(0)
                p.paragraph_format.space_after=Pt(0)
                for r in p.runs: r.font.size=Pt(7.5)
    tp=t._tbl.tblPr
    cm=OxmlElement('w:tblCellMar')
    for side in ('left','right'):
        e=OxmlElement('w:'+side); e.set(qn('w:w'),'40'); e.set(qn('w:type'),'dxa'); cm.append(e)
    tp.append(cm)
# note paragraphs under tables ("Note.") smaller
for p in ps:
    if p.text.startswith('Note.') or p.text.startswith('Note '):
        p.paragraph_format.first_line_indent=Inches(0)
        for r in p.runs: r.font.size=Pt(7.5)
        p.paragraph_format.space_after=Pt(4)
# figures to column width (3.45 in)
W=int(3.45*914400)
for ext in d.element.body.iter(qn('wp:extent')):
    cx,cy=int(ext.get('cx')),int(ext.get('cy')); nh=int(cy*W/cx)
    ext.set('cx',str(W)); ext.set('cy',str(nh))
for e in d.element.body.iter(qn('a:ext')):
    if e.get('cx'):
        cx,cy=int(e.get('cx')),int(e.get('cy')); e.set('cx',str(W)); e.set('cy',str(int(cy*W/cx)))

# ---- sections: single-column front matter, two-column body ----
def setpage(sp):
    for k,v in (('w:w','12240'),('w:h','15840')):
        pass
    pg=sp.find(qn('w:pgMar'))
    for k,v in dict(left=900,right=900,top=1080,bottom=1080,header=540,footer=540).items(): pg.set(qn('w:'+k),str(v))
body_sp=d.element.body.find(qn('w:sectPr'))
setpage(body_sp)
cols=OxmlElement('w:cols'); cols.set(qn('w:num'),'2'); cols.set(qn('w:space'),'360')
body_sp.append(cols)
front=copy.deepcopy(body_sp)
for c in front.findall(qn('w:cols')): front.remove(c)
c1=OxmlElement('w:cols'); c1.set(qn('w:num'),'1'); front.append(c1)
t=OxmlElement('w:type'); t.set(qn('w:val'),'continuous')
# front section break is stored in the last front-matter paragraph (Index Terms); the *following* section carries type=continuous
body_sp.insert(0,t) if body_sp.find(qn('w:type')) is None else None
for f in front.findall(qn('w:footerReference')): front.remove(f)
pPr=ps[3]._p.get_or_add_pPr(); pPr.append(front)
set_text(ps[0],'Event-Driven LLM-Agent Simulation for Pre-Testing Social Media Seeding Strategies')
d.core_properties.title=ps[0].text
set_text(ps[1],'Ook Lee')
a=copy.deepcopy(ps[1]._p); ps[1]._p.addnext(a)
from docx.text.paragraph import Paragraph
ap=Paragraph(a,ps[1]._parent); set_text(ap,'Department of Information System, Hanyang University, Seoul, Korea')
for r in ap.runs: r.font.size=Pt(9); r.font.italic=True
ap.paragraph_format.space_after=Pt(8)
ps[1].paragraph_format.space_after=Pt(0)
last=d.paragraphs[-1]
bio=copy.deepcopy(d.paragraphs[100]._p) if False else None
BIO=('is a professor in the Department of Information System at Hanyang University, Seoul, Korea. He received the Ph.D. degree in management of information systems from Claremont Graduate University, USA. He was a faculty member at the University of Queensland, Australia, the University of Nevada at Las Vegas, USA, and North Carolina A&T State University, USA. His research interests include information system economics, information system maintenance, large language models, and critical social theory.')
tb=d.add_table(rows=1,cols=2)
tb.autofit=False
for c,w in zip(tb.rows[0].cells,(Inches(1.1),Inches(2.35))): c.width=w
pc,tc=tb.rows[0].cells
pc.paragraphs[0].add_run().add_picture('/tmp/claude-0/-home-user-fluffy-meme/85836436-e025-583a-b00a-2fcdeefc7114/images/1.jpg',width=Inches(1.0),height=Inches(1.25))
tp=tc.paragraphs[0]
r=tp.add_run('Ook Lee '); r.bold=True; r.font.size=Pt(8)
r=tp.add_run(BIO); r.font.size=Pt(8)
tp.alignment=3
for c in (pc,tc):
    for p in c.paragraphs: p.paragraph_format.space_after=Pt(0); p.paragraph_format.first_line_indent=Inches(0)
tp_=tb._tbl.tblPr
for e in tp_.findall(qn('w:tblW')): tp_.remove(e)
w=OxmlElement('w:tblW'); w.set(qn('w:w'),str(int(3.45*1440))); w.set(qn('w:type'),'dxa'); tp_.append(w)
lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tp_.append(lay)
for gc,wd in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')),(1.1,2.35)): gc.set(qn('w:w'),str(int(wd*1440)))
tb._tbl.addprevious(OxmlElement('w:p'))
# keep section properties last
body=d.element.body; body.append(body.find(qn('w:sectPr')))
d.core_properties.author='Ook Lee'
d.save(OUT)
