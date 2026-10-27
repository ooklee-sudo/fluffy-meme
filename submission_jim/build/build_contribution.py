import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
SUB = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
d = Document(); s = d.sections[0]; s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
d.styles['Normal'].font.name = 'Times New Roman'; d.styles['Normal'].font.size = Pt(11.5)
def P(t, bold=False, italic=False, after=6):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic; p.paragraph_format.space_after = Pt(after); return p
def B(items):
    for t in items: d.add_paragraph(t, style='List Bullet')
t = P('Statement of Contribution', bold=True); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
P('When Can Marketers Trust Synthetic Consumers? Validity Regions and a Validation Protocol for LLM Digital Twins in Digital Experimentation', bold=True)
P('1. Short version (about 100 words, for submission forms)', bold=True)
P('Marketers are starting to use LLM-based synthetic consumers to pre-screen digital experiments, but when such agents can be trusted is unknown. We offer a falsifiable validity-region account, test it with pre-registered studies on new simulated respondents, a reasoning-by-temperature grid, and 999 public headline A/B tests, and map its stability across 13 model versions. Defaults reverse with paradigm familiarity, reasoning lowers factual errors but raises bias, and click predictions are weak and compressed. From these results we derive a version-specific validation protocol for synthetic evidence in digital experimentation.')
P('2. Full statement', bold=True)
P('The problem. Digital experimentation is core infrastructure for interactive marketing, and synthetic consumers built from large language models are being proposed as a cheap way to screen designs and run market research before customer traffic is spent. Marketing research lacks a theory of when such agents are informative about consumer behavior, and managers lack rules for using them. Existing studies mostly report whether twins reproduce classic effects on average.')
P('What the paper contributes.', after=2)
B(['Theory. A validity-region account that states four conditions under which a synthetic agent is informative about a treatment effect: the effect operates through textual cues (V1), it is not mediated by effort or transaction costs (V2), the paradigm is not memorized (V3), and the agent\'s reasoning regime matches the human one (V4). The account yields seven falsifiable propositions. It was built from exploratory data and tested out of sample; the paper reports which parts survive (V3 and the cost-based dissociation) and which do not (universal overshooting of cue-based effects; reasoning as a debiasing mechanism).',
   'Evidence. A validity map across effect types, persona architectures, reasoning settings, and prompt templates, based on three pre-registered confirmatory studies (twin study, reasoning grid, field test on public A/B tests) and one exploratory stability study over 13 model versions. Main findings: default effects reverse with paradigm familiarity in all five model families; naming the theory inflates canonical effects; more reasoning reduces factual errors but increases cue-based bias; prompt wording matters little but each model version has its own error profile; headline-click prediction beats chance by about nine points only without personas, and predicted effects are compressed.',
   'Design knowledge. A five-step validation protocol for managers and marketing researchers: classify the effect type, consult the validity map, validate the exact model version and configuration, never rely on synthetic magnitudes, and treat textual field predictions as a weak prior, with documentation duties for auditability.'])
P('Relation to prior work. Peng et al. (2026) show that twins predict individuals poorly and behave hyper-rationally; Hullman et al. (2026) argue for validating simulations against held-out human data; Persson et al. (2026) and Maiorano (2026) evaluate LLM prediction of headline tests. Our contribution is to specify validity conditions that predict where fidelity should hold, test them with pre-registered out-of-sample propositions across effect types, reasoning regimes, and model generations, and turn the results into a protocol. We do not claim that synthetic agents model human cognition, or that the map generalizes beyond the seven effects and one US panel studied.')
P('Who benefits. Managers deciding whether to use synthetic consumers in digital experimentation and market research; marketing researchers who use synthetic respondents and need standards for validation and reporting; and those concerned with governance of AI-mediated choice architectures, since the largest failure we document concerns default design.')
P('3. Author contributions (CRediT taxonomy; to be completed by the authors)', bold=True)
rows = [('Conceptualization', '[Author(s)]'), ('Methodology', '[Author(s)]'), ('Software', '[Author(s)]'), ('Validation', '[Author(s)]'), ('Formal analysis', '[Author(s)]'), ('Investigation', '[Author(s)]'),
        ('Data curation', '[Author(s)]'), ('Writing: original draft', '[Author(s)]'), ('Writing: review and editing', '[Author(s)]'), ('Visualization', '[Author(s)]'), ('Supervision', '[Author(s)]'), ('Project administration', '[Author(s)]'), ('Funding acquisition', '[Author(s) or none]')]
t = d.add_table(rows=1, cols=2); t.style = 'Table Grid'; t.rows[0].cells[0].text = 'Role'; t.rows[0].cells[1].text = 'Contributor(s)'
for a, b in rows: c = t.add_row().cells; c[0].text = a; c[1].text = b
d.add_paragraph()
P('Use of AI assistance. An AI model (Anthropic Claude) assisted with study design discussion, writing and running simulation and analysis code, and drafting text. It is not an author. The authors are responsible for the content and verified [to be confirmed] the reported numbers against the repository. [Confirm wording against the journal\'s AI-use policy.]', italic=True)
d.save(os.path.join(SUB, '07_Contribution_Statement.docx')); print('ok')
