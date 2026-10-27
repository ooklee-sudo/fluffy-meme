"""Builds the one-page cover letter (05_Cover_Letter.docx). Usage: python3 build_cover_letter.py"""
import os
from docx import Document
from docx.shared import Pt, Inches
SUB = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
TITLE = 'When Can Marketers Trust Synthetic Consumers? Validity Regions and a Validation Protocol for LLM Digital Twins in Digital Experimentation'
d = Document(); s = d.sections[0]; s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(0.8)
d.styles['Normal'].font.name = 'Times New Roman'; d.styles['Normal'].font.size = Pt(10.5)
def P(t, b=False, after=5):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = b; p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0); return p
P('[Date]'); P('[Name], Editor-in-Chief\nJournal of Interactive Marketing', after=7)
P('Re: Submission of "' + TITLE + '"', b=True, after=7)
P('Dear Editor,')
P('We submit the enclosed manuscript for consideration as a [article type] in the Journal of Interactive Marketing. Marketers increasingly use LLM-based "synthetic consumers" to pre-screen defaults, choice sets, and headlines before spending customer traffic, yet it is unclear when such agents are informative. We propose a validity-region account (textual channel, cost insensitivity, novelty, reasoning-regime match) and test seven propositions in studies preregistered at AsPredicted (#320110, #320124): a confirmatory twin study with five model families and three persona architectures, a reasoning-effort by temperature grid, and a field test on 999 public headline A/B tests.')
P('Three findings matter for practice. (1) Default effects reverse with paradigm familiarity: twins overstate the canonical organ-donation default and understate a novel default. (2) More reasoning reduces factual errors but increases cue-based bias, contrary to our registered prediction. (3) In the field test, twins beat chance by about nine percentage points only without a persona and compress effects (slope about 0.26). Because error profiles change between model generations, we offer a version-specific validation protocol.')
P('The work extends recent evaluations of digital twins (Peng et al., 2026) and of LLM prediction of headline tests (Persson et al., 2026; Maiorano, 2026) by testing a falsifiable account out of sample and by reporting rejected predictions on equal footing. All confirmatory analysis code was committed before outputs were analyzed; one multiplicity-correction indexing error was corrected without changing any verdict and is documented.')
P('Files: title page; anonymized main document; Web Appendix (PDF); replication package (code, prompts, summary results, raw model outputs). No new human-subject data were collected; all human data are public secondary data. In line with the journal\'s generative-AI policy, an AI model (Claude, Anthropic) wrote most of the code, executed the analyses, and drafted the text under our direction; this is disclosed in the main document and the title page, and we take responsibility for the content. [Authors to confirm.] The manuscript is not under review elsewhere, and all authors approved the submission. Suggested reviewers: [names].')
P('Thank you for considering our work.', after=7); P('Sincerely,\n[Corresponding author, affiliation, email], on behalf of all authors', after=0)
d.save(os.path.join(SUB, '05_Cover_Letter.docx'))
