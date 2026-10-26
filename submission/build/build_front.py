"""Builds the title page and cover letter (Word), the code package zip, and the raw-outputs zip. Usage: python3 build_front.py <scratch_dir_with_raw_outputs>"""
import sys, os, zipfile, gzip, shutil
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.abspath(os.path.join(HERE, '..', '..')); SUB = os.path.join(ROOT, 'submission'); SCR = sys.argv[1]
TITLE = 'When Can Platforms Trust Synthetic Agents? Validity Regions and Decision Rules for LLM Digital Twins in Experimentation Pipelines'
def newdoc(size=12):
    d = Document(); s = d.sections[0]; s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
    d.styles['Normal'].font.name = 'Times New Roman'; d.styles['Normal'].font.size = Pt(size); return d
def para(d, text, bold=False, italic=False, align=None, after=6):
    p = d.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic; p.paragraph_format.space_after = Pt(after)
    if align == 'c': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p
def bullets(d, items):
    for t in items: p = d.add_paragraph(style='List Bullet'); p.add_run(t)

# ---------------- title page
d = newdoc(); para(d, 'TITLE PAGE (not for blind review)', bold=True, align='c')
para(d, TITLE, bold=True, align='c', after=12)
para(d, 'Running title: Validity regions for synthetic agents', italic=True, align='c', after=12)
para(d, 'Authors and affiliations', bold=True)
para(d, '[Author 1 name], [department, institution, city, country], [email], [ORCID]\n[Author 2 name], [department, institution, city, country], [email], [ORCID]')
para(d, 'Corresponding author', bold=True); para(d, '[Name, postal address, email, telephone]')
para(d, 'Keywords', bold=True); para(d, 'synthetic agents; digital twins; large language models; experimentation platforms; validity; behavioral economics; A/B testing; design science')
para(d, 'Manuscript length', bold=True); para(d, 'About 5,700 words including tables and references (main manuscript); online appendix provided separately. [Check against the journal limit.]')
para(d, 'Pre-registration', bold=True); para(d, 'AsPredicted #320110 (twin study, reasoning grid) and AsPredicted #320124 (field test; overlapping amendment to #320110). Analysis code for confirmatory studies was committed before the confirmatory outputs were analyzed.')
para(d, 'Data and code availability', bold=True); para(d, 'Code, pre-registrations, summary results, and raw model outputs are provided as supplementary files (code package and raw-output archive). Human data are public secondary data: Twin-2K-500 and its mega-study (Hugging Face, CC BY 4.0 and Apache-2.0) and the Upworthy Research Archive (OSF jd64p). [Add repository DOI / permanent link when deposited.]')
para(d, 'Ethics', bold=True); para(d, 'No new human-subject data were collected. [Institutional determination of exemption to be added by the authors.]')
para(d, 'Funding', bold=True); para(d, '[Funding sources, or "none"]. API compute for the simulations (about 306 USD) was paid by [the authors / institution].')
para(d, 'Conflicts of interest', bold=True); para(d, '[Authors to complete.]')
para(d, 'Use of AI assistance', bold=True); para(d, 'An AI model (Anthropic Claude) assisted with study design discussion, code for the simulations and analyses, and drafting of the manuscript, appendix, and cover letter. The authors are responsible for all content. [Confirm wording against the journal\'s AI-use policy.]')
para(d, 'Acknowledgments', bold=True); para(d, '[To be completed.]')
d.save(os.path.join(SUB, '01_Title_Page.docx'))

# ---------------- cover letter
d = newdoc(11.5)
para(d, '[Date]'); para(d, '[Name], Editor-in-Chief\nInformation Systems Research\n[Address / submission system]\n', after=10)
para(d, 'Re: Submission of "' + TITLE + '"', bold=True, after=10)
para(d, 'Dear Editor,')
para(d, 'We submit the enclosed manuscript for consideration as [article type, e.g., Research Article; confirm] in Information Systems Research.')
para(d, 'The question. Experimentation is core infrastructure for digital platforms, and large language models now offer "synthetic agents" that could pre-screen designs before human traffic is spent. Whether and when such agents can be trusted in an experimentation pipeline is unresolved. Existing work asks whether twins reproduce effects on average. Our paper asks the pipeline question: for which kind of treatment effect, in which model configuration, and for how long does a synthetic agent give information about direction, magnitude, or neither?')
para(d, 'What we do. We propose a validity-region account in which a synthetic agent is informative about a treatment effect only if the effect operates through textual cues, does not depend on effort costs, lies outside memorized paradigms, and matches the human reasoning regime. We derive seven propositions and test them under two AsPredicted pre-registrations (#320110, #320124): a confirmatory twin study with new simulated respondents and a new prompt template (five model families, three persona architectures), a reasoning-effort by temperature grid (36 cells), and a field test on 999 public headline A/B tests. A separate exploratory study varies prompt templates and model generations (13 model versions).')
para(d, 'What we find.', after=2)
bullets(d, ['Default effects reverse with paradigm familiarity in all five model families: twins overstate the canonical organ-donation default and understate a novel default that humans respond to strongly; naming the theory in the prompt inflates canonical effects.',
            'More reasoning reduces factual errors but increases cue-based bias, and the two do not co-move; this contradicts our registered prediction.',
            'In the field test, twin predictions beat chance by about nine percentage points only when no persona is supplied, and predicted effects are compressed (slope about 0.26), not inflated; the theory-derived moderator did not hold.',
            'Prompt wording matters little across 13 model versions, but each model has its own error profile that changes between generations. The practical implication is a version-specific validation protocol for synthetic evidence, which we state as design knowledge.'])
para(d, 'Contribution and relation to prior work. The closest studies are Peng et al. (2026), who show that twins fail on individual prediction and behave hyper-rationally, Hullman et al. (2026), who argue for validating simulations against held-out human data, and Persson et al. (2026) and Maiorano (2026), who evaluate LLM prediction of headline tests. We differ in testing a falsifiable account with pre-registered, out-of-sample propositions across effect types, persona architectures, reasoning regimes, and model generations, and in translating the results into a validation protocol for platform teams. We acknowledge that the theory built from exploratory data was only partly confirmed, and we report rejected predictions on equal footing.')
para(d, 'Transparency. All analysis code for the confirmatory studies was committed before the confirmatory outputs were analyzed; deviations from the registrations are listed in the appendix; an indexing error in our multiplicity correction was found and corrected without changing any verdict, and is documented. Code, prompts, summary results, and raw outputs are supplied.')
para(d, 'Limitations we state in the paper. Human benchmarks come from one US online panel and seven effects; results describe model versions available in October 2026; the hallucination measure showed almost no variation; and the stability study was not pre-registered.')
para(d, 'Required statements. [Authors to confirm each:] The manuscript is not published and not under review elsewhere. All authors approved the submission. We have no conflicts of interest to declare [or list them]. No new human-subject data were collected; all human data are public secondary data [institutional exemption statement]. An AI model assisted with code, analysis execution, and drafting, and the authors take responsibility for the content, as disclosed on the title page.')
para(d, 'Suggested editor/department: [to be completed]. Suggested reviewers: [names, affiliations, emails, with reasons; avoid collaborators and the authors of the closest prior work unless disclosed]. Reviewers we ask to exclude: [none / names].')
para(d, 'Thank you for considering our work.', after=14); para(d, 'Sincerely,\n\n[Corresponding author name]\n[Affiliation, email]\non behalf of all authors')
d.save(os.path.join(SUB, '05_Cover_Letter.docx'))

# ---------------- code zip
README = '''# Code package: When Can Platforms Trust Synthetic Agents?

Contents
- pilot/            simulation runners (run_full.py, run_h6.py, run_c1.py, run_h6c.py, run_c2.py, run_gen.py), analysis scripts (analyze_*.py), factual battery (h6_battery.py)
- benchmarks/       compute_human_benchmarks.py: human effects from the public Twin-2K-500 data and mega-study
- prereg/           pre-registration texts (#320110, #320124), amendment, working document
- results/          summary tables of every study (c1, c2, e1, e2, h6c, s)
- paper_src/        manuscript source (content.js, make.js) and submission/build/ builders for appendix, title page, cover letter

Mapping of study names: E1 = full-run simulations (pilot/run_full.py), E2 = reasoning pilot (run_h6.py), C1 = run_c1.py, H6c = run_h6c.py, C2 = run_c2.py, S = run_gen.py.

Setup
  python3 -m pip install -r requirements.txt
  An API key for an LLM aggregator is needed for the runners (they call https://openrouter.ai/api/v1/chat/completions; set credentials in your environment or proxy).
  Download the public data (not redistributed here): huggingface.co/datasets/LLM-Digital-Twin/Twin-2K-500 (question_catalog_and_human_response_csv/wave1_3_response_label.csv, wave4_response_label.csv, full_persona/chunks/*.parquet),
  huggingface.co/datasets/LLM-Digital-Twin/Twin-2K-500-Mega-Study (data/default_eric, context_effects, targeting_fairness parquet), and the Upworthy Research Archive (OSF jd64p; exploratory and confirmatory package CSVs).
  persona_summary.parquet is built from the persona chunks (columns pid, persona_summary; PID = "pid_" + pid).

Run order (examples; <data> is the data directory, <out> an output directory)
  python3 benchmarks/compute_human_benchmarks.py <data> <data> [<twin_csv_dir>]
  python3 pilot/run_c1.py <data> <out>/c1 600 120      then  python3 pilot/analyze_c1.py <out>/c1/raw.jsonl <out>/c1/final 1000
  python3 pilot/run_h6c.py <data> <out>/h6c 100 60     then  python3 pilot/analyze_h6c.py <out>/h6c/raw.jsonl <out>/h6c/final
  python3 pilot/run_c2.py <data> <upworthy confirmatory csv> <out>/c2conf 1000 2026 5 100 persona,nopersona
  python3 pilot/analyze_c2.py <out>/c2conf <out>/c2conf/final2 500   and   python3 pilot/analyze_c2_verdict.py <out>/c2conf <out>/c2conf/final 1000
  python3 pilot/run_gen.py <data> <out>/gen 200 130    then  python3 pilot/analyze_gen.py <out>/gen/raw.jsonl <out>/gen/res 400
Seeds: 2026 (respondent draws, test draw, bootstrap), 11 (C2 development), 2027 (Study S).
Rebuild the manuscript: cd paper_src && node make.js ../manuscript.docx (needs Node.js and the docx package). Rebuild the appendix: python3 submission/build/build_appendix.py <data> out.docx.

Notes
- Raw outputs are in the separate archive 06_Raw_Model_Outputs.zip (one gzip JSON-lines file per study). Each line holds respondent id, model, settings, response text, cost, and the questions shown.
- Confirmatory analysis scripts were committed before outputs were analyzed; see the git history of the project repository.
- Known issues: fairness condition tag inverted in the public mega-study file (we code from displayed text); an indexing error in the BH routine was corrected (no verdict changed).
- License: [to be chosen by the authors].
'''
REQ = 'pandas>=2.2\nnumpy>=1.26\nscipy>=1.11\nrequests>=2.31\npyarrow>=15\npython-docx>=1.1\n'
zp = os.path.join(SUB, '04_Code_Package.zip')
with zipfile.ZipFile(zp, 'w', zipfile.ZIP_DEFLATED) as z:
    z.writestr('README_CODE.md', README); z.writestr('requirements.txt', REQ)
    for sub in ['pilot', 'benchmarks', 'prereg', 'paper_src']:
        for dp, dn, fn in os.walk(os.path.join(ROOT, sub)):
            dn[:] = [x for x in dn if x != '__pycache__']
            for f in fn:
                if f.endswith('.pyc'): continue
                full = os.path.join(dp, f); z.write(full, os.path.relpath(full, ROOT))
    for sub in ['c1', 'c2', 'e1', 'e2', 'h6c', 's']:
        for dp, dn, fn in os.walk(os.path.join(ROOT, 'results', sub)):
            for f in fn: full = os.path.join(dp, f); z.write(full, os.path.relpath(full, ROOT))
    for f in ['build_appendix.py', 'build_front.py']: z.write(os.path.join(HERE, f), 'submission/build/' + f)
    z.write(os.path.join(ROOT, 'submission', '02_Manuscript_blind.docx'), 'manuscript/02_Manuscript_blind.docx')
# ---------------- raw outputs zip
RAWREAD = '''Raw model outputs (JSON lines, gzip). One line per API call.
Fields: study-specific. Common: pid (Twin-2K-500 participant id, "pid_<n>"), model, effort/temp/arch/template/arm (where applicable), content (model response text), cost (USD), qs/conds (questions shown and condition assignment).
Files: E1 (E1_full_runs.jsonl.gz), E2 (E2_reasoning_pilot.jsonl.gz), C1, H6c, C2 (C2_field_confirmatory), C2dev (development run on the exploratory package), S (stability study). tests.json is the selected Upworthy tests (headlines, impressions, clicks) for C2 and C2dev.
Persona text is not included (Twin-2K-500 CC BY 4.0 data are available from the original source).
'''
rp = os.path.join(SUB, '06_Raw_Model_Outputs.zip')
with zipfile.ZipFile(rp, 'w', zipfile.ZIP_STORED) as z:
    z.writestr('README_RAW.txt', RAWREAD)
    for name, rel in [('E1_full_runs', 'full2'), ('E2_reasoning_pilot', 'h6'), ('C1_confirmatory_twin', 'c1'), ('H6c_reasoning_grid', 'h6c'), ('C2_field_confirmatory', 'c2conf'), ('C2dev_development', 'c2dev'), ('S_stability', 'gen')]:
        src = os.path.join(SCR, rel, 'raw.jsonl'); tmp = os.path.join('/tmp', name + '.jsonl.gz')
        with open(src, 'rb') as fi, gzip.open(tmp, 'wb', 9) as fo: shutil.copyfileobj(fi, fo)
        z.write(tmp, name + '.jsonl.gz'); os.remove(tmp)
        tj = os.path.join(SCR, rel, 'tests.json')
        if os.path.exists(tj): z.write(tj, name + '_tests.json')
print('built', [f for f in os.listdir(SUB) if not f.startswith('build')])
