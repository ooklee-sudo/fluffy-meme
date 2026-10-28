from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
d = Document(); s = d.sections[0]
s.page_width, s.page_height = Inches(8.27), Inches(11.69)
for a in ("left_margin", "right_margin"): setattr(s, a, Inches(0.7))
s.top_margin = s.bottom_margin = Inches(0.55)
st = d.styles["Normal"]; st.font.name = "Malgun Gothic"; st.font.size = Pt(8.8); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
st.paragraph_format.space_after = Pt(1.5); st.paragraph_format.line_spacing = 1.0
def para(text, lead=None, size=None, align=None, after=None):
    p = d.add_paragraph()
    if lead: r = p.add_run(lead); r.bold = True
    r = p.add_run(text)
    if size: [setattr(x.font, "size", Pt(size)) for x in p.runs]
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
def head(t):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(10); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = t.add_run("실패는 실패를 부르는가? LLM 에이전트 오류의 자기 흥분 점과정 모형과 사람 인계 규칙"); r.bold = True; r.font.size = Pt(12.5)
para("Hawkes 과정으로 본 에이전트 신뢰성: 연쇄의 측정, 인과 식별, 의사결정 규칙  ·  연구 설계안 v0 (한 장)", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
head("1. 문제와 연구 질문")
para("도구를 쓰는 LLM 에이전트는 한 번의 오류 뒤에 오류를 연달아 내는 경향이 보고되지만(AgentDebug, 2025; 연쇄 증폭 계수, ICML 2026), 이 '연쇄'를 확률 모형으로 측정하고 사람에게 넘길 시점을 정하는 근거 있는 규칙은 없다. 실무의 인계 기준은 '오류 3번' 같은 경험칙이다.")
para("RQ1 에이전트의 도구 오류는 포아송(독립)보다 몰려서 일어나는가? 자기 흥분의 크기(분기비 n: 오류 하나가 평균적으로 유발하는 추가 오류 수)는 얼마인가?  RQ2 몰림은 오류가 오류를 부르는 '전염'인가, 원래 어려운 작업이라 오류가 많은 '이질성'인가?  RQ3 조건부 강도 λ(t)는 최종 실패를 오류 횟수보다 더 일찍, 더 정확히 예측하는가?  RQ4 비용(사람 처리비, 실패 손실)을 반영한 λ 기반 인계 규칙은 고정 규칙보다 나은가?")
head("2. 모형")
para("시간 축은 에이전트의 단계(step) 번호. 단계 t의 오류 강도 λ(t) = μ_task + Σ_{t_i<t} α·g(t − t_i), g는 지수 감쇠 커널, 분기비 n = α·Σg. 비교 모형: 포아송(상수 μ), 비동질 포아송(단계에 따른 추세), 음이항/작업별 무작위효과(이질성만), Hawkes(전염), Hawkes+무작위효과(둘 다). 표시(mark): 오류 유형(실행 오류, 잘못된 인자, 시간 초과, 정책 위반)별 교차 흥분. 적합도: 시간 재척도화 검정, 우도비, 예측 로그우도.")
head("3. 식별 전략 (핵심)")
para("관측 데이터만으로는 전염과 이질성이 섞인다(같은 작업 안에서 오류가 몰리는 이유가 둘 다 될 수 있음). 그래서 오류 주입 실험을 둔다: 무작위로 정한 단계에서 도구 호출을 강제로 실패시키면, 그 뒤 자연 오류 강도의 증가가 곧 인과적 흥분 효과다. 주입 확률·위치·오류 형식(원시 오류 페이지 대 형식화된 오류)을 무작위로 바꿔 흥분이 오류 정보의 형식에 따라 달라지는지도 본다.")
head("4. 데이터")
para("(a) 공개 궤적: SWE-bench 계열 에이전트 궤적(수만 건, 단계별 명령과 오류 기록), τ-bench/τ²-bench 고객 응대 궤적. (b) 직접 실행: τ²-bench 소매·항공 도메인에서 모델 3~4종 × 조건(기준, 무작위 오류 주입, 형식화된 오류 주입). 최종 성공 여부는 벤치마크 평가기가 판정.", lead="")
head("5. 가설 (사전 등록)")
para("H1 오류는 포아송보다 몰리고, 이질성을 통제한 뒤에도 분기비 n > 0이다(전염 존재).  H2 n은 약한 모델일수록 크고 1에 가까운 경우가 있다(임계 근처 = 연쇄 붕괴 위험).  H3 주입 오류는 이후 자연 오류 강도를 올린다(인과적 흥분), 형식화된 오류는 그 효과를 줄인다.  H4 처음 k단계의 λ(t)는 오류 횟수보다 최종 실패를 더 잘 예측한다(AUC 비교).  H5 λ가 임계값 c*/(손실) 를 넘을 때 인계하는 규칙이 '오류 k회' 규칙보다 같은 인계율에서 더 많은 실패를 막는다.")
head("6. 기여와 투고처")
para("에이전트 신뢰성의 점과정 측정 틀(분기비라는 한 숫자로 연쇄 위험 요약), 오류 주입을 통한 전염의 인과 식별, 비용 기반 사람-AI 인계 규칙. 경영대: ISR, MISQ(사람-AI 위임과 통제), Management Science(운영), JMIS. 방법 중심으로 쓰면 IEEE TAI, NeurIPS/ICML 평가 트랙.")
head("7. 비용·일정·위험")
para("공개 궤적 분석은 무료, 직접 실행은 API 비용 대략 수백 달러, 6~8주(추정치). 위험: 단계 번호를 시간으로 쓰는 가정, 궤적마다 다른 오류 기록 방식(오류 정의 표준화 필요), 공개 궤적의 오류 표시가 불완전할 수 있음, 선행연구(연쇄 증폭 계수)와의 차별화 필요.")
d.save("Agent_Hawkes_Design.docx")
