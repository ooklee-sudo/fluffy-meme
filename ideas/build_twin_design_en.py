from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
d = Document(); s = d.sections[0]
s.page_width, s.page_height = Inches(8.27), Inches(11.69)
for a in ("left_margin", "right_margin"): setattr(s, a, Inches(0.7))
s.top_margin = s.bottom_margin = Inches(0.55)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(9.5); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(1.5); st.paragraph_format.line_spacing = 1.0
def para(text, bold_lead=None, size=None, after=None, align=None):
    p = d.add_paragraph()
    if bold_lead: r = p.add_run(bold_lead); r.bold = True; (setattr(r.font, "size", Pt(size)) if size else None)
    r = p.add_run(text); (setattr(r.font, "size", Pt(size)) if size else None)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    return p
def head(t):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(10); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = t.add_run("Do Digital Twins Behave as Behavioral-Economic Theory Predicts?"); r.bold = True; r.font.size = Pt(13)
para("A theory-level fidelity test of LLM twins: classic experiments, rewritten items and digital-platform contexts  ·  Research design v0 (one page)  ·  Ook Lee, Hanyang University", size=8.5, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
head("1. Research questions")
para("Do LLM-based digital twins show the effects that behavioral-economic theories predict, as people do? The criterion is not the accuracy of individual answers but the direction, size and conditionality of the effects a theory predicts. Prior work (Horton 2023; Toubia et al. 2025, Twin-2K-500; Binz et al. 2025, Centaur) focused on reproducing classic effects and paid less attention to contamination (models may have seen the original experiments in training) and to platform contexts.")
para("RQ1 Do twins reproduce the direction and size of theory-predicted effects?  RQ2 Do they reproduce moderator interactions (conditionality)?  RQ3 Do effects shrink on new items with the same structure as classic items (contamination)?  RQ4 In platform contexts (defaults, price presentation, dark patterns), do the effect sizes twins predict fall within the range reported by published field experiments and meta-analyses?  RQ5 How does bias vary with the model's training stage (base pretraining → instruction tuning → reasoning model) and with resource limits (thinking budget, context length, sampling temperature), and does it move in the same direction as hallucination (confabulation) under the same manipulations?")
head("2. Theories (six: well-documented effect sizes, relevant to platforms, and with differing replication records)")
para("Prospect theory / loss framing (loss aversion)  ·  default effect  ·  anchoring  ·  decoy (attraction) effect  ·  choice overload (contested replication; serves as a contrast)  ·  fairness (ultimatum game)")
head("3. Design")
para("Persona from demographics only vs twin built from an individual's answers (Twin-2K-500 format) vs a prompt that names the theory (leakage check). Four to five current models; repeated responses per item give a distribution.", bold_lead="Twins. ")
para("Compare paired public base and instruction-tuned models of the same family for training stage, and vary thinking budget, context and temperature within one model. In each condition measure the bias effect size together with the hallucination rate on a set of factual questions. Training a new model from scratch is excluded for cost.", bold_lead="RQ5 manipulations. ")
para("(1) classic items (contamination risk); (2) rewritten items with the same structure and new content; (3) preregistered experiments published after the model's training cutoff; (4) newly written platform items (e.g., subscription defaults, price anchors, recommendation order).", bold_lead="Stimuli. ")
para("No new human experiment. All human benchmarks come from public sources: (i) item-level comparisons use Twin-2K-500 and Many Labs / preregistered replication data; (ii) platform items (4) are benchmarked against pooled effect sizes and confidence intervals from published field experiments and meta-analyses (e.g., a meta-analysis of default effects), testing whether twin effect sizes fall within those intervals and whether the ordering of effects matches. Items are designed to match the conditions those meta-analyses cover (opt-in vs opt-out, type of default option, etc.).", bold_lead="Human benchmark. ")
para("Effect size (twin vs human), sign agreement, reproduction of moderator interactions, distributional distance, individual-level correlation, contamination index = (effect on rewritten items / effect on classic items).", bold_lead="Measures. ")
head("4. Hypotheses (preregistered)")
para("H1 Twins reproduce the direction of most classic effects (80% or more) but overstate their size.  H2 Effects are smaller on rewritten than on classic items (contamination index < 1).  H3 Moderator interactions are reproduced less often than main effects.  H4 Adding individual answers improves average effects but barely improves the reproduction of heterogeneity.  H5 On platform items, the twins' ordering of effects (which design produces the larger effect) correlates positively with the condition-wise ordering in meta-analyses, while sizes fall systematically outside the meta-analytic intervals (overstatement).  H6 Bias shrinks after instruction tuning and in reasoning models but does not vanish, and tighter resource limits (smaller thinking budget, higher temperature) increase bias and hallucination together. This is consistent with the hypothesis that both arise from a shared computational principle, approximate prediction under limited resources (resource rationality, predictive processing); it does not establish the mechanism in the human brain (functional similarity ≠ the same mechanism; bias inherited through training data is only partly separated by the base-model comparison).  Success and failure criteria (signs, lower bounds on correlations) are fixed before the data are seen.")
head("5. Contribution and target journals")
para("A theory-by-theory fidelity map, a design that controls contamination, and a predictive test in platform contexts, framed as boundary conditions for using generative-AI twins as surrogate users in IS design evaluation. Business-school targets: MIS Quarterly, Information Systems Research (theory contribution required), Management Science, Marketing Science. Nature Machine Intelligence only with a new methodological contribution (e.g., a new way to build twins).")
head("6. Cost, schedule and risks")
para("Public data only: API costs of roughly a few hundred US dollars, four to six weeks, no IRB review needed (estimates). Risks: results date quickly as models change; contamination control is imperfect; competition from related work; meta-analytic effect sizes come from conditions that may not match the items, so only comparable conditions are selected and preregistered.")
d.save("Digital_Twin_BehavioralEcon_Design_EN.docx")
