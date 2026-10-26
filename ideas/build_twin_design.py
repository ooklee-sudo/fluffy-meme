from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
d = Document(); s = d.sections[0]
s.page_width, s.page_height = Inches(8.27), Inches(11.69)
for a in ("left_margin", "right_margin"): setattr(s, a, Inches(0.7))
s.top_margin = s.bottom_margin = Inches(0.55)
st = d.styles["Normal"]; st.font.name = "Malgun Gothic"; st.font.size = Pt(8.8); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
st.paragraph_format.space_after = Pt(1.5); st.paragraph_format.line_spacing = 1.0
def para(text, bold_lead=None, size=None, after=None, align=None):
    p = d.add_paragraph()
    if bold_lead: r = p.add_run(bold_lead); r.bold = True; (setattr(r.font, "size", Pt(size)) if size else None)
    r = p.add_run(text); (setattr(r.font, "size", Pt(size)) if size else None)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    return p
def head(t):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(10); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = t.add_run("디지털 트윈은 행동경제학 이론대로 작동하는가?"); r.bold = True; r.font.size = Pt(13)
para("LLM 트윈의 이론 수준 충실도 검증: 고전 실험, 변형 문항, 디지털 플랫폼 맥락  ·  연구 설계안 v0 (한 장)", size=8.5, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
head("1. 연구 질문")
para("LLM 기반 디지털 트윈이 행동경제학 이론이 예측하는 효과를 사람처럼 보이는가? 개별 응답 정확도가 아니라 이론이 예측하는 효과의 방향, 크기, 조건 의존성을 기준으로 본다. 선행(Horton 2023; Toubia 외 2025 Twin-2K-500; Binz 외 2025 Centaur)은 고전 효과의 재현에 집중했고, 모델이 원 실험을 학습했을 가능성(오염)과 플랫폼 맥락은 덜 다뤘다.")
para("RQ1 트윈은 이론 효과의 방향과 크기를 재현하는가?  RQ2 조절변수 상호작용(조건 의존성)까지 재현하는가?  RQ3 고전 문항과 구조가 같은 새 문항에서 효과가 줄어드는가(오염)?  RQ4 플랫폼 맥락(기본값, 가격 표시, 다크 패턴)에서 트윈이 예측한 효과가 실제 사람의 효과와 일치하는가?")
head("2. 이론 선정 (6개: 효과 문헌이 풍부하고, 플랫폼과 관련되며, 재현 가능성이 서로 다르게 알려진 것)")
para("전망이론/손실 프레이밍(손실 회피)  ·  기본값 효과  ·  닻 내리기  ·  유인(decoy) 효과  ·  선택 과부하(재현 논란 있음, 대조군 역할)  ·  공정성(최후통첩 게임)")
head("3. 설계")
para("(a) 트윈: 인구통계 페르소나만  vs  개인 응답 기반 트윈(Twin-2K-500 형식)  vs  이론 명시 프롬프트(누설 검사). 모델 4~5종(최신), 문항당 반복 응답으로 분포를 얻는다.", bold_lead="트윈. ")
para("① 고전 문항(오염 위험), ② 구조는 같고 내용을 새로 쓴 변형 문항, ③ 모델 학습 시점 이후 공개된 사전등록 실험, ④ 새로 만든 플랫폼 문항(예: 구독 기본값, 가격 닻, 추천 순서).", bold_lead="자극. ")
para("공개 데이터(Twin-2K-500, Many Labs 계열과 사전등록 복제 데이터)를 우선 사용. 새 실험(Prolific 등, IRB 필요)은 플랫폼 문항 ④ 한정 소규모로 선택.", bold_lead="사람 기준. ")
para("효과크기(트윈 대 사람), 부호 일치율, 조절변수 상호작용의 재현, 분포 거리, 개인 수준 상관, 오염 지표 = (변형 효과 / 고전 효과).", bold_lead="지표. ")
head("4. 가설 (사전 등록)")
para("H1 트윈은 고전 효과의 방향을 대부분(80% 이상) 재현하지만 크기는 사람보다 크다(과대 재현).  H2 변형 문항에서는 효과가 고전 문항보다 작다(오염 지표 < 1).  H3 조절변수 상호작용의 재현율은 주효과보다 낮다.  H4 개인 응답 정보를 더 주면 평균 효과는 개선되나 이질성 재현은 거의 개선되지 않는다.  H5 플랫폼 문항에서 트윈이 예측한 효과 순서(어떤 설계가 더 큰 효과를 내는가)는 사람의 순서와 양의 상관을 보이되, 크기는 체계적으로 편향된다.  성공/실패 기준(부호, 상관 하한)은 데이터를 보기 전에 고정한다.")
head("5. 기여와 투고처")
para("이론별 충실도 지도, 오염 통제 설계, 플랫폼 맥락의 예측 시험. 경영대 목표: Marketing Science, Management Science, ISR, MISQ(이론 기여 필요). NMI는 트윈 구성법 등 새 방법 기여가 있을 때.")
head("6. 비용·일정·위험")
para("공개 데이터만 사용 시 API 비용 대략 수백 달러, 4~6주. 새 사람 실험을 더하면 수천 달러와 IRB 수 주 추가(추정치). 위험: 모델 변화로 결과가 빨리 낡음, 오염 통제의 불완전성, 선점 경쟁, 사람 실험의 재현성.")
d.save("Digital_Twin_BehavioralEcon_Design.docx")
