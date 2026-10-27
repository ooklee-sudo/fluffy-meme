"""IEEE Transactions on Artificial Intelligence version of the defence paper (two-column Word layout).
Numbers and figures come from the same analysis prefix as build_nmi3.py."""
import sys; sys.path.insert(0, "/home/user/fluffy-meme/isr_submission/manuscript_build")
src = open("build_nmi3.py").read(); exec(src[:src.index("REFS = {}")])          # data, helpers, figures, tables T1-T5
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as _qn
REF = {
 "liang25": 'W. Liang et al., "Quantifying large language model usage in scientific papers," Nat. Hum. Behav., 2025, doi: 10.1038/s41562-025-02273-8.',
 "kobak": 'D. Kobak, R. Gonzalez-Marquez, E.-A. Horvát, and J. Lause, "Delving into LLM-assisted writing in biomedical publications through excess vocabulary," Sci. Adv., vol. 11, no. 27, 2025.',
 "gltr": 'S. Gehrmann, H. Strobelt, and A. Rush, "GLTR: Statistical detection and visualization of generated text," in Proc. 57th Annu. Meeting ACL: Syst. Demonstrations, 2019, pp. 111–116.',
 "detectgpt": 'E. Mitchell, Y. Lee, A. Khazatsky, C. D. Manning, and C. Finn, "DetectGPT: Zero-shot machine-generated text detection using probability curvature," in Proc. 40th Int. Conf. Mach. Learn., 2023.',
 "fastdetect": 'G. Bao, Y. Zhao, Z. Teng, L. Yang, and Y. Zhang, "Fast-DetectGPT: Efficient zero-shot detection of machine-generated text via conditional probability curvature," in Proc. Int. Conf. Learn. Represent., 2024.',
 "binoculars": 'A. Hans et al., "Spotting LLMs with Binoculars: Zero-shot detection of machine-generated text," in Proc. 41st Int. Conf. Mach. Learn., 2024.',
 "ghostbuster": 'V. Verma, E. Fleisig, N. Tomlin, and D. Klein, "Ghostbuster: Detecting text ghostwritten by large language models," in Proc. Conf. North Amer. Chapter ACL (NAACL), 2024.',
 "solaiman": 'I. Solaiman et al., "Release strategies and the social impacts of language models," 2019, arXiv:1908.09203.',
 "pu": 'X. Pu, J. Zhang, X. Han, Y. Tsvetkov, and T. He, "On the zero-shot generalization of machine-generated text detectors," in Findings ACL: EMNLP 2023, 2023, pp. 4799–4808.',
 "dipper": 'K. Krishna, Y. Song, M. Karpinska, J. Wieting, and M. Iyyer, "Paraphrasing evades detectors of AI-generated text, but retrieval is an effective defense," in Proc. Adv. Neural Inf. Process. Syst., vol. 36, 2023.',
 "sadasivan": 'V. S. Sadasivan, A. Kumar, S. Balasubramanian, W. Wang, and S. Feizi, "Can AI-generated text be reliably detected?" Trans. Mach. Learn. Res., 2025.',
 "shi": 'Z. Shi, Y. Wang, F. Yin, X. Chen, K.-W. Chang, and C.-J. Hsieh, "Red teaming language model detectors with language models," Trans. Assoc. Comput. Linguist., vol. 12, pp. 174–189, 2024.',
 "stumbling": 'Y. Wang et al., "Stumbling blocks: Stress testing the robustness of machine-generated text detectors under attacks," in Proc. 62nd Annu. Meeting ACL, 2024, pp. 2894–2925.',
 "huang": 'G. Huang et al., "Are AI-generated text detectors robust to adversarial perturbations?" in Proc. 62nd Annu. Meeting ACL, 2024, pp. 6005–6024.',
 "zhou": 'Y. Zhou, B. He, and L. Sun, "Navigating the shadows: Unveiling effective disturbances for modern AI content detectors," in Proc. 62nd Annu. Meeting ACL, 2024, pp. 10847–10861.',
 "silverspeak": 'A. Creo and S. Pudasaini, "SilverSpeak: Evading AI-generated text detectors using homoglyphs," in Proc. 1st Workshop GenAI Content Detection, 2025, pp. 1–46.',
 "outfox": 'R. Koike, M. Kaneko, and N. Okazaki, "OUTFOX: LLM-generated essay detection through in-context learning with adversarially generated examples," in Proc. AAAI Conf. Artif. Intell., vol. 38, 2024, pp. 21258–21266.',
 "advpara": 'Y. Cheng, V. S. Sadasivan, M. Saberi, S. Saha, and S. Feizi, "Adversarial paraphrasing: A universal attack for humanizing AI-generated text," in Proc. Adv. Neural Inf. Process. Syst., vol. 38, 2025.',
 "authormist": 'I. David and A. Gervais, "AuthorMist: Evading AI text detectors with reinforcement learning," 2025, arXiv:2503.08716.',
 "stealthrl": 'S. Ranganath and A. Ramesh, "StealthRL: Reinforcement learning paraphrase attacks for multi-detector evasion of AI-text detectors," 2026, arXiv:2602.08934.',
 "raid": 'L. Dugan et al., "RAID: A shared benchmark for robust evaluation of machine-generated text detectors," in Proc. 62nd Annu. Meeting ACL, 2024, pp. 12463–12492.',
 "mage": 'Y. Li et al., "MAGE: Machine-generated text detection in the wild," in Proc. 62nd Annu. Meeting ACL, 2024, pp. 36–53.',
 "m4": 'Y. Wang et al., "M4: Multi-generator, multi-domain, and multi-lingual black-box machine-generated text detection," in Proc. 18th Conf. Eur. Chapter ACL, 2024, pp. 1369–1407.',
 "detectrl": 'J. Wu et al., "DetectRL: Benchmarking LLM-generated text detection in real-world scenarios," in Proc. Adv. Neural Inf. Process. Syst., vol. 37, Datasets and Benchmarks Track, 2024.',
 "survey": 'J. Wu, S. Yang, R. Zhan, Y. Yuan, L. S. Chao, and D. F. Wong, "A survey on LLM-generated text detection: Necessity, methods, and future directions," Comput. Linguist., vol. 51, pp. 275–338, 2025.',
 "contam": 'P. Dingfelder and C. Riess, "Contamination in generated text detection benchmarks," 2025, arXiv:2511.09200.',
 "whyfail": 'S. Pudasaini, L. Miralles-Pechuán, D. Lillis, and M. Llorens Salvador, "Why AI-generated text detection fails: Evidence from explainable AI beyond benchmark accuracy," 2026, arXiv:2603.23146.',
 "gritsai": 'G. Gritsai, A. Voznyuk, A. Grabovoy, and Y. Chekhovich, "Are AI detectors good enough? A survey on quality of datasets with machine-generated texts," 2024, arXiv:2410.14677.',
 "bias": 'W. Liang, M. Yuksekgonul, Y. Mao, E. Wu, and J. Zou, "GPT detectors are biased against non-native English writers," Patterns, vol. 4, Art. no. 100779, 2023.',
 "watermark": 'J. Kirchenbauer, J. Geiping, Y. Wen, J. Katz, I. Miers, and T. Goldstein, "A watermark for large language models," in Proc. 40th Int. Conf. Mach. Learn., 2023, pp. 17061–17084.',
 "sand": 'H. Zhang et al., "Watermarks in the sand: Impossibility of strong watermarking for generative models," in Proc. 41st Int. Conf. Mach. Learn., 2024.',
 "judge": 'L. Zheng et al., "Judging LLM-as-a-judge with MT-Bench and Chatbot Arena," in Proc. Adv. Neural Inf. Process. Syst., vol. 36, 2023.',
 "selfpref": 'A. Panickssery, S. R. Bowman, and S. Feng, "LLM evaluators recognize and favor their own generations," in Proc. Adv. Neural Inf. Process. Syst., vol. 37, 2024.',
 "mosca": 'E. Mosca, M. H. I. Abdalla, P. Basso, M. Musumeci, and G. Groh, "Distinguishing fact from fiction: A benchmark dataset for identifying machine-generated scientific papers in the LLM era," in Proc. 3rd Workshop Trustworthy NLP, 2023, pp. 190–207.',
 "desaire": 'H. Desaire, A. E. Chua, M. Isom, R. Jarosova, and D. Hua, "Distinguishing academic science writing from humans or ChatGPT with over 99% accuracy using off-the-shelf machine learning tools," Cell Rep. Phys. Sci., vol. 4, Art. no. 101426, 2023.',
 "peerreview": 'S. Yu, M. Luo, A. Madasu, V. Lal, and P. Howard, "Is your paper being reviewed by an LLM? Benchmarking AI text detection in peer review," in Proc. Int. Conf. Learn. Represent., 2026.',
 "cameron": 'A. C. Cameron and P. K. Trivedi, Regression Analysis of Count Data, 2nd ed. Cambridge, U.K.: Cambridge Univ. Press, 2013.',
}
order = []
def C(*keys):
    for k in keys:
        if k not in order: order.append(k)
    return " [" + "], [".join(str(order.index(k) + 1) for k in keys) + "]"
def cols(sec, n):
    sp = sec._sectPr; c = sp.find(_qn("w:cols"))
    if c is None: c = OxmlElement("w:cols"); sp.append(c)
    c.set(_qn("w:num"), str(n)); c.set(_qn("w:space"), "340")
def section(n): s = d.add_section(WD_SECTION.CONTINUOUS); cols(s, n); return s
ROM = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX"]; hc = [0]; sc = [0]
def SEC(t): d.add_paragraph(t, style=ST("Heading1"))
def SUB(t): d.add_paragraph(t, style=ST("Heading2"))
def _rich(p, t):
    for part in _re.split(r"(_\{.*?\}|\*[^*]+\*)", t):
        if not part: continue
        if part.startswith("_{"): r = p.add_run(part[2:-1]); r.font.subscript = True
        elif part.startswith("*"): r = p.add_run(part[1:-1]); r.italic = True
        else: p.add_run(part)
def TXT(t, **k): _rich(d.add_paragraph(style=ST("Text")), t)
def BUL2(items):
    for it in items: _rich(d.add_paragraph(style=ST("Text")), "• " + it)
def EQN(t, num):
    p = d.add_paragraph(style=ST("Equation")); _rich(p, "\t" + t + "\t(" + num + ")")
def FIGT(path, cap, width):
    p = d.add_paragraph(); p.alignment = 1; p.add_run().add_picture(path, width=Inches(width))
    d.add_paragraph(cap, style=ST("FigureCaption"))
def TAB(cap, rows, widths, size=8):
    num, title = cap.split("\n", 1)
    p = d.add_paragraph(num, style=ST("TableTitle")); p.alignment = 1
    p = d.add_paragraph(title, style=ST("TableTitle")); p.alignment = 1
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.alignment = 1
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Inches(widths[j]); c.text = ""; pp = c.paragraphs[0]; _rich(pp, str(v))
            for r in pp.runs: r.font.size = Pt(size); r.bold = (i == 0)
    d.add_paragraph(style=ST("Text"))
ABSTRACT = f"Screens for machine-written manuscripts are usually evaluated on text that no author has tried to disguise. We evaluate nine screens on 474 human scientific papers and papers from four large language models under eight attacks, two of which query the detector's own score and keep the candidate it rates most human. Every single detector has a blind spot: keeping the least machine-like of four sentence-mixed candidates, as judged by Fast-DetectGPT, lowers its AUC on Claude papers from {cell('fd','claude','orig'):.2f} to {cell('fd','claude','bok_mix'):.2f}, and probability-based scores invert because spliced text is less probable than genuine human text. We propose a parameter-free defence: a two-sided likelihood score calibrated on a held-out set of human papers, combined with detectors whose failures differ, by averaging robustly standardized scores. It needs no attacked training data and no knowledge of the attack. Across four generators and eight attacks, the worst-case AUC rises from {wfd:.2f} for the best single detector to {wm4:.2f}–{wt2:.2f}. The gain is smaller at a 5% false-positive operating point, and the defence is not evaluated against attackers who optimize against the combination. We also show that supervised detectors partly rely on structural artifacts of corpus construction."
IMPACT = "Publishers and research-integrity offices increasingly screen submissions for machine-written text, and a wrong decision either misses fabricated papers or accuses honest authors. This study shows that every detector in common use can be pushed below chance by an author who can query it, and that a simple combination of detectors that fail in different ways keeps the worst case far higher without training on attacks. The method is transparent, cheap to deploy and easy to audit, and the paper reports where it falls short, including its operating-point limits and its untested robustness to ensemble-aware attackers. The results, code and data can help integrity offices choose and evaluate screens and help researchers report detector robustness more honestly."
from docx import Document as _Doc
import copy, re as _re
d = _Doc("tai_template/TAI_Word_Template.docx")
body = d.element.body; paras = d.paragraphs
def ST(i): return next(x for x in d.styles if x.style_id == i)
def _set(p, lead, text):
    for el in list(p._p):
        if el.tag.endswith("}r") or el.tag.endswith("}hyperlink"): p._p.remove(el)
    if lead:
        r = p.add_run(lead); r.bold = True; r.italic = True
    p.add_run(text)
_set(paras[1], None, "Combining Detectors That Fail Differently Defends Screens for Machine-Written Scholarly Text Against Adaptive Authors")
_set(paras[2], None, "")
_br = paras[1].add_run(); _br.add_break(); _ar = paras[1].add_run("Ook Lee"); _ar.font.size = Pt(11)
for el in list(paras[2]._p):
    if el.tag.endswith("}pPr"):
        fp = el.find(_qn("w:framePr"))
        if fp is not None: el.remove(fp)
_set(paras[3], "Abstract—", ABSTRACT)
_set(paras[5], "Impact Statement—", IMPACT)
_set(paras[8], "Index Terms—", "Adversarial robustness, AI-generated text detection, ensemble methods, large language models, research integrity, zero-shot detection.")
keep = set(id(p._p) for p in paras[:10])
for el in list(body):
    if el.tag.endswith("}sectPr") or id(el) in keep: continue
    body.remove(el)
for s_ in d.sections: cols(s_, 2)
for fp in body.iter(_qn("w:framePr")): fp.set(_qn("w:hRule"), "auto")             # let the title frame grow
FOOT = "Manuscript received [date]. O. Lee is with the Department of Information Systems, Hanyang University, Seoul, Republic of Korea (e-mail: ooklee@hanyang.ac.kr)."
for rel in d.part.rels.values():
    if rel.reltype.endswith("/footnotes"):
        fx = rel.target_part
        from docx.oxml import parse_xml
        root = parse_xml(fx.blob) if not hasattr(fx, "_element") else fx._element
        W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        for fnote in root.findall(W + "footnote"):
            if fnote.get(W + "type") in ("separator", "continuationSeparator"): continue
            ps = fnote.findall(W + "p")
            for p_ in ps[1:]: fnote.remove(p_)
            p0 = ps[0]; ts = [t for t in p0.iter(W + "t")]
            if ts:
                ts[0].text = FOOT
                for t in ts[1:]: t.text = ""
            for hl in list(p0.iter(W + "hyperlink")): hl.getparent().remove(hl)
            for hl in list(p0.iter(W + "highlight")): hl.getparent().remove(hl)
        if not hasattr(fx, "_element"):
            from lxml import etree as _et
            fx._blob = _et.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
SEC("Introduction")
TXT("Large language models (LLMs) make fluent scientific prose cheap, and corpus-level estimates suggest that a growing share of published papers already show signs of their use" + C("liang25", "kobak") + ". Publishers, preprint servers and integrity offices therefore screen manuscripts. Zero-shot detectors based on token probabilities" + C("gltr", "detectgpt", "fastdetect", "binoculars") + " and supervised classifiers" + C("solaiman", "ghostbuster") + " report high accuracy, but usually on unmodified output of the generators they were designed for. Evaluations that include adversaries show that paraphrasing, perturbation and detector-guided rewriting defeat them" + C("dipper", "sadasivan", "stumbling", "huang", "advpara") + ", that accuracy falls on unseen generators and domains" + C("mage", "raid", "m4") + ", and that benchmarks can contain artifacts that classifiers learn as shortcuts" + C("contam", "whyfail") + ".")
TXT("A screen deployed at a journal faces an author who can try drafts against public detectors. Two questions follow. How badly does each kind of detector fail when the author can query its score? And is there a defence that does not require knowing the attack in advance? We study both on scientific papers, where false accusations are costly and detectors are already in use" + C("desaire", "mosca", "peerreview") + ".")
TXT("Our contributions are:")
BUL2([f"An adversarial evaluation of nine screens, from a transparent count-based screen to Fast-DetectGPT, supervised classifiers and an LLM judge, on four generators and eight attacks, including two detector-guided attacks; guided sentence mixing drives Fast-DetectGPT below chance on Claude papers.",
      "An analysis of why probability-based scores invert under perturbation, and a two-sided score calibrated on held-out humans that turns the inversion into a signal.",
      f"A parameter-free ensemble of standardized scores of detectors that fail differently, which raises the worst-case AUC from {wfd:.2f} to {wm4:.2f}–{wt2:.2f} without attacked training data.",
      "A control showing that supervised accuracy partly rests on structural artifacts of how human and machine texts were obtained, and an out-of-field test whose high numbers we argue should not be trusted for this reason.",
      "Code, per-document scores and machine-generated corpora released for reproduction."])
SEC("Related Work")
SUB("Detectors")
TXT("Probability-based detectors score how expected a text is under a language model: log-likelihood and rank statistics" + C("gltr") + ", probability curvature" + C("detectgpt", "fastdetect") + " and the ratio of perplexity to cross-perplexity" + C("binoculars") + ". Supervised detectors fine-tune a classifier, such as RoBERTa, on labelled human and machine text" + C("solaiman", "ghostbuster") + "; their transfer to new generators is limited" + C("pu", "m4") + ". LLM judges have been used to assess text quality" + C("judge") + ", and they can recognize their own generations" + C("selfpref") + ". Watermarking" + C("watermark") + " avoids detection but needs generator cooperation, and strong watermarks may be impossible" + C("sand") + ". Detectors can also be biased against non-native writers" + C("bias") + ". A survey is given in" + C("survey") + ".")
SUB("Attacks")
TXT("Paraphrasing with a dedicated model" + C("dipper") + " or recursively" + C("sadasivan") + " evades most detectors; character-level noise and homoglyphs" + C("huang", "zhou", "silverspeak") + ", red-teaming by LLMs" + C("shi") + " and co-adaptation" + C("outfox") + " are further threats. Attacks that use the detector's own output include detector-guided paraphrasing" + C("advpara") + " and reinforcement-learning paraphrasers trained on detector rewards" + C("authormist", "stealthrl") + ", one of which targets an ensemble of detectors" + C("stealthrl") + ". Benchmarks such as RAID" + C("raid") + ", MAGE" + C("mage") + ", M4" + C("m4") + " and DetectRL" + C("detectrl") + " collect attacks, generators and domains.")
SUB("Data artifacts")
TXT("Detection datasets can carry structural cues unrelated to authorship, such as chatbot preambles" + C("contam") + " or format and length differences that feature-based detectors exploit" + C("whyfail", "gritsai") + ". To our knowledge, no earlier study combines detector-guided attacks with an artifact control on scholarly text, although each element has precedents" + C("advpara", "stealthrl", "contam") + ".")
SEC("Threat Model and Problem")
TXT("A screen assigns each manuscript a score and flags manuscripts above a threshold set so that a target share of human manuscripts (here 5%) is flagged. The defender knows human papers of the target fields but not which generator or attack an author used. The author holds a machine-written paper and may (i) edit counts of surface features, (ii) rewrite it with an LLM, (iii) replace a share of sentences with sentences from human papers of the same field, (iv) inject typographical noise, or (v) generate several candidates with (ii) or (iii) and submit the one a public detector scores as least machine-like. We call (v) a guided attack; the author queries Fast-DetectGPT and keeps the best of four candidates. The author does not know or query the defender's ensemble.")
SEC("Proposed Defence")
TXT("Let s_{j}(x) be the score of detector *j* for document *x*, with larger values indicating machine authorship, and let C be a set of human calibration documents used for nothing else. We standardize each score robustly,")
EQN("z_{j}(x) = (s_{j}(x) − m_{j}) / (1.4826 · MAD_{j}),", "1")
TXT("where m_{j} and MAD_{j} are the median and the median absolute deviation of s_{j} over C. The *two-sided likelihood score* replaces the mean token log-likelihood ℓ(x) by its absolute deviation from the calibration median,")
EQN("s_{dl}(x) = | ℓ(x) − median_{C} ℓ |,", "2")
TXT("so that text that is too improbable is as suspicious as text that is too probable. The *mean ensemble* averages the standardized scores of four components with different failure modes: Fast-DetectGPT (*fd*), a RoBERTa classifier (*rb*), the count screen (*sc*) and the two-sided score (*dl*):")
EQN("E(x) = ¼ [ z_{fd}(x) + z_{rb}(x) + z_{sc}(x) + z_{dl}(x) ].", "3")
TXT("The *top-2 ensemble* averages the two largest of the four standardized scores. Neither rule has fitted weights or uses attacked text; the threshold is the 95th percentile of E over C. The choice of components follows the observation that probability-based detectors fail under mixing and noise, supervised detectors under shifts of generator and under mixing, and the count screen under edits of its features, whereas the two-sided score gains under mixing (Section VI).")
SEC("Experimental Setup")
SUB("Data")
TXT("We sampled 500 human papers from 2015 to 2019, before LLMs were widely used: 330 from the open-access subset of PubMed Central and 170 from arXiv (cs.*); 474 had at least 2,000 tokens after cleaning. For each sampled paper a generator wrote a paper on the same topic from its title and abstract, with one of three prompts of increasing detail, in four requests (introduction, methods, results, discussion). The generators were GPT-4o, Claude Sonnet 5.5, Llama 3.3 70B Instruct and DeepSeek V4.1 Flash. For a field test we sampled 130 economics papers from arXiv (econ.*, 2017–2019) and generated papers from them. Documents were truncated to 2,000 whitespace tokens with markdown removed.")
SUB("Splits")
TXT("Two hundred human papers were fixed as evaluation negatives (186 have at least 2,000 tokens) and 100 further human papers as calibration set C. For each target generator, supervised detectors and the count screen were trained on the remaining human papers and on GPT-4o and Claude papers other than the target, excluding Llama (the rewriting model of the attacks), DeepSeek and every paper whose source is an evaluation paper. Attacks were applied to 100 papers per generator.")
SUB("Detectors")
TXT("The count screen counts parentheses, square brackets, semicolons, paired quotation marks, discourse markers, passive constructions and nominalizations in four 500-token windows and uses each event's rate and the mean and slope of its log Fano factor across window sizes" + C("cameron") + " (21 features, logistic regression). Log-likelihood, Binoculars and Fast-DetectGPT used Falcon-7B and Falcon-7B-instruct on the first 1,024 tokens. The TF-IDF classifier uses word unigrams and bigrams with logistic regression; the RoBERTa classifier is roberta-base fine-tuned for two epochs. Gemini 2.5 Pro served as a zero-shot judge on a subset of conditions.")
SUB("Attacks")
TXT("Feature edit: counts of five events are moved to the human mean. LLM rewrite: Llama 3.3 70B rewrites the paper as an experienced researcher would. Rewrite + edit, and rewrite twice with different wording. 50% human sentences: half the sentences are replaced by human sentences of the same field outside the evaluation set. 10% typos: each word of four or more letters receives a random typo with probability 0.1. Guided rewrite and guided mixing: four candidates per paper, the one with the lowest Fast-DetectGPT score is kept.")
SUB("Metrics")
TXT("We report the AUC against the evaluation humans with 95% intervals from 400 bootstrap resamples of documents within each class, and the share of machine papers flagged at the calibrated threshold together with the false-positive rate on the evaluation humans.")
SEC("Results")
SUB("Detection without attack")
TXT(f"Table I gives the AUC on unattacked papers. Fast-DetectGPT reaches {cell('fd','claude','orig'):.2f}–{cell('fd','gpt4o','orig'):.2f}, the supervised classifiers 0.92 or more, the count screen {min(cell('sc',g,'orig') for g,_ in GENS):.2f}–{max(cell('sc',g,'orig') for g,_ in GENS):.2f}, and the judge {jcell('claude','orig'):.2f} on Claude papers. One-sided log-likelihood is near chance on Claude papers ({cell('ll','claude','orig'):.2f}). Claude is the hardest generator for every method.")
section(1)
TAB("TABLE I\nAUC WITHOUT ATTACK (186 EVALUATION HUMANS, 100 MACHINE PAPERS PER GENERATOR; * PROPOSED COMPONENTS)", T1, widths=[2.3, 1.1, 1.1, 1.1, 1.1], size=8.5)
section(2)
SUB("Attacks and guided attacks")
TXT(f"Fig. 1 and Table II give the AUC under every attack. Guided rewriting barely affects Fast-DetectGPT ({min(cell('fd',g,'bok_rw') for g in ('claude','gpt4o','llama')):.2f} or higher). Guided mixing lowers it on Claude papers from {ci('fd','claude','orig')} to {ci('fd','claude','bok_mix')}, below chance, and on GPT-4o papers to {cell('fd','gpt4o','bok_mix'):.2f}. The attack transfers in part to RoBERTa ({cell('rb','claude','bok_mix'):.2f} on the same papers), no more than unguided mixing ({cell('rb','claude','mix50'):.2f}). The judge is not fooled ({jcell('claude','bok_mix'):.2f}).")
section(1)
FIGT("fig3_audit.png", "Fig. 1. AUC of every method under every attack, by generator. Rows below the white line are the proposed components and the judge. Blank: not run.", width=7.0)
TAB("TABLE II\nWORST CASE UNDER ATTACK: AUC WITHOUT ATTACK → LOWEST AUC OVER EIGHT ATTACKS (ATTACK)", T2, widths=[1.7, 1.35, 1.35, 1.35, 1.35], size=7.5)
section(2)
SUB("Inversion and the two-sided score")
TXT(f"Mixing in human sentences makes text less probable than genuine human text, so one-sided log-likelihood falls to {cell('ll','claude','mix50'):.2f} on Claude papers. The two-sided score (2) reaches {ci('dl','claude','bok_mix')} under guided mixing, exactly where Fast-DetectGPT fails, but is near chance on unattacked Claude papers ({cell('dl','claude','orig'):.2f}).")
SUB("Ensembles")
TXT(f"Over four generators and eight attacks the lowest AUC is {wm4:.2f} for the mean ensemble (3) and {wt2:.2f} for the top-2 ensemble, against {wfd:.2f} for Fast-DetectGPT, {wrb:.2f} for RoBERTa, {wtf:.2f} for TF-IDF and {wsc:.2f} for the count screen (Fig. 2). Without the two-sided component the mean ensemble's worst case is {wm3:.2f}. On unattacked Claude papers the mean ensemble loses little ({cell('mean4','claude','orig'):.2f} against {cell('fd','claude','orig'):.2f}).")
section(1)
FIGT("fig4_defence.png", "Fig. 2. (a) Lowest AUC over four generators and eight attacks per method; proposed components in blue. (b) AUC on Claude papers by attack for Fast-DetectGPT, the two-sided score and the mean ensemble.", width=6.8)
TAB("TABLE III\nSHARE FLAGGED AT THE CALIBRATED 5% THRESHOLD (MEAN OVER GENERATORS) AND FPR ON EVALUATION HUMANS", T3, widths=[1.7, 0.9, 0.7, 0.8, 0.8, 0.9, 0.9], size=8)
section(2)
SUB("Operating point")
TXT(f"At the calibrated threshold the mean ensemble flags {tpr('mean4','bok_mix'):.2f} of guided-mixing papers against {tpr('fd','bok_mix'):.2f} for Fast-DetectGPT, with false-positive rates of {FPR['mean4']:.3f} and {FPR['fd']:.3f} (Table III). The two-sided score alone flags {tpr('dl','bok_mix'):.2f} but with a false-positive rate of {FPR['dl']:.3f}. Averaging dilutes a strong component at strict thresholds.")
SUB("A new field and the artifact control")
TXT(f"On economics papers supervised detectors trained on medicine and computer science score near 1.00 (Table IV), but all economics human papers were converted from PDF and carry structural artifacts. A TF-IDF control (Table V) shows why this matters: on an unseen generator its AUC for Claude papers falls from {acp.loc[('unseen-generator','claude'),'raw']:.2f} to {acp.loc[('unseen-generator','claude'),'clean']:.2f} when headers, figure labels and similar tokens are removed and to {acp.loc[('unseen-generator','claude'),'function']:.2f} with function words only.")
section(1)
TAB("TABLE IV\nECONOMICS PAPERS: AUC WITHOUT ATTACK / WITH 50% HUMAN SENTENCES", T4, widths=[1.7, 1.35, 1.35, 1.35, 1.35], size=8)
TAB("TABLE V\nARTIFACT CONTROL FOR THE TF-IDF CLASSIFIER", T5, widths=[2.0, 1.0, 1.1, 1.4, 1.5], size=8)
section(2)
SEC("Discussion and Limitations")
TXT("No detector survives every attack, and different attacks hurt different detectors; this is what the defence exploits. The combination raises the cost of evasion but does not make it impossible, which theory suggests cannot be achieved for sufficiently good generators" + C("sadasivan") + ". Our attacker adapts to one public detector, not to the ensemble; RL paraphrasers trained against an ensemble already evade one that includes Fast-DetectGPT and RoBERTa" + C("stealthrl") + ". The guided attacks used four candidates; more candidates would push further. The gain is larger in AUC than at a strict operating point, and the two-sided score needs a larger calibration set than our 100 papers.")
TXT("Each cell compares 100 attacked papers with 186 human papers, so AUC differences below about 0.07 are within sampling error. The corpus covers English papers from three fields; human and machine texts were obtained differently, which our control addresses only for TF-IDF. Zero-shot scores used one open model pair, and the judge was run on a subset of attacks. We did not evaluate fairness for non-native writers" + C("bias") + ", light human editing of machine text, or peer reviews" + C("peerreview") + ". This is an independent re-implementation of an earlier unpublished study whose data were unavailable; all corpora and results were rebuilt.")
SEC("Conclusion")
TXT("Screens for machine-written scholarly text should be evaluated against authors who can query them. In our data every single detector fails under some attack, and a parameter-free combination of detectors that fail differently, anchored by a two-sided likelihood score, keeps the worst case far higher. We release code, scores and corpora to support evaluations of this kind.")
d.add_paragraph("Data Availability", style=ST("Style11"))
TXT("Code, machine-generated corpora, per-document scores and result tables: Zenodo, https://doi.org/10.5281/zenodo.23664508 (version 2). Human papers are public and are re-collected from the identifier lists provided.")
d.add_paragraph("Acknowledgment", style=ST("Style11"))
TXT("Generative AI disclosure: the machine-written papers studied here were produced by GPT-4o, Claude Sonnet 5.5, Llama 3.3 70B Instruct and DeepSeek V4.1 Flash, and Gemini 2.5 Pro served as a judge. An AI coding assistant (Claude Code) was used to write and run the analysis code and to draft and edit the manuscript text; the author reviewed the results and takes responsibility for the content.")
d.add_paragraph("References", style=ST("ReferenceHead"))
for k in order: d.add_paragraph(REF[k], style=ST("References"))
d.add_paragraph("Ook Lee is with the Department of Information Systems, Hanyang University, Seoul, Republic of Korea. [Biography to be completed.]", style=ST("Normal"))
d.save("TAI_Manuscript_template.docx")
print("words", sum(len(p.text.split()) for p in d.paragraphs), "refs", len(order), "unused", set(REF) - set(order))
